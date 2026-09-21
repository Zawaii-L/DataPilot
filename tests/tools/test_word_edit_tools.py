import tempfile
from pathlib import Path

from docx import Document

from document_tools import read_word_file
from word_edit_tools import (
    append_paragraph,
    apply_word_edits,
    delete_paragraphs,
    replace_paragraph,
    replace_text,
    update_table_cell,
)


def create_test_word(path: Path):
    document = Document()

    document.add_heading(
        "2025年度工作总结",
        level=1,
    )

    document.add_paragraph(
        "本报告总结2025年度重点工作。"
    )

    document.add_paragraph(
        "临时说明：本段后续需要删除。"
    )

    table = document.add_table(
        rows=2,
        cols=2,
    )

    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "年度"
    table.cell(1, 0).text = "DataPilot"
    table.cell(1, 1).text = "2025年度"

    document.save(
        str(path)
    )


def assert_exists(path):
    path = Path(path)

    assert path.exists()
    assert path.is_file()

    return path


def main():
    print("=" * 70)
    print("DataPilot v3.2 word_edit_tools 自动化测试")
    print("=" * 70)

    with tempfile.TemporaryDirectory() as temp_dir:
        folder = Path(temp_dir)
        source = folder / "原始报告.docx"

        create_test_word(
            source
        )

        print("\n[测试 1] 全文替换")
        result = replace_text(
            file_path=source,
            old_text="2025年度",
            new_text="2026年度",
            output_path=folder / "全文替换.docx",
        )

        output = assert_exists(
            result["output_path"]
        )
        text = read_word_file(
            output
        )

        assert "2025年度" not in text
        assert "2026年度" in text
        assert result["total_replacements"] == 3

        print("全文 + 表格替换通过。")

        print("\n[测试 2] 整段替换")
        result = replace_paragraph(
            file_path=source,
            search_text="2025年度工作总结",
            new_text="2026年度工作总结",
            output_path=folder / "整段替换.docx",
            match_mode="exact",
        )

        output = assert_exists(
            result["output_path"]
        )
        text = read_word_file(
            output
        )

        assert "2026年度工作总结" in text
        assert result["matched_paragraphs"] == 1

        print("整段替换通过。")

        print("\n[测试 3] 追加段落")
        result = append_paragraph(
            file_path=source,
            text="DataPilot 自动追加内容。",
            output_path=folder / "追加段落.docx",
        )

        output = assert_exists(
            result["output_path"]
        )
        text = read_word_file(
            output
        )

        assert "DataPilot 自动追加内容。" in text

        print("追加段落通过。")

        print("\n[测试 4] 删除指定段落")
        result = delete_paragraphs(
            file_path=source,
            search_text="临时说明",
            output_path=folder / "删除段落.docx",
        )

        output = assert_exists(
            result["output_path"]
        )
        text = read_word_file(
            output
        )

        assert "临时说明" not in text
        assert result["deleted_paragraphs"] == 1

        print("删除段落通过。")

        print("\n[测试 5] 修改表格单元格")
        result = update_table_cell(
            file_path=source,
            table_index=0,
            row_index=1,
            column_index=1,
            new_text="2026年度",
            output_path=folder / "修改表格.docx",
        )

        output = assert_exists(
            result["output_path"]
        )
        text = read_word_file(
            output
        )

        assert "DataPilot | 2026年度" in text

        print("表格单元格修改通过。")

        print("\n[测试 6] 多操作一次执行")
        result = apply_word_edits(
            file_path=source,
            output_path=folder / "综合编辑.docx",
            operations=[
                {
                    "action": "replace_text",
                    "old_text": "2025年度",
                    "new_text": "2026年度",
                },
                {
                    "action": "delete_paragraphs",
                    "search_text": "临时说明",
                },
                {
                    "action": "append_paragraph",
                    "text": "审核状态：已完成。",
                },
                {
                    "action": "update_table_cell",
                    "table_index": 0,
                    "row_index": 1,
                    "column_index": 0,
                    "new_text": "DataPilot v3.2",
                },
            ],
        )

        output = assert_exists(
            result["output_path"]
        )
        text = read_word_file(
            output
        )

        assert "2025年度" not in text
        assert "2026年度" in text
        assert "临时说明" not in text
        assert "审核状态：已完成。" in text
        assert "DataPilot v3.2 | 2026年度" in text
        assert result["operation_count"] == 4

        print("多操作连续编辑通过。")

        print("\n[测试 7] 默认不覆盖原文件")
        result = replace_text(
            file_path=source,
            old_text="2025年度",
            new_text="2027年度",
        )

        output = assert_exists(
            result["output_path"]
        )

        assert output != source.resolve()
        assert source.exists()

        original_text = read_word_file(
            source
        )

        assert "2025年度" in original_text
        assert "2027年度" not in original_text

        print("源文件保护通过。")

    print("\n" + "=" * 70)
    print("全部测试通过！")
    print()
    print("DataPilot v3.2 Word 编辑底层已具备：")
    print("1. 正文和表格全文替换")
    print("2. 整段替换")
    print("3. 追加段落")
    print("4. 删除段落")
    print("5. 表格单元格修改")
    print("6. 多操作一次执行")
    print("7. 默认另存新文件，不覆盖源文件")
    print("=" * 70)


if __name__ == "__main__":
    main()
