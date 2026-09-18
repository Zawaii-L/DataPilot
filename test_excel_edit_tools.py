import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

from excel_edit_tools import (
    apply_excel_edits,
    inspect_excel_workbook,
    read_excel_sheet_records,
)


def create_test_excel(path: Path):
    workbook = Workbook()

    sales = workbook.active
    sales.title = "销售数据"

    sales.append(
        [
            "城市",
            "部门",
            "销售额",
            "备注",
        ]
    )

    sales.append(
        [
            "珠海",
            "部门A",
            120,
            "正常",
        ]
    )

    sales.append(
        [
            "澳门",
            "部门B",
            180,
            None,
        ]
    )

    sales.append(
        [
            "中山",
            "部门A",
            90,
            "待核对",
        ]
    )

    sales.append(
        [
            "珠海",
            "部门A",
            120,
            "正常",
        ]
    )

    info = workbook.create_sheet(
        "说明"
    )

    info["A1"] = "文件说明"
    info["A2"] = "这个 Sheet 必须被保留。"

    workbook.save(
        path
    )


def main():
    print("=" * 70)
    print("DataPilot v3.2 excel_edit_tools 自动化测试")
    print("=" * 70)

    with tempfile.TemporaryDirectory() as temp_dir:
        folder = Path(
            temp_dir
        )

        source = (
            folder
            / "原始销售表.xlsx"
        )

        create_test_excel(
            source
        )

        print("\n[测试 1] 工作簿结构检查")

        info = inspect_excel_workbook(
            source
        )

        assert info["sheet_names"] == [
            "销售数据",
            "说明",
        ]

        assert (
            info["sheets"][0]["headers"]
            == [
                "城市",
                "部门",
                "销售额",
                "备注",
            ]
        )

        print("多 Sheet 结构识别通过。")

        print("\n[测试 2] 综合 Excel 编辑")

        output = (
            folder
            / "编辑后销售表.xlsx"
        )

        result = apply_excel_edits(
            file_path=source,
            output_path=output,
            operations=[
                {
                    "action": "replace_values",
                    "sheet_name": "销售数据",
                    "column": "部门",
                    "old_value": "部门A",
                    "new_value": "华南一部",
                },
                {
                    "action": "fill_missing_values",
                    "sheet_name": "销售数据",
                    "columns": ["备注"],
                    "fill_value": "无",
                },
                {
                    "action": "delete_duplicate_rows",
                    "sheet_name": "销售数据",
                    "columns": [
                        "城市",
                        "部门",
                        "销售额",
                        "备注",
                    ],
                    "keep": "first",
                },
                {
                    "action": "add_column",
                    "sheet_name": "销售数据",
                    "column_name": "审核状态",
                    "value": "已审核",
                },
                {
                    "action": "sort_rows",
                    "sheet_name": "销售数据",
                    "column": "销售额",
                    "ascending": False,
                },
                {
                    "action": "update_cell",
                    "sheet_name": "销售数据",
                    "row": 2,
                    "column_name": "审核状态",
                    "value": "重点审核",
                },
                {
                    "action": "append_rows",
                    "sheet_name": "销售数据",
                    "rows": [
                        {
                            "城市": "横琴",
                            "部门": "华南一部",
                            "销售额": 150,
                            "备注": "新增",
                            "审核状态": "已审核",
                        }
                    ],
                },
                {
                    "action": "rename_columns",
                    "sheet_name": "销售数据",
                    "rename_map": {
                        "销售额": "销售额_万元"
                    },
                },
                {
                    "action": "delete_columns",
                    "sheet_name": "销售数据",
                    "columns": [
                        "备注"
                    ],
                },
            ],
        )

        assert result["success"] is True
        assert result["operation_count"] == 9
        assert output.exists()

        print("9 项连续编辑通过。")

        print("\n[测试 3] 编辑结果回读")

        records = read_excel_sheet_records(
            output,
            sheet_name="销售数据",
        )

        assert records["headers"] == [
            "城市",
            "部门",
            "销售额_万元",
            "审核状态",
        ]

        assert records["row_count"] == 4

        rows = records["records"]

        assert rows[0][
            "销售额_万元"
        ] == 180

        assert rows[0][
            "审核状态"
        ] == "重点审核"

        assert any(
            row["城市"] == "横琴"
            and row["销售额_万元"] == 150
            for row in rows
        )

        assert all(
            row["部门"] != "部门A"
            for row in rows
        )

        assert all(
            "备注" not in row
            for row in rows
        )

        print("编辑结果确定性验证通过。")

        print("\n[测试 4] 其他 Sheet 保留")

        workbook = load_workbook(
            output,
            data_only=False,
        )

        assert workbook.sheetnames == [
            "销售数据",
            "说明",
        ]

        assert (
            workbook["说明"]["A2"].value
            == "这个 Sheet 必须被保留。"
        )

        print("未修改 Sheet 完整保留。")

        print("\n[测试 5] 源文件保护")

        original = load_workbook(
            source,
            data_only=False,
        )

        assert original.sheetnames == [
            "销售数据",
            "说明",
        ]

        assert (
            original["销售数据"]["B2"].value
            == "部门A"
        )

        assert (
            original["销售数据"]["D3"].value
            is None
        )

        assert (
            original["销售数据"].max_row
            == 5
        )

        print("原 Excel 未被覆盖。")

        print("\n[测试 6] 默认输出路径")

        result = apply_excel_edits(
            file_path=source,
            operations=[
                {
                    "action": "replace_values",
                    "sheet_name": "销售数据",
                    "column": "部门",
                    "old_value": "部门B",
                    "new_value": "华南二部",
                }
            ],
        )

        default_output = Path(
            result["output_path"]
        )

        assert default_output.exists()
        assert (
            default_output
            != source.resolve()
        )

        print("默认另存新文件通过。")

    print("\n" + "=" * 70)
    print("全部测试通过！")
    print()
    print("DataPilot v3.2 Excel 编辑底层已具备：")
    print("1. 多 Sheet 工作簿结构识别")
    print("2. 指定 Sheet 编辑")
    print("3. 批量值替换")
    print("4. 新增 / 删除 / 重命名字段")
    print("5. 数据排序")
    print("6. 重复行删除")
    print("7. 缺失值填充")
    print("8. 追加数据行")
    print("9. 指定单元格修改")
    print("10. 未修改 Sheet 保留")
    print("11. 默认另存新文件，不覆盖源文件")
    print("=" * 70)


if __name__ == "__main__":
    main()
