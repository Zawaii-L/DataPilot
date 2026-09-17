from pathlib import Path
from typing import Optional

import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


# ============================================================
# Matplotlib 中文显示
# ============================================================

plt.rcParams["font.sans-serif"] = [
    "Microsoft YaHei",
    "SimHei",
    "Arial Unicode MS",
]

plt.rcParams["axes.unicode_minus"] = False


# ============================================================
# 工具函数
# ============================================================

def _ensure_parent(path):
    """
    确保输出文件的父目录存在。
    """
    path = Path(path)

    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    return path


def _safe_text(value):
    """
    将各种数据安全转换成 Word 可写入的字符串。
    """
    if pd.isna(value):
        return ""

    return str(value)


# ============================================================
# 自动选择图表
# ============================================================

def create_office_chart(
    dataframe: pd.DataFrame,
    output_path="outputs/DataPilot_办公分析图.png",
    title="DataPilot 办公数据分析",
):
    """
    根据最终办公处理结果自动选择合适的图表。

    当前规则：

    1. 分类字段 + 数值字段
       -> 柱状图

    2. 只有多个数值字段
       -> 数值字段平均值柱状图

    3. 只有一个数值字段
       -> 行号趋势折线图

    如果无法找到数值字段，则返回 None。
    """

    if dataframe is None:
        return None

    if dataframe.empty:
        return None

    output_path = _ensure_parent(
        output_path
    )

    df = dataframe.copy()

    numeric_columns = (
        df.select_dtypes(
            include="number"
        )
        .columns
        .tolist()
    )

    non_numeric_columns = [
        column
        for column in df.columns
        if column not in numeric_columns
    ]

    if not numeric_columns:
        return None

    plt.figure(
        figsize=(10, 6)
    )

    # --------------------------------------------------------
    # 情况 1：
    # 分类字段 + 数值字段
    # --------------------------------------------------------

    if non_numeric_columns:
        category_column = (
            non_numeric_columns[0]
        )

        value_column = (
            numeric_columns[0]
        )

        chart_df = df[
            [
                category_column,
                value_column,
            ]
        ].copy()

        chart_df[value_column] = (
            pd.to_numeric(
                chart_df[value_column],
                errors="coerce",
            )
        )

        chart_df = (
            chart_df
            .dropna(
                subset=[value_column]
            )
        )

        if chart_df.empty:
            plt.close()
            return None

        # 如果分类值重复，再自动按分类求平均
        if chart_df[
            category_column
        ].duplicated().any():

            chart_df = (
                chart_df
                .groupby(
                    category_column,
                    dropna=False,
                )[value_column]
                .mean()
                .reset_index()
            )

        x_values = (
            chart_df[
                category_column
            ]
            .astype(str)
        )

        y_values = (
            chart_df[
                value_column
            ]
        )

        plt.bar(
            x_values,
            y_values,
        )

        plt.xlabel(
            category_column
        )

        plt.ylabel(
            value_column
        )

        plt.title(
            title
        )

        plt.xticks(
            rotation=30,
            ha="right",
        )

        # 显示数值标签
        for index, value in enumerate(
            y_values
        ):
            if pd.notna(value):
                plt.text(
                    index,
                    value,
                    f"{value:.2f}",
                    ha="center",
                    va="bottom",
                )

    # --------------------------------------------------------
    # 情况 2：
    # 多个数值字段
    # --------------------------------------------------------

    elif len(numeric_columns) > 1:

        selected_columns = (
            numeric_columns[:8]
        )

        averages = (
            df[selected_columns]
            .mean(
                numeric_only=True
            )
        )

        plt.bar(
            averages.index,
            averages.values,
        )

        plt.title(
            title
        )

        plt.ylabel(
            "平均值"
        )

        plt.xticks(
            rotation=30,
            ha="right",
        )

        for index, value in enumerate(
            averages.values
        ):
            if pd.notna(value):
                plt.text(
                    index,
                    value,
                    f"{value:.2f}",
                    ha="center",
                    va="bottom",
                )

    # --------------------------------------------------------
    # 情况 3：
    # 单一数值字段
    # --------------------------------------------------------

    else:
        value_column = (
            numeric_columns[0]
        )

        values = pd.to_numeric(
            df[value_column],
            errors="coerce",
        )

        plt.plot(
            range(
                1,
                len(values) + 1,
            ),
            values,
            marker="o",
        )

        plt.title(
            title
        )

        plt.xlabel(
            "记录序号"
        )

        plt.ylabel(
            value_column
        )

    plt.grid(
        alpha=0.25
    )

    plt.tight_layout()

    plt.savefig(
        output_path,
        dpi=160,
        bbox_inches="tight",
    )

    plt.close()

    return str(
        output_path
    )


# ============================================================
# Word 字体
# ============================================================

def _set_run_font(
    run,
    font_name="Microsoft YaHei",
    font_size=10.5,
    bold=False,
):
    """
    设置 Word 中英文和中文字体。
    """

    run.font.name = (
        font_name
    )

    run._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        font_name,
    )

    run.font.size = Pt(
        font_size
    )

    run.bold = bold


def _add_paragraph(
    document,
    text="",
    font_size=10.5,
    bold=False,
):
    """
    添加统一字体的普通段落。
    """

    paragraph = (
        document.add_paragraph()
    )

    run = paragraph.add_run(
        str(text)
    )

    _set_run_font(
        run,
        font_size=font_size,
        bold=bold,
    )

    return paragraph


# ============================================================
# DataFrame 写入 Word
# ============================================================

def _add_dataframe_table(
    document,
    dataframe: pd.DataFrame,
    max_rows=30,
):
    """
    将 DataFrame 写入 Word 表格。

    为避免超大文件，默认最多显示前 30 行。
    """

    if dataframe is None:
        _add_paragraph(
            document,
            "无数据。",
        )
        return

    if dataframe.empty:
        _add_paragraph(
            document,
            "无数据。",
        )
        return

    display_df = (
        dataframe
        .head(max_rows)
        .copy()
    )

    table = document.add_table(
        rows=1,
        cols=len(
            display_df.columns
        ),
    )

    table.style = (
        "Table Grid"
    )

    # 表头
    for index, column in enumerate(
        display_df.columns
    ):
        cell = (
            table.rows[0]
            .cells[index]
        )

        cell.text = ""

        paragraph = (
            cell.paragraphs[0]
        )

        run = paragraph.add_run(
            str(column)
        )

        _set_run_font(
            run,
            font_size=9,
            bold=True,
        )

    # 数据
    for _, row in display_df.iterrows():

        cells = (
            table.add_row().cells
        )

        for index, value in enumerate(
            row
        ):
            cells[index].text = ""

            paragraph = (
                cells[index]
                .paragraphs[0]
            )

            run = paragraph.add_run(
                _safe_text(value)
            )

            _set_run_font(
                run,
                font_size=9,
            )

    if len(dataframe) > max_rows:
        _add_paragraph(
            document,
            (
                f"注：结果共有 {len(dataframe)} 行，"
                f"报告仅展示前 {max_rows} 行。"
            ),
            font_size=9,
        )


# ============================================================
# Office Word 报告
# ============================================================

def generate_office_word_report(
    user_task: str,
    dataframe: pd.DataFrame,
    execution_log=None,
    source_files=None,
    chart_path: Optional[str] = None,
    output_path="outputs/DataPilot_办公分析报告.docx",
):
    """
    根据 Office Agent 的最终处理结果生成 Word 报告。
    """

    output_path = _ensure_parent(
        output_path
    )

    execution_log = (
        execution_log or []
    )

    source_files = (
        source_files or []
    )

    document = Document()

    # --------------------------------------------------------
    # 默认字体
    # --------------------------------------------------------

    normal_style = (
        document.styles["Normal"]
    )

    normal_style.font.name = (
        "Microsoft YaHei"
    )

    normal_style._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        "Microsoft YaHei",
    )

    normal_style.font.size = Pt(
        10.5
    )

    # --------------------------------------------------------
    # 标题
    # --------------------------------------------------------

    title = (
        document.add_paragraph()
    )

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = title.add_run(
        "DataPilot 办公数据分析报告"
    )

    _set_run_font(
        run,
        font_size=18,
        bold=True,
    )

    # --------------------------------------------------------
    # 1. 任务说明
    # --------------------------------------------------------

    heading = (
        document.add_heading(
            "1. 任务说明",
            level=1,
        )
    )

    for run in heading.runs:
        _set_run_font(
            run,
            font_size=14,
            bold=True,
        )

    _add_paragraph(
        document,
        user_task,
    )

    # --------------------------------------------------------
    # 2. 输入文件
    # --------------------------------------------------------

    heading = (
        document.add_heading(
            "2. 输入文件",
            level=1,
        )
    )

    for run in heading.runs:
        _set_run_font(
            run,
            font_size=14,
            bold=True,
        )

    if source_files:
        for file_path in source_files:
            _add_paragraph(
                document,
                f"• {file_path}",
            )
    else:
        _add_paragraph(
            document,
            "未记录输入文件。",
        )

    # --------------------------------------------------------
    # 3. Agent 执行步骤
    # --------------------------------------------------------

    heading = (
        document.add_heading(
            "3. Agent 执行步骤",
            level=1,
        )
    )

    for run in heading.runs:
        _set_run_font(
            run,
            font_size=14,
            bold=True,
        )

    if execution_log:

        for item in execution_log:

            step = item.get(
                "step",
                ""
            )

            action = item.get(
                "action",
                ""
            )

            rows_after = item.get(
                "rows_after",
                ""
            )

            columns_after = item.get(
                "columns_after",
                ""
            )

            operation = item.get(
                "operation",
                {},
            )

            text = (
                f"步骤 {step}：{action}"
            )

            if operation:
                text += (
                    f"；参数：{operation}"
                )

            if rows_after != "":
                text += (
                    f"；执行后 {rows_after} 行"
                )

            if columns_after != "":
                text += (
                    f"、{columns_after} 列"
                )

            _add_paragraph(
                document,
                text,
            )

    else:
        _add_paragraph(
            document,
            "未记录执行步骤。",
        )

    # --------------------------------------------------------
    # 4. 最终结果概况
    # --------------------------------------------------------

    heading = (
        document.add_heading(
            "4. 最终结果概况",
            level=1,
        )
    )

    for run in heading.runs:
        _set_run_font(
            run,
            font_size=14,
            bold=True,
        )

    _add_paragraph(
        document,
        (
            f"最终结果共有 "
            f"{len(dataframe)} 行、"
            f"{len(dataframe.columns)} 列。"
        ),
    )

    _add_paragraph(
        document,
        (
            "结果字段："
            + "、".join(
                [
                    str(column)
                    for column
                    in dataframe.columns
                ]
            )
        ),
    )

    # --------------------------------------------------------
    # 5. 数据结果
    # --------------------------------------------------------

    heading = (
        document.add_heading(
            "5. 数据结果",
            level=1,
        )
    )

    for run in heading.runs:
        _set_run_font(
            run,
            font_size=14,
            bold=True,
        )

    _add_dataframe_table(
        document,
        dataframe,
        max_rows=30,
    )

    # --------------------------------------------------------
    # 6. 图表
    # --------------------------------------------------------

    if (
        chart_path
        and Path(chart_path).exists()
    ):

        heading = (
            document.add_heading(
                "6. 数据图表",
                level=1,
            )
        )

        for run in heading.runs:
            _set_run_font(
                run,
                font_size=14,
                bold=True,
            )

        document.add_picture(
            str(chart_path),
            width=Inches(6.2),
        )

        paragraph = (
            document.paragraphs[-1]
        )

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

    # --------------------------------------------------------
    # 保存
    # --------------------------------------------------------

    document.save(
        output_path
    )

    return str(
        output_path
    )


# ============================================================
# 一次生成 Office 交付文件
# ============================================================

def generate_office_deliverables(
    user_task: str,
    dataframe: pd.DataFrame,
    execution_log=None,
    source_files=None,
    output_dir="outputs",
    need_chart=True,
    need_word_report=True,
):
    """
    一次生成：

    - PNG 图表
    - Word 报告

    返回统一结果字典。
    """

    output_dir = Path(
        output_dir
    )

    output_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    chart_path = None
    word_path = None

    if need_chart:
        chart_path = (
            create_office_chart(
                dataframe=dataframe,
                output_path=(
                    output_dir
                    / "DataPilot_办公分析图.png"
                ),
                title=(
                    "DataPilot 办公数据分析"
                ),
            )
        )

    if need_word_report:
        word_path = (
            generate_office_word_report(
                user_task=user_task,
                dataframe=dataframe,
                execution_log=(
                    execution_log
                ),
                source_files=(
                    source_files
                ),
                chart_path=chart_path,
                output_path=(
                    output_dir
                    / "DataPilot_办公分析报告.docx"
                ),
            )
        )

    return {
        "chart_path": chart_path,
        "plot_path": chart_path,
        "word_path": word_path,
    }


if __name__ == "__main__":
    print(
        "office_report_tools.py 已加载成功。"
    )