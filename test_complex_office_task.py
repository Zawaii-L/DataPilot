from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook

from agent_loop import AgentLoop


TEST_ROOT = (
    Path("outputs")
    / "complex_office_task_test"
)
SOURCE_DIR = TEST_ROOT / "待处理"
REFERENCE_DIR = TEST_ROOT / "参考资料"


def reset_test_directory() -> None:
    if TEST_ROOT.exists():
        shutil.rmtree(TEST_ROOT)

    SOURCE_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )
    REFERENCE_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )


def create_sales_workbook(
    path: Path,
    rows: list[list],
    note_text: str,
) -> None:
    workbook = Workbook()

    sales = workbook.active
    sales.title = "销售数据"
    sales.append(
        ["城市", "月份", "销售额", "订单数"]
    )

    for row in rows:
        sales.append(row)

    note = workbook.create_sheet("说明")
    note["A1"] = "销售额单位：万元"
    note["A2"] = note_text

    workbook.save(path)


def create_report(
    path: Path,
    month_text: str,
    total_text: str,
    champion_city: str,
    champion_sales: str,
    ranking_text: str,
    status_text: str,
) -> None:
    document = Document()

    document.add_heading(
        f"{month_text}销售月报",
        level=1,
    )
    document.add_paragraph(
        f"本报告汇总{month_text}三个城市的销售情况。"
    )
    document.add_paragraph(
        f"{month_text}总销售额为{total_text}。"
    )
    document.add_paragraph(
        f"销售额最高的城市是{champion_city}，"
        f"销售额为{champion_sales}。"
    )
    document.add_paragraph(
        f"城市销售额排名：{ranking_text}。"
    )
    document.add_paragraph(
        f"报告状态：{status_text}。"
    )

    table = document.add_table(
        rows=4,
        cols=2,
    )
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "内容"
    table.cell(1, 0).text = "报告月份"
    table.cell(1, 1).text = month_text
    table.cell(2, 0).text = "总销售额"
    table.cell(2, 1).text = total_text
    table.cell(3, 0).text = "销售冠军"
    table.cell(3, 1).text = champion_city

    document.save(path)


def create_source_files() -> dict[str, Path]:
    reset_test_directory()

    draft_excel = (
        SOURCE_DIR
        / "业务资料_A.xlsx"
    )
    official_excel = (
        SOURCE_DIR
        / "业务资料_B.xlsx"
    )
    unrelated_excel = (
        SOURCE_DIR
        / "业务资料_C.xlsx"
    )
    source_word = (
        SOURCE_DIR
        / "销售月报.docx"
    )
    reference_word = (
        REFERENCE_DIR
        / "销售月报_参考.docx"
    )

    create_sales_workbook(
        draft_excel,
        [
            ["珠海", "2026-10", 205, 24],
            ["澳门", "2026-10", 198, 23],
            ["横琴", "2026-10", 176, 20],
        ],
        (
            "内部预测草稿：2026-10 尚未审核批准，"
            "不得用于正式月报。"
        ),
    )

    create_sales_workbook(
        official_excel,
        [
            ["珠海", "2026-09", 128, 16],
            ["澳门", "2026-09", 186, 21],
            ["横琴", "2026-09", 154, 18],
        ],
        (
            "财务已审核：2026-09 为当前正式月报"
            "唯一批准数据源。"
        ),
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "客户资料"
    worksheet.append(
        ["客户名称", "城市", "联系人"]
    )
    worksheet.append(
        ["客户A", "珠海", "张三"]
    )
    worksheet.append(
        ["客户B", "澳门", "李四"]
    )
    workbook.save(unrelated_excel)

    create_report(
        source_word,
        month_text="2026年8月",
        total_text="390万元",
        champion_city="珠海",
        champion_sales="150万元",
        ranking_text="珠海第一、澳门第二、横琴第三",
        status_text="待更新",
    )

    create_report(
        reference_word,
        month_text="2026年7月",
        total_text="360万元",
        champion_city="珠海",
        champion_sales="140万元",
        ranking_text="珠海第一、澳门第二、横琴第三",
        status_text="已归档",
    )

    return {
        "draft_excel": draft_excel,
        "official_excel": official_excel,
        "unrelated_excel": unrelated_excel,
        "source_word": source_word,
        "reference_word": reference_word,
    }


def read_excel_rows(
    excel_path: Path,
    sheet_name: str = "销售数据",
) -> list[tuple]:
    workbook = load_workbook(
        excel_path,
        data_only=False,
    )
    worksheet = workbook[sheet_name]

    return list(
        worksheet.iter_rows(
            values_only=True,
        )
    )


def read_word_text(
    word_path: Path,
) -> str:
    document = Document(word_path)
    parts = []

    for paragraph in document.paragraphs:
        value = paragraph.text.strip()
        if value:
            parts.append(value)

    for table in document.tables:
        for row in table.rows:
            parts.append(
                " | ".join(
                    cell.text.strip()
                    for cell in row.cells
                )
            )

    return "\n".join(parts)


def find_new_words(
    excluded_paths: set[Path],
) -> list[Path]:
    excluded = {
        path.resolve()
        for path in excluded_paths
    }

    results = []

    for path in TEST_ROOT.rglob("*.docx"):
        if path.resolve() not in excluded:
            results.append(path)

    return sorted(
        results,
        key=lambda item: item.stat().st_mtime,
        reverse=True,
    )


def find_new_excels(
    excluded_paths: set[Path],
) -> list[Path]:
    excluded = {
        path.resolve()
        for path in excluded_paths
    }

    results = []

    for path in TEST_ROOT.rglob("*.xlsx"):
        if path.resolve() not in excluded:
            results.append(path)

    return sorted(
        results,
        key=lambda item: item.stat().st_mtime,
        reverse=True,
    )


def workbook_contains_expected_summary(
    path: Path,
) -> bool:
    workbook = load_workbook(
        path,
        data_only=False,
    )

    values = []

    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows(
            values_only=True,
        ):
            values.extend(
                value
                for value in row
                if value is not None
            )

    text = "\n".join(
        str(value)
        for value in values
    )

    required_text_values = [
        "2026-09",
        "澳门",
        "横琴",
        "珠海",
    ]

    if not all(
        value in text
        for value in required_text_values
    ):
        return False

    numeric_values = {
        value
        for value in values
        if isinstance(
            value,
            (int, float),
        )
    }

    required_numbers = {
        128,
        154,
        186,
        468,
    }

    return required_numbers.issubset(
        numeric_values
    )


def main() -> None:
    print("=" * 70)
    print(
        "DataPilot v3.5 复杂办公任务第五阶段基准测试"
    )
    print("=" * 70)

    paths = create_source_files()

    draft_excel = paths["draft_excel"]
    official_excel = paths["official_excel"]
    unrelated_excel = paths["unrelated_excel"]
    source_word = paths["source_word"]
    reference_word = paths["reference_word"]

    source_snapshots = {
        "draft_excel": read_excel_rows(
            draft_excel
        ),
        "official_excel": read_excel_rows(
            official_excel
        ),
        "source_word": read_word_text(
            source_word
        ),
        "reference_word": read_word_text(
            reference_word
        ),
    }

    print()
    print("测试工作区：", TEST_ROOT.resolve())
    print()
    print("本阶段要求 Agent 一次完成两个交付物：")
    print(
        "1. 根据正式批准数据生成/更新销售月报 Word"
    )
    print(
        "2. 同时生成一份销售汇总 Excel"
    )
    print(
        "3. 两个交付物必须使用同一个正式数据源"
    )
    print(
        "4. 两个交付物生成后都必须自主回读核验"
    )
    print()

    task = (
        "帮我把这个工作区的销售资料整理好。"
        "这里有几份名字很像的业务资料，其中有的数据月份更晚，"
        "但正式结果只能使用已经审核批准的数据。"
        "你自己读取文件内容和说明，判断本次应该采用哪份正式数据源。"
        "然后完成两个交付物："
        "第一，把销售月报更新成正式数据；"
        "第二，再做一份销售汇总Excel，"
        "里面要能看出月份、三个城市的销售额和订单数、"
        "总销售额、销售冠军以及城市销售排名。"
        "两个交付物必须使用同一份正式数据，"
        "未批准草稿、参考资料和其他无关文件不要改，"
        "所有原文件都不要覆盖。"
        "完成后你自己重新检查两个新文件，"
        "确认Word和Excel中的关键数字彼此一致，"
        "也和正式数据源一致。"
    )

    print("用户任务：")
    print(task)
    print()
    print(
        "注意：本阶段不提供 input_paths，"
        "而且要求 Agent 自主规划 Word + Excel 两个交付物。"
    )

    print()
    print("=" * 70)
    print("开始运行现有 DataPilot Agent Loop")
    print("=" * 70)

    agent = AgentLoop(
        max_iterations=18,
    )

    result = agent.run(
        task,
        context={
            "working_directory": str(
                TEST_ROOT.resolve()
            ),
        },
    )

    print()
    print("=" * 70)
    print("Agent Loop 返回结果")
    print("=" * 70)
    print("SUCCESS =", result.success)
    print("STOP_REASON =", result.stop_reason)
    print("ITERATIONS =", result.iterations)
    print(
        "TOOL_COUNT =",
        len(result.tool_results),
    )

    print()
    print("工具调用链：")

    for index, item in enumerate(
        result.tool_results,
        start=1,
    ):
        print(
            f"{index}. "
            f"{item.tool_name} "
            f"SUCCESS={item.success}"
        )

    print()
    print("Agent 最终回答：")
    print(result.final_answer)

    if not result.success:
        raise AssertionError(
            "Agent 没有完成第五阶段多交付物任务。"
        )

    new_words = find_new_words(
        {
            source_word,
            reference_word,
        }
    )
    new_excels = find_new_excels(
        {
            draft_excel,
            official_excel,
            unrelated_excel,
        }
    )

    if not new_words:
        raise AssertionError(
            "Agent 没有生成新的 Word 月报。"
        )

    if not new_excels:
        raise AssertionError(
            "Agent 没有生成新的 Excel 汇总表。"
        )

    generated_word = new_words[0]

    matching_excels = [
        path
        for path in new_excels
        if workbook_contains_expected_summary(
            path
        )
    ]

    if not matching_excels:
        raise AssertionError(
            "Agent 虽然生成了 Excel，"
            "但没有找到包含完整正式汇总结果的新 Excel。"
        )

    generated_excel = matching_excels[0]

    print()
    print(
        "检测到新 Word：",
        generated_word.resolve(),
    )
    print(
        "检测到新 Excel：",
        generated_excel.resolve(),
    )

    word_text = read_word_text(
        generated_word
    )

    print()
    print("新 Word 内容：")
    print("-" * 70)
    print(word_text)
    print("-" * 70)

    print()
    print("开始确定性业务验证……")

    required_word_values = [
        "2026年9月",
        "468万元",
        "澳门",
        "186万元",
        "澳门第一",
        "横琴第二",
        "珠海第三",
    ]

    for value in required_word_values:
        if value not in word_text:
            raise AssertionError(
                f"新 Word 缺少正式业务结果：{value}"
            )

    forbidden_word_values = [
        "2026年10月",
        "579万元",
        "205万元",
    ]

    for value in forbidden_word_values:
        if value in word_text:
            raise AssertionError(
                f"新 Word 错误使用未批准预测数据：{value}"
            )

    if (
        read_excel_rows(draft_excel)
        != source_snapshots["draft_excel"]
    ):
        raise AssertionError(
            "Agent 修改了未批准预测草稿。"
        )

    if (
        read_excel_rows(official_excel)
        != source_snapshots["official_excel"]
    ):
        raise AssertionError(
            "Agent 修改了正式数据源。"
        )

    if (
        read_word_text(source_word)
        != source_snapshots["source_word"]
    ):
        raise AssertionError(
            "Agent 覆盖或修改了源月报。"
        )

    if (
        read_word_text(reference_word)
        != source_snapshots["reference_word"]
    ):
        raise AssertionError(
            "Agent 修改了参考月报。"
        )

    if not unrelated_excel.exists():
        raise AssertionError(
            "Agent 删除了无关业务文件。"
        )

    successful_results = [
        item
        for item in result.tool_results
        if item.success
    ]
    successful_tools = [
        item.tool_name
        for item in successful_results
    ]

    discovery_tools = {
        "discover_data_files",
        "scan_document_files",
        "discover_files",
        "find_files",
        "list_files",
        "search_files",
        "inspect_data_files",
        "inspect_documents",
    }

    if not any(
        name in discovery_tools
        for name in successful_tools
    ):
        raise AssertionError(
            "Agent 没有自主发现工作区文件。"
        )

    note_reads = [
        item
        for item in successful_results
        if (
            item.tool_name
            == "read_office_data"
            and str(
                item.arguments.get(
                    "sheet_name",
                    "",
                )
            )
            == "说明"
        )
    ]

    note_paths = {
        str(
            Path(
                item.arguments.get(
                    "file_path",
                    ""
                )
            ).resolve()
        )
        for item in note_reads
    }

    required_note_paths = {
        str(draft_excel.resolve()),
        str(official_excel.resolve()),
    }

    if not required_note_paths.issubset(
        note_paths
    ):
        raise AssertionError(
            "Agent 没有读取两个候选 Excel 的说明 Sheet，"
            "无法证明其根据审核证据选择正式数据源。"
        )

    word_output_tools = {
        "apply_word_edits",
        "generate_document_summary_report",
    }

    excel_output_tools = {
        "apply_excel_edits",
        "export_excel",
        "export_multi_sheet_excel",
        "generate_excel",
        "generate_excel_report",
        "create_excel_report",
    }

    word_output_indexes = [
        index
        for index, item in enumerate(
            result.tool_results
        )
        if (
            item.success
            and item.tool_name
            in word_output_tools
        )
    ]

    excel_output_indexes = [
        index
        for index, item in enumerate(
            result.tool_results
        )
        if (
            item.success
            and item.tool_name
            in excel_output_tools
        )
    ]

    if not word_output_indexes:
        raise AssertionError(
            "Agent 没有调用 Word 输出工具。"
        )

    if not excel_output_indexes:
        raise AssertionError(
            "Agent 没有调用 Excel 输出工具。"
        )

    word_output_index = word_output_indexes[-1]
    excel_output_index = excel_output_indexes[-1]

    word_reread = any(
        item.success
        and item.tool_name == "read_document"
        for item in result.tool_results[
            word_output_index + 1:
        ]
    )

    if not word_reread:
        raise AssertionError(
            "Agent 生成 Word 后没有重新读取 Word 核验。"
        )

    generated_excel_resolved = str(
        generated_excel.resolve()
    )

    excel_reread = any(
        item.success
        and item.tool_name
        in {
            "read_office_data",
            "inspect_data_files",
        }
        and (
            generated_excel_resolved
            in str(
                item.arguments.get(
                    "file_path",
                    "",
                )
            )
            or generated_excel_resolved
            in str(
                item.arguments.get(
                    "file_paths",
                    "",
                )
            )
        )
        for item in result.tool_results[
            excel_output_index + 1:
        ]
    )

    if not excel_reread:
        raise AssertionError(
            "Agent 生成 Excel 后没有重新读取或检查 Excel 核验。"
        )

    print()
    print("=" * 70)
    print(
        "DataPilot v3.5 第五阶段多交付物基准测试通过！"
    )
    print("=" * 70)
    print()
    print("已验证：")
    print("1. 用户只给一个自然语言办公目标")
    print("2. Agent 自主发现并甄别候选文件")
    print("3. Agent 根据审核说明选择正式数据源")
    print("4. Agent 排除月份更晚但未批准的草稿")
    print("5. Agent 自主规划 Word + Excel 两个交付物")
    print("6. 两个交付物均使用 2026-09 正式数据")
    print("7. Word 与 Excel 的关键业务数字一致")
    print("8. 所有源文件均未覆盖")
    print("9. Agent 生成后自主回读 Word")
    print("10. Agent 生成后自主回读/检查 Excel")
    print("=" * 70)


if __name__ == "__main__":
    main()
