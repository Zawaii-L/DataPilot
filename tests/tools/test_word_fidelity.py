from __future__ import annotations

import shutil
from copy import copy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from word_edit_tools import apply_word_edits


TEST_ROOT = Path("outputs") / "word_fidelity_test"
SOURCE_WORD = TEST_ROOT / "word_fidelity_source.docx"
OUTPUT_WORD = TEST_ROOT / "word_fidelity_output.docx"


def reset_test_directory() -> None:
    if TEST_ROOT.exists():
        shutil.rmtree(TEST_ROOT)
    TEST_ROOT.mkdir(parents=True, exist_ok=True)


def run_signature(run):
    color = None
    if run.font.color is not None and run.font.color.rgb is not None:
        color = str(run.font.color.rgb)

    return {
        "text": run.text,
        "bold": run.bold,
        "italic": run.italic,
        "underline": run.underline,
        "font_name": run.font.name,
        "font_size": run.font.size.pt if run.font.size else None,
        "color": color,
    }


def paragraph_signature(paragraph):
    return {
        "style": paragraph.style.name if paragraph.style else None,
        "alignment": paragraph.alignment,
        "runs": [run_signature(run) for run in paragraph.runs],
    }


def create_source_word() -> None:
    document = Document()

    title = document.add_paragraph()
    title.style = document.styles["Title"]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    title_run = title.add_run("2026年8月销售月报")
    title_run.bold = True
    title_run.font.name = "微软雅黑"
    title_run.font.size = Pt(20)
    title_run.font.color.rgb = RGBColor(31, 78, 121)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    run1 = paragraph.add_run("本月销售冠军是")
    run1.font.name = "宋体"
    run1.font.size = Pt(11)

    run2 = paragraph.add_run("珠海")
    run2.bold = True
    run2.font.name = "微软雅黑"
    run2.font.size = Pt(14)
    run2.font.color.rgb = RGBColor(192, 0, 0)

    run3 = paragraph.add_run("，销售额为")
    run3.font.name = "宋体"
    run3.font.size = Pt(11)

    run4 = paragraph.add_run("150万元")
    run4.italic = True
    run4.font.name = "Calibri"
    run4.font.size = Pt(12)
    run4.font.color.rgb = RGBColor(0, 112, 192)

    run5 = paragraph.add_run("。")
    run5.font.name = "宋体"
    run5.font.size = Pt(11)

    table = document.add_table(rows=2, cols=2)
    table.style = "Table Grid"

    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "内容"
    table.cell(1, 0).text = "销售冠军"

    value_cell = table.cell(1, 1)
    value_paragraph = value_cell.paragraphs[0]
    value_paragraph.clear()

    value_run = value_paragraph.add_run("珠海")
    value_run.bold = True
    value_run.font.name = "微软雅黑"
    value_run.font.size = Pt(12)
    value_run.font.color.rgb = RGBColor(112, 48, 160)

    document.save(SOURCE_WORD)


def main() -> None:
    print("=" * 70)
    print("DataPilot v3.4 Word 高保真编辑基准测试")
    print("=" * 70)

    reset_test_directory()
    create_source_word()

    source_document = Document(SOURCE_WORD)

    source_title = paragraph_signature(source_document.paragraphs[0])
    source_body = paragraph_signature(source_document.paragraphs[1])
    source_table_style = source_document.tables[0].style.name
    source_table_value = paragraph_signature(
        source_document.tables[0].cell(1, 1).paragraphs[0]
    )

    print()
    print("[测试 1] 单 Run 文本替换保持字体格式")

    result = apply_word_edits(
        str(SOURCE_WORD),
        operations=[
            {
                "action": "replace_text",
                "old_text": "2026年8月",
                "new_text": "2026年9月",
                "include_tables": True,
            },
            {
                "action": "replace_text",
                "old_text": "珠海",
                "new_text": "澳门",
                "include_tables": True,
            },
            {
                "action": "replace_text",
                "old_text": "150万元",
                "new_text": "186万元",
                "include_tables": True,
            },
        ],
        output_path=str(OUTPUT_WORD),
    )

    if not result.get("success"):
        raise AssertionError("Word 编辑执行失败。")

    output_document = Document(OUTPUT_WORD)
    output_title = paragraph_signature(output_document.paragraphs[0])
    output_body = paragraph_signature(output_document.paragraphs[1])
    output_table_value = paragraph_signature(
        output_document.tables[0].cell(1, 1).paragraphs[0]
    )

    if output_title["style"] != source_title["style"]:
        raise AssertionError("标题段落样式被破坏。")

    if output_title["alignment"] != source_title["alignment"]:
        raise AssertionError("标题对齐方式被破坏。")

    if output_title["runs"][0]["bold"] != source_title["runs"][0]["bold"]:
        raise AssertionError("标题加粗格式被破坏。")

    if output_title["runs"][0]["font_size"] != source_title["runs"][0]["font_size"]:
        raise AssertionError("标题字号被破坏。")

    if output_title["runs"][0]["color"] != source_title["runs"][0]["color"]:
        raise AssertionError("标题颜色被破坏。")

    if output_body["runs"][1]["text"] != "澳门":
        raise AssertionError("正文销售冠军没有正确替换。")

    for key in ("bold", "italic", "underline", "font_name", "font_size", "color"):
        if output_body["runs"][1][key] != source_body["runs"][1][key]:
            raise AssertionError(f"正文冠军 Run 的 {key} 格式被破坏。")

    if output_body["runs"][3]["text"] != "186万元":
        raise AssertionError("正文销售额没有正确替换。")

    for key in ("bold", "italic", "underline", "font_name", "font_size", "color"):
        if output_body["runs"][3][key] != source_body["runs"][3][key]:
            raise AssertionError(f"正文销售额 Run 的 {key} 格式被破坏。")

    if output_document.tables[0].style.name != source_table_style:
        raise AssertionError("Word 表格样式被破坏。")

    if output_table_value["runs"][0]["text"] != "澳门":
        raise AssertionError("表格销售冠军没有正确替换。")

    for key in ("bold", "italic", "underline", "font_name", "font_size", "color"):
        if output_table_value["runs"][0][key] != source_table_value["runs"][0][key]:
            raise AssertionError(f"表格单元格 Run 的 {key} 格式被破坏。")

    print("单 Run 替换格式保持通过。")

    print()
    print("[测试 2] 跨 Run 文本替换保持未修改文本的格式边界")

    cross_source = TEST_ROOT / "cross_run_source.docx"
    cross_output = TEST_ROOT / "cross_run_output.docx"

    document = Document()
    paragraph = document.add_paragraph()

    first = paragraph.add_run("销售额最高的城市是")
    first.font.name = "宋体"
    first.font.size = Pt(11)

    second = paragraph.add_run("珠")
    second.bold = True
    second.font.name = "微软雅黑"
    second.font.color.rgb = RGBColor(192, 0, 0)

    third = paragraph.add_run("海")
    third.bold = True
    third.font.name = "微软雅黑"
    third.font.color.rgb = RGBColor(192, 0, 0)

    fourth = paragraph.add_run("，销售额为150万元。")
    fourth.font.name = "宋体"
    fourth.font.size = Pt(11)

    document.save(cross_source)

    before = Document(cross_source)
    before_paragraph = before.paragraphs[0]
    before_first = run_signature(before_paragraph.runs[0])
    before_fourth = run_signature(before_paragraph.runs[3])

    apply_word_edits(
        str(cross_source),
        operations=[
            {
                "action": "replace_text",
                "old_text": "珠海",
                "new_text": "澳门",
                "include_tables": True,
            }
        ],
        output_path=str(cross_output),
    )

    after = Document(cross_output)
    after_paragraph = after.paragraphs[0]

    if after_paragraph.text != "销售额最高的城市是澳门，销售额为150万元。":
        raise AssertionError(
            "跨 Run 替换后的正文文本不正确："
            + after_paragraph.text
        )

    non_empty_runs = [
        run_signature(run)
        for run in after_paragraph.runs
        if run.text
    ]

    if not non_empty_runs:
        raise AssertionError("跨 Run 替换后所有 Run 均为空。")

    if non_empty_runs[0]["text"] != "销售额最高的城市是":
        raise AssertionError(
            "跨 Run 替换把前方未修改文本合并或破坏了。"
        )

    for key in ("font_name", "font_size", "bold", "italic", "color"):
        if non_empty_runs[0][key] != before_first[key]:
            raise AssertionError(
                f"跨 Run 替换破坏了前方未修改文本的 {key} 格式。"
            )

    if non_empty_runs[-1]["text"] != "，销售额为150万元。":
        raise AssertionError(
            "跨 Run 替换把后方未修改文本合并或破坏了。"
        )

    for key in ("font_name", "font_size", "bold", "italic", "color"):
        if non_empty_runs[-1][key] != before_fourth[key]:
            raise AssertionError(
                f"跨 Run 替换破坏了后方未修改文本的 {key} 格式。"
            )

    replacement_runs = [
        item
        for item in non_empty_runs
        if "澳门" in item["text"]
    ]

    if not replacement_runs:
        raise AssertionError("没有找到跨 Run 替换后的“澳门”。")

    if replacement_runs[0]["bold"] is not True:
        raise AssertionError(
            "跨 Run 替换没有继承原目标文本的加粗格式。"
        )

    if replacement_runs[0]["color"] != "C00000":
        raise AssertionError(
            "跨 Run 替换没有继承原目标文本的颜色。"
        )

    print("跨 Run 替换格式保持通过。")

    print()
    print("[测试 3] 源文件保护")

    source_check = Document(SOURCE_WORD)
    if "2026年8月销售月报" not in source_check.paragraphs[0].text:
        raise AssertionError("源 Word 被覆盖。")

    if "珠海" not in source_check.paragraphs[1].text:
        raise AssertionError("源 Word 正文被覆盖。")

    print("源文件保护通过。")

    print()
    print("=" * 70)
    print("DataPilot v3.4 Word 高保真编辑基准测试通过！")
    print("=" * 70)


if __name__ == "__main__":
    main()
