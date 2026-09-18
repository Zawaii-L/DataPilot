from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from docx.oxml.ns import qn

from professional_word_report_tools import (
    create_professional_word_report,
    inspect_professional_word_report,
)


def assert_true(condition, message):
    if not condition:
        raise AssertionError(message)


def main():
    print("=" * 72)
    print("DataPilot v5.0 Professional Word Visual Baseline 测试")
    print("=" * 72)

    with TemporaryDirectory(prefix="datapilot_word_visual_") as temp_dir:
        output_path = Path(temp_dir) / "visual_report.docx"

        result = create_professional_word_report(
            output_path=str(output_path),
            report_title="销售数据城市汇总分析汇报",
            subtitle="按城市汇总销售额 · 销售冠军 KPI · 统计期间 2026-09",
            metadata={
                "报告对象": "公司管理层",
                "数据来源": "销售数据.xlsx（工作表：销售数据）",
                "统计期间": "2026-09",
                "统计维度": "城市（珠海、澳门、横琴）",
            },
            executive_summary=(
                "本报告基于真实销售记录形成城市汇总。三城合计销售额 468，"
                "澳门以 186 位居第一，横琴为 154，珠海为 128。"
            ),
            kpis=[
                {"label": "销售冠军城市", "value": "澳门"},
                {"label": "冠军销售额", "value": 186},
                {"label": "冠军订单数", "value": 21},
                {"label": "冠军销售额占比", "value": "39.74%"},
                {"label": "冠军平均单笔订单金额", "value": "8.86"},
                {"label": "三城合计销售额", "value": 468},
                {"label": "三城合计订单数", "value": 55},
            ],
            sections=[
                {
                    "title": "一、分析口径与数据范围",
                    "type": "paragraphs",
                    "content": [
                        "统计口径：按城市对销售额与订单数分别求和。",
                        "本期数据覆盖珠海、澳门、横琴。",
                    ],
                },
                {
                    "title": "二、城市销售额汇总表",
                    "type": "table",
                    "columns": [
                        "排名", "城市", "销售额",
                        "订单数", "销售额占比", "平均单笔订单金额",
                    ],
                    "rows": [
                        {
                            "排名": 1, "城市": "澳门", "销售额": 186,
                            "订单数": 21, "销售额占比": "39.74%",
                            "平均单笔订单金额": "8.86",
                        },
                        {
                            "排名": 2, "城市": "横琴", "销售额": 154,
                            "订单数": 18, "销售额占比": "32.91%",
                            "平均单笔订单金额": "8.56",
                        },
                        {
                            "排名": 3, "城市": "珠海", "销售额": 128,
                            "订单数": 16, "销售额占比": "27.35%",
                            "平均单笔订单金额": "8.00",
                        },
                        {
                            "排名": "合计", "城市": "三城合计", "销售额": 468,
                            "订单数": 55, "销售额占比": "100.00%",
                            "平均单笔订单金额": "8.51",
                        },
                    ],
                },
            ],
            source_note="数据来源：销售数据.xlsx；仅只读使用。",
        )

        print("\n测试 1：Visual Baseline 文件真实生成")
        assert_true(result["success"] is True and output_path.exists(), result)
        print("PASS")

        document = Document(str(output_path))

        print("\n测试 2：7 个 KPI 自动变成两行、每行最多 4 列")
        assert_true(len(document.tables) >= 2, "缺少 KPI 表或业务表。")
        kpi_table = document.tables[1] if len(document.tables) >= 3 else document.tables[0]
        # metadata is the first table; KPI is second when metadata exists
        if "报告对象" in document.tables[0].cell(0, 0).text:
            kpi_table = document.tables[1]
        assert_true(len(kpi_table.rows) == 2, f"KPI 行数错误：{len(kpi_table.rows)}")
        assert_true(len(kpi_table.columns) == 4, f"KPI 列数错误：{len(kpi_table.columns)}")
        print("PASS")

        print("\n测试 3：KPI 内容完整，没有因换行布局丢失")
        kpi_text = " ".join(
            cell.text
            for row in kpi_table.rows
            for cell in row.cells
        )
        for expected in (
            "销售冠军城市", "澳门", "冠军销售额", "186",
            "冠军订单数", "21", "39.74%", "8.86", "468", "55",
        ):
            assert_true(expected in kpi_text, f"KPI 缺失：{expected}")
        print("PASS")

        print("\n测试 4：元数据使用紧凑网格而非超长单行")
        metadata_table = document.tables[0]
        assert_true(
            "报告对象" in metadata_table.cell(0, 0).text,
            "元数据表未生成。",
        )
        assert_true(
            len(metadata_table.columns) == 2,
            "元数据应使用两列网格。",
        )
        print("PASS")

        print("\n测试 5：业务表表头设置为跨页重复")
        business_table = document.tables[-1]
        tr_pr = business_table.rows[0]._tr.get_or_add_trPr()
        repeat = tr_pr.find(qn("w:tblHeader"))
        assert_true(repeat is not None, "业务表表头未设置重复。")
        print("PASS")

        print("\n测试 6：业务表各行禁止跨页拆分")
        for row in business_table.rows:
            tr_pr = row._tr.get_or_add_trPr()
            assert_true(
                tr_pr.find(qn("w:cantSplit")) is not None,
                "发现允许跨页拆分的业务表行。",
            )
        print("PASS")

        print("\n测试 7：Heading 设置 keep-with-next")
        headings = [
            p for p in document.paragraphs
            if getattr(p.style, "name", "").startswith("Heading")
        ]
        assert_true(headings, "未找到 Heading。")
        for paragraph in headings:
            p_pr = paragraph._p.get_or_add_pPr()
            assert_true(
                p_pr.find(qn("w:keepNext")) is not None,
                f"Heading 未 keep-with-next：{paragraph.text}",
            )
        print("PASS")

        print("\n测试 8：Visual profile 元数据返回正确")
        assert_true(
            result.get("visual_profile")
            == "v5_0_professional_word_baseline",
            result,
        )
        assert_true(result.get("kpi_columns_per_row") == 4, result)
        print("PASS")

        print("\n测试 9：inspect 仍可重新读取升级后的 Word")
        inspection = inspect_professional_word_report(str(output_path))
        assert_true(inspection["success"] is True, inspection)
        assert_true("执行摘要" in inspection["headings"], inspection["headings"])
        assert_true("二、城市销售额汇总表" in inspection["headings"], inspection["headings"])
        print("PASS")

        print("\n测试 10：页眉 DataPilot 与 portrait 页面保持")
        assert_true(
            inspection["sections"][0]["header_text"] == "DataPilot",
            inspection["sections"],
        )
        assert_true(
            inspection["sections"][0]["orientation"] == "portrait",
            inspection["sections"],
        )
        print("PASS")

    print("\n" + "=" * 72)
    print("Professional Word Visual Baseline：10/10 PASS")
    print("=" * 72)


if __name__ == "__main__":
    main()
