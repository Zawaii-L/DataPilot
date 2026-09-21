from __future__ import annotations

import tempfile
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from excel_report_tools import create_professional_excel_report


def assert_true(value, message):
    if not value:
        raise AssertionError(message)


def main():
    print("=" * 72)
    print("DataPilot v4.5 Professional Excel Visual Baseline 测试")
    print("=" * 72)

    with tempfile.TemporaryDirectory(
        prefix="datapilot_excel_visual_"
    ) as temp_dir:
        output_path = Path(temp_dir) / "城市销售额分析报告.xlsx"

        df = pd.DataFrame(
            {
                "城市": ["澳门", "横琴", "珠海"],
                "销售额合计": [186, 154, 128],
            }
        )

        create_professional_excel_report(
            output_path=output_path,
            dataframe=df,
            default_sheet_name="城市销售额汇总",
            report_title="销售数据按城市汇总报告",
            subtitle=(
                "数据来源：销售数据.xlsx（工作表：销售数据）｜"
                "汇总口径：按城市合计销售额，降序排列"
            ),
            kpis=[
                {"label": "销售冠军城市", "value": "澳门"},
                {"label": "冠军城市销售额", "value": 186},
                {"label": "城市销售额合计", "value": 468},
            ],
            charts=[
                {
                    "type": "bar",
                    "sheet_name": "城市销售额汇总",
                    "category_column": "城市",
                    "value_column": "销售额合计",
                    "title": "城市销售额对比",
                }
            ],
        )

        wb = load_workbook(output_path)
        ws = wb["城市销售额汇总"]
        chart = ws._charts[0]

        print("\n测试 1：KPI 名称列宽与自动换行")
        for col in ("A", "C", "E"):
            assert_true(
                ws.column_dimensions[col].width >= 18,
                f"{col} 列不足以容纳 KPI 名称。",
            )
        for ref in ("A4", "C4", "E4"):
            assert_true(
                ws[ref].alignment.wrap_text is True,
                f"{ref} KPI 未启用自动换行。",
            )
        print("PASS")

        print("\n测试 2：数据表实体填充与清晰边框")
        for ref in ("A6", "B6", "A7", "B7"):
            cell = ws[ref]
            assert_true(cell.fill.fill_type == "solid", f"{ref} 填充不是实体。")
            assert_true(cell.border.left.style == "thin", f"{ref} 左边框缺失。")
            assert_true(cell.border.right.style == "thin", f"{ref} 右边框缺失。")
        print("PASS")

        print("\n测试 3：柱状图不自动多色")
        assert_true(chart.varyColors is False, "柱状图仍会自动使用多种颜色。")
        assert_true(len(chart.series) == 1, "城市销售额对比应只有一个数据系列。")
        print("PASS")

        print("\n测试 4：柱状图不显示冗余图例")
        assert_true(chart.legend is None, "单系列柱状图仍显示图例。")
        print("PASS")

        print("\n测试 5：不显示坐标轴标题")
        assert_true(chart.x_axis.title is None, "横轴标题仍存在。")
        assert_true(chart.y_axis.title is None, "纵轴标题仍存在。")
        print("PASS")

        print("\n测试 6：关闭纵轴主网格线")
        assert_true(
            chart.y_axis.majorGridlines is None,
            "纵轴横向网格线仍存在。",
        )
        print("PASS")

        print("\n测试 7：柱顶只显示纯数值")
        assert_true(chart.dLbls is not None, "缺少数据标签配置。")
        assert_true(chart.dLbls.showVal is True, "没有显示柱顶数值。")
        assert_true(
            chart.dLbls.showSerName is False,
            "柱顶不应显示系列名。",
        )
        assert_true(
            chart.dLbls.showCatName is False,
            "柱顶不应重复显示城市名。",
        )
        assert_true(
            chart.dLbls.showLegendKey is False,
            "柱顶不应显示图例键。",
        )
        assert_true(
            chart.dLbls.numFmt == "0",
            "柱顶数值应使用整数格式。",
        )
        print("PASS")

        print("\n测试 8：横轴分类标签明确保持可见")
        assert_true(
            chart.x_axis.delete is False,
            "横轴被隐藏，城市名称将无法显示。",
        )
        assert_true(
            chart.x_axis.tickLblPos == "low",
            "城市名称没有固定在横轴底部。",
        )
        print("PASS")

        print("\n测试 9：图表位置和尺寸保持克制")
        assert_true(chart.anchor._from.col == 3, "图表默认应从 D 列开始。")
        assert_true(chart.anchor._from.row == 5, "图表默认应从第 6 行开始。")
        # openpyxl 重新读取工作簿后，chart.width / chart.height 会回到
        # ChartBase 的默认值，不能用它们验证已保存文件的真实尺寸。
        # 实际保存尺寸位于 anchor.ext，单位为 EMU；1 cm = 360000 EMU。
        anchor_width_cm = chart.anchor.ext.cx / 360000
        anchor_height_cm = chart.anchor.ext.cy / 360000

        assert_true(
            12 <= anchor_width_cm <= 13,
            f"图表实际宽度不符合紧凑基准：{anchor_width_cm:.2f} cm。",
        )
        assert_true(
            6 <= anchor_height_cm <= 7,
            f"图表实际高度不符合紧凑基准：{anchor_height_cm:.2f} cm。",
        )
        print("PASS")

        print("\n测试 10：冻结窗格与 AutoFilter 保留")
        assert_true(str(ws.freeze_panes) == "A7", "冻结窗格被破坏。")
        assert_true(ws.auto_filter.ref == "A6:B9", "AutoFilter 被破坏。")
        print("PASS")

    print("\n" + "=" * 72)
    print("Professional Excel Visual Baseline 测试通过！")
    print("=" * 72)


if __name__ == "__main__":
    main()
