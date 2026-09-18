from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Optional

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


_INVALID_SHEET_CHARS = set(r'\/:*?[]')


def _normalize_output_path(output_path) -> Path:
    path = Path(output_path).expanduser()

    if path.suffix.lower() != ".xlsx":
        raise ValueError("专业 Excel 报告输出路径必须以 .xlsx 结尾。")

    try:
        path = path.resolve()
    except Exception:
        path = path.absolute()

    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def _clean_sheet_name(name: str, used_names: set) -> str:
    text = str(name or "").strip() or "Sheet"
    text = "".join("_" if char in _INVALID_SHEET_CHARS else char for char in text)
    text = text[:31] or "Sheet"

    candidate = text
    index = 2

    while candidate in used_names:
        suffix = f"_{index}"
        candidate = f"{text[:31 - len(suffix)]}{suffix}"
        index += 1

    used_names.add(candidate)
    return candidate


def _validate_dataframe(df: Any, label: str) -> pd.DataFrame:
    if not isinstance(df, pd.DataFrame):
        raise TypeError(f"{label} 必须是 pandas DataFrame。")

    return df.copy()


def _normalize_sheets(
    dataframe: Optional[pd.DataFrame],
    sheets: Optional[Dict[str, pd.DataFrame]],
    default_sheet_name: str,
) -> Dict[str, pd.DataFrame]:
    if dataframe is None and not sheets:
        raise ValueError("至少需要提供 dataframe 或 sheets。")

    result: Dict[str, pd.DataFrame] = {}

    if dataframe is not None:
        result[str(default_sheet_name or "数据明细")] = _validate_dataframe(
            dataframe,
            "dataframe",
        )

    if sheets is not None:
        if not isinstance(sheets, dict):
            raise TypeError("sheets 必须是 {Sheet名称: DataFrame} 字典。")

        for sheet_name, df in sheets.items():
            result[str(sheet_name)] = _validate_dataframe(
                df,
                f"Sheet {sheet_name}",
            )

    return result


def _normalize_kpis(kpis: Optional[List[Dict[str, Any]]]) -> List[Dict[str, Any]]:
    if not kpis:
        return []

    if not isinstance(kpis, list):
        raise TypeError("kpis 必须是列表。")

    normalized = []

    for index, item in enumerate(kpis, start=1):
        if not isinstance(item, dict):
            raise TypeError(f"第 {index} 个 KPI 必须是字典。")

        label = str(item.get("label") or "").strip()

        if not label:
            raise ValueError(f"第 {index} 个 KPI 缺少 label。")

        normalized.append(
            {
                "label": label,
                "value": item.get("value"),
                "number_format": item.get("number_format"),
            }
        )

    return normalized


def _normalize_charts(charts: Optional[List[Dict[str, Any]]]) -> List[Dict[str, Any]]:
    if not charts:
        return []

    if not isinstance(charts, list):
        raise TypeError("charts 必须是列表。")

    normalized = []

    for index, item in enumerate(charts, start=1):
        if not isinstance(item, dict):
            raise TypeError(f"第 {index} 个 chart 必须是字典。")

        chart_type = str(item.get("type") or "bar").strip().lower()

        if chart_type not in {"bar", "line", "pie"}:
            raise ValueError(
                f"第 {index} 个 chart type 暂不支持：{chart_type}。"
                "只支持 bar / line / pie。"
            )

        sheet_name = str(item.get("sheet_name") or "").strip()
        category_column = str(item.get("category_column") or "").strip()
        value_column = str(item.get("value_column") or "").strip()

        if not sheet_name or not category_column or not value_column:
            raise ValueError(
                f"第 {index} 个 chart 必须提供 "
                "sheet_name、category_column、value_column。"
            )

        normalized.append(
            {
                "type": chart_type,
                "sheet_name": sheet_name,
                "category_column": category_column,
                "value_column": value_column,
                "title": str(item.get("title") or "").strip(),
                "anchor": str(item.get("anchor") or "D6").strip() or "D6",
            }
        )

    return normalized


def _is_percentage_column(column_name: str) -> bool:
    text = str(column_name).lower()
    return any(
        token in text
        for token in [
            "百分比",
            "占比",
            "比例",
            "率",
            "percent",
            "ratio",
            "rate",
        ]
    )


def _is_currency_column(column_name: str) -> bool:
    text = str(column_name).lower()
    return any(
        token in text
        for token in [
            "销售额",
            "金额",
            "收入",
            "成本",
            "利润",
            "价格",
            "单价",
            "revenue",
            "amount",
            "sales",
            "cost",
            "profit",
            "price",
        ]
    )


def _is_date_column(column_name: str) -> bool:
    text = str(column_name).lower()
    return any(
        token in text
        for token in [
            "日期",
            "时间",
            "date",
            "time",
        ]
    )


def _apply_number_format(cell, column_name: str):
    value = cell.value

    if value is None:
        return

    if _is_percentage_column(column_name) and isinstance(value, (int, float)):
        cell.number_format = "0.0%"
        return

    if _is_currency_column(column_name) and isinstance(value, (int, float)):
        cell.number_format = '#,##0.00'
        return

    if _is_date_column(column_name):
        if hasattr(value, "year") and hasattr(value, "month"):
            cell.number_format = "yyyy-mm-dd"
            return

    if isinstance(value, float):
        cell.number_format = '#,##0.00'
    elif isinstance(value, int):
        cell.number_format = '#,##0'


def _display_length(value: Any) -> int:
    if value is None:
        return 0

    text = str(value)
    length = 0

    for char in text:
        length += 2 if ord(char) > 127 else 1

    return length


def _style_data_sheet(
    worksheet,
    header_row: int,
    data_end_row: int,
    data_column_count: int,
):
    """
    企业财务报表风格的数据区。

    目标：
    - 单元格边界清晰，不出现“透明/消失”的视觉效果；
    - 表头深色、正文浅色，层级明确；
    - 文本自动换行，避免较长字段名被遮挡；
    - 保留冻结窗格和 AutoFilter。
    """
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )
    header_font = Font(
        bold=True,
        color="FFFFFF",
        size=11,
    )
    body_fill = PatternFill(
        fill_type="solid",
        fgColor="FFFFFF",
    )
    alternate_fill = PatternFill(
        fill_type="solid",
        fgColor="F4F8FC",
    )
    border_side = Side(
        style="thin",
        color="B8C6D1",
    )
    border = Border(
        left=border_side,
        right=border_side,
        top=border_side,
        bottom=border_side,
    )

    for column_index in range(1, data_column_count + 1):
        cell = worksheet.cell(
            row=header_row,
            column=column_index,
        )
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = border

    for row_index in range(header_row + 1, data_end_row + 1):
        row_fill = (
            alternate_fill
            if (row_index - header_row) % 2 == 0
            else body_fill
        )

        for column_index in range(1, data_column_count + 1):
            cell = worksheet.cell(
                row=row_index,
                column=column_index,
            )
            cell.fill = row_fill
            cell.border = border
            cell.alignment = Alignment(
                horizontal=(
                    "left"
                    if isinstance(cell.value, str)
                    else "right"
                ),
                vertical="center",
                wrap_text=True,
            )

    worksheet.freeze_panes = f"A{header_row + 1}"

    if data_end_row >= header_row:
        worksheet.auto_filter.ref = (
            f"A{header_row}:"
            f"{get_column_letter(data_column_count)}{data_end_row}"
        )

    worksheet.row_dimensions[header_row].height = 26

    for row_index in range(header_row + 1, data_end_row + 1):
        worksheet.row_dimensions[row_index].height = 22


def _write_dataframe(
    worksheet,
    df: pd.DataFrame,
    start_row: int,
) -> Dict[str, Any]:
    header_row = start_row

    for column_index, column_name in enumerate(df.columns, start=1):
        worksheet.cell(
            row=header_row,
            column=column_index,
            value=str(column_name),
        )

    for row_offset, row_values in enumerate(
        df.itertuples(index=False, name=None),
        start=1,
    ):
        excel_row = header_row + row_offset

        for column_index, value in enumerate(row_values, start=1):
            if pd.isna(value):
                value = None

            cell = worksheet.cell(
                row=excel_row,
                column=column_index,
                value=value,
            )
            _apply_number_format(
                cell,
                str(df.columns[column_index - 1]),
            )

    data_end_row = header_row + len(df)
    _style_data_sheet(
        worksheet,
        header_row=header_row,
        data_end_row=max(data_end_row, header_row),
        data_column_count=len(df.columns),
    )

    for column_index, column_name in enumerate(df.columns, start=1):
        max_length = _display_length(column_name)

        for row_index in range(header_row + 1, data_end_row + 1):
            max_length = max(
                max_length,
                _display_length(
                    worksheet.cell(
                        row=row_index,
                        column=column_index,
                    ).value
                ),
            )

        # 中文表头与业务名称需要比纯字符长度预留更多空间。
        # 这里主动增加缓冲，并提高最大列宽，避免名称被截断/遮挡。
        worksheet.column_dimensions[
            get_column_letter(column_index)
        ].width = min(
            max(max_length + 5, 14),
            42,
        )

    return {
        "header_row": header_row,
        "data_start_row": header_row + 1,
        "data_end_row": data_end_row,
        "column_count": len(df.columns),
        "row_count": len(df),
    }


def _add_chart(
    worksheet,
    chart_config: Dict[str, Any],
    header_row: int,
    data_end_row: int,
):
    """
    添加克制的企业财务图表。

    v4.5 视觉基准：
    - 单系列柱状图统一企业蓝，不让每个城市自动变成不同颜色；
    - 不显示冗余图例；
    - 不显示横/纵坐标轴标题；
    - 关闭纵轴主网格线；
    - 柱顶直接显示数值；
    - 去掉明显图表外框。
    """
    headers = {
        str(cell.value).strip(): cell.column
        for cell in worksheet[header_row]
        if cell.value is not None
    }

    category_column = chart_config["category_column"]
    value_column = chart_config["value_column"]

    if category_column not in headers:
        raise ValueError(
            f"Sheet {worksheet.title} 中不存在图表分类字段：{category_column}"
        )

    if value_column not in headers:
        raise ValueError(
            f"Sheet {worksheet.title} 中不存在图表数值字段：{value_column}"
        )

    if data_end_row < header_row + 1:
        raise ValueError(
            f"Sheet {worksheet.title} 没有数据行，无法生成图表。"
        )

    chart_type = chart_config["type"]

    if chart_type == "bar":
        chart = BarChart()
        chart.type = "col"
        chart.varyColors = False
        chart.legend = None
        chart.gapWidth = 85
        chart.dLbls = DataLabelList()
        chart.dLbls.showVal = True
        chart.dLbls.showSerName = False
        chart.dLbls.showCatName = False
        chart.dLbls.showLegendKey = False
        chart.dLbls.showPercent = False
        chart.dLbls.position = "outEnd"
        chart.dLbls.numFmt = "0"

    elif chart_type == "line":
        chart = LineChart()
        chart.varyColors = False
        chart.legend = None
        chart.dLbls = DataLabelList()
        chart.dLbls.showVal = True

    else:
        chart = PieChart()
        chart.legend = None
        chart.dLbls = DataLabelList()
        chart.dLbls.showPercent = True
        chart.dLbls.showLeaderLines = True

    values = Reference(
        worksheet,
        min_col=headers[value_column],
        min_row=header_row,
        max_row=data_end_row,
    )
    categories = Reference(
        worksheet,
        min_col=headers[category_column],
        min_row=header_row + 1,
        max_row=data_end_row,
    )

    chart.add_data(values, titles_from_data=True)
    chart.set_categories(categories)
    chart.title = (
        chart_config["title"]
        or f"{category_column} - {value_column}"
    )

    if chart_type in {"bar", "line"}:
        if chart.series:
            try:
                series = chart.series[0]
                series.graphicalProperties.solidFill = "4472C4"
                series.graphicalProperties.line.solidFill = "4472C4"
            except Exception:
                pass

        chart.x_axis.title = None
        chart.y_axis.title = None
        chart.y_axis.majorGridlines = None

        try:
            # 明确要求 Excel Desktop 在横轴底部显示分类标签。
            chart.x_axis.tickLblPos = "low"
            chart.x_axis.delete = False
            chart.x_axis.noMultiLvlLbl = True
        except Exception:
            pass


    try:
        chart.graphical_properties = GraphicalProperties(noFill=True)
    except Exception:
        pass

    chart.height = 6.8
    chart.width = 12.6

    worksheet.add_chart(
        chart,
        chart_config["anchor"],
    )



def create_professional_excel_report(
    output_path,
    dataframe=None,
    sheets=None,
    report_title="DataPilot 分析报告",
    subtitle=None,
    kpis=None,
    charts=None,
    default_sheet_name="数据明细",
) -> Dict[str, Any]:
    """
    创建新的专业 Excel 报告。

    设计边界：
    - 这是“新报告生成器”，不是已有 Excel 高保真编辑器；
    - dataframe / sheets 必须是已经分析完成的 DataFrame；
    - Python 负责稳定排版，LLM 只负责业务内容与图表意图；
    - 不读取或修改源 Excel，因此不会覆盖业务源文件。
    """
    target = _normalize_output_path(output_path)
    normalized_sheets = _normalize_sheets(
        dataframe=dataframe,
        sheets=sheets,
        default_sheet_name=default_sheet_name,
    )
    normalized_kpis = _normalize_kpis(kpis)
    normalized_charts = _normalize_charts(charts)

    workbook = Workbook()
    workbook.remove(workbook.active)

    used_names = set()
    sheet_name_map: Dict[str, str] = {}
    sheet_layouts: Dict[str, Dict[str, Any]] = {}

    for requested_name, df in normalized_sheets.items():
        actual_name = _clean_sheet_name(
            requested_name,
            used_names,
        )
        sheet_name_map[requested_name] = actual_name

        worksheet = workbook.create_sheet(actual_name)
        worksheet.sheet_view.showGridLines = False

        # 第一张报告页需要同时容纳 KPI 与图表，因此标题横跨至少 8 列；
        # 这样不会出现标题区太窄、图表挤到页面外侧的问题。
        is_primary_sheet = requested_name == next(iter(normalized_sheets))
        visual_columns = 8 if is_primary_sheet else max(len(df.columns), 4)
        max_columns = max(len(df.columns), visual_columns)
        last_title_column = get_column_letter(max_columns)

        worksheet.merge_cells(
            f"A1:{last_title_column}1"
        )
        title_cell = worksheet["A1"]
        title_cell.value = str(report_title or "DataPilot 分析报告")
        title_cell.font = Font(
            bold=True,
            size=20,
            color="17365D",
        )
        title_cell.fill = PatternFill(
            fill_type="solid",
            fgColor="FFFFFF",
        )
        title_cell.alignment = Alignment(
            horizontal="left",
            vertical="center",
            wrap_text=True,
        )
        worksheet.row_dimensions[1].height = 34

        current_row = 2

        if subtitle:
            worksheet.merge_cells(
                f"A2:{last_title_column}2"
            )
            subtitle_cell = worksheet["A2"]
            subtitle_cell.value = str(subtitle)
            subtitle_cell.font = Font(
                size=10,
                color="44546A",
            )
            subtitle_cell.fill = PatternFill(
                fill_type="solid",
                fgColor="FFFFFF",
            )
            subtitle_cell.alignment = Alignment(
                horizontal="left",
                vertical="center",
                wrap_text=True,
            )
            worksheet.row_dimensions[2].height = 30
            current_row = 4
        else:
            current_row = 3

        if normalized_kpis and is_primary_sheet:
            kpi_start_row = current_row
            kpi_fill = PatternFill(
                fill_type="solid",
                fgColor="EAF2F8",
            )
            kpi_border_side = Side(
                style="thin",
                color="A9C4DA",
            )
            kpi_border = Border(
                left=kpi_border_side,
                right=kpi_border_side,
                top=kpi_border_side,
                bottom=kpi_border_side,
            )

            for index, kpi in enumerate(normalized_kpis, start=1):
                label_column = (index - 1) * 2 + 1
                value_column = label_column + 1

                label_cell = worksheet.cell(
                    row=kpi_start_row,
                    column=label_column,
                    value=kpi["label"],
                )
                value_cell = worksheet.cell(
                    row=kpi_start_row,
                    column=value_column,
                    value=kpi["value"],
                )

                label_cell.font = Font(
                    bold=True,
                    size=9,
                    color="44546A",
                )
                value_cell.font = Font(
                    bold=True,
                    size=12,
                    color="17365D",
                )

                for cell in (label_cell, value_cell):
                    cell.fill = kpi_fill
                    cell.border = kpi_border
                    cell.alignment = Alignment(
                        horizontal="center",
                        vertical="center",
                        wrap_text=True,
                    )

                if kpi["number_format"]:
                    value_cell.number_format = str(kpi["number_format"])

                # KPI 区使用固定的业务友好宽度，防止中文名称溢出。
                worksheet.column_dimensions[
                    get_column_letter(label_column)
                ].width = max(
                    worksheet.column_dimensions[
                        get_column_letter(label_column)
                    ].width or 0,
                    18,
                )
                worksheet.column_dimensions[
                    get_column_letter(value_column)
                ].width = max(
                    worksheet.column_dimensions[
                        get_column_letter(value_column)
                    ].width or 0,
                    14,
                )

            worksheet.row_dimensions[kpi_start_row].height = 34
            current_row += 2

        # 先保存 KPI 所需宽度；数据表自动列宽后再取较大值，
        # 避免数据表的窄列把 KPI 名称重新压窄。
        preserved_widths = {
            column_letter: dimension.width
            for column_letter, dimension in worksheet.column_dimensions.items()
            if dimension.width
        }

        layout = _write_dataframe(
            worksheet,
            df,
            start_row=current_row,
        )

        for column_letter, width in preserved_widths.items():
            current_width = (
                worksheet.column_dimensions[column_letter].width
                or 0
            )
            worksheet.column_dimensions[column_letter].width = max(
                current_width,
                width,
            )

        # 图表区域预留稳定宽度，使默认 D6 锚点后的图表不会挤压。
        if is_primary_sheet:
            for column_letter in ("D", "E", "F", "G", "H"):
                current_width = (
                    worksheet.column_dimensions[column_letter].width
                    or 0
                )
                worksheet.column_dimensions[column_letter].width = max(
                    current_width,
                    13,
                )

        sheet_layouts[actual_name] = layout

    for chart_config in normalized_charts:
        requested_sheet = chart_config["sheet_name"]

        if requested_sheet in sheet_name_map:
            actual_sheet = sheet_name_map[requested_sheet]
        elif requested_sheet in workbook.sheetnames:
            actual_sheet = requested_sheet
        else:
            raise ValueError(
                f"图表指定的 Sheet 不存在：{requested_sheet}。"
                f"可用 Sheet：{', '.join(workbook.sheetnames)}"
            )

        worksheet = workbook[actual_sheet]
        layout = sheet_layouts[actual_sheet]

        _add_chart(
            worksheet,
            chart_config,
            header_row=layout["header_row"],
            data_end_row=layout["data_end_row"],
        )

    workbook.save(target)

    return {
        "success": True,
        "output_path": str(target),
        "sheet_names": list(workbook.sheetnames),
        "sheet_name_map": sheet_name_map,
        "sheet_count": len(workbook.sheetnames),
        "kpi_count": len(normalized_kpis),
        "chart_count": len(normalized_charts),
        "report_title": str(report_title or "DataPilot 分析报告"),
    }


def inspect_professional_excel_report(
    file_path,
    max_preview_rows=10,
) -> Dict[str, Any]:
    """
    重新打开专业 Excel 报告并返回可供 Agent / Verification 使用的结构证据。
    """
    path = Path(file_path).expanduser()

    try:
        path = path.resolve()
    except Exception:
        path = path.absolute()

    if not path.exists():
        raise FileNotFoundError(f"Excel 报告不存在：{path}")

    if path.suffix.lower() != ".xlsx":
        raise ValueError("专业 Excel 报告检查器只支持 .xlsx。")

    workbook = load_workbook(
        path,
        data_only=False,
    )

    sheets = []

    for worksheet in workbook.worksheets:
        header_row = None
        headers = []

        search_limit = min(worksheet.max_row, 20)

        for row_index in range(1, search_limit + 1):
            row_cells = list(worksheet[row_index])
            non_empty_cells = [
                cell
                for cell in row_cells
                if cell.value is not None
            ]

            if len(non_empty_cells) < 2:
                continue

            # KPI 行通常是“标签/值/标签/值”，不能被误判成数据表头。
            # 真正的数据表头应当紧邻至少一行具有相同有效列宽的数据。
            candidate_last_column = max(
                cell.column
                for cell in non_empty_cells
            )

            next_row_index = row_index + 1

            if next_row_index > worksheet.max_row:
                continue

            next_row_values = [
                worksheet.cell(
                    row=next_row_index,
                    column=column_index,
                ).value
                for column_index in range(
                    1,
                    candidate_last_column + 1,
                )
            ]

            if not any(
                value is not None
                for value in next_row_values
            ):
                continue

            candidate_headers = [
                str(
                    worksheet.cell(
                        row=row_index,
                        column=column_index,
                    ).value
                ).strip()
                for column_index in range(
                    1,
                    candidate_last_column + 1,
                )
            ]

            if any(
                not header
                or header == "None"
                for header in candidate_headers
            ):
                continue

            header_row = row_index
            headers = candidate_headers
            break

        preview = []

        if header_row is not None:
            end_row = min(
                worksheet.max_row,
                header_row + int(max_preview_rows),
            )

            for row_index in range(header_row + 1, end_row + 1):
                record = {}

                for column_index, header in enumerate(headers, start=1):
                    record[header] = worksheet.cell(
                        row=row_index,
                        column=column_index,
                    ).value

                preview.append(record)

        sheets.append(
            {
                "sheet_name": worksheet.title,
                "max_row": worksheet.max_row,
                "max_column": worksheet.max_column,
                "title": worksheet["A1"].value,
                "header_row": header_row,
                "headers": headers,
                "freeze_panes": (
                    str(worksheet.freeze_panes)
                    if worksheet.freeze_panes
                    else None
                ),
                "auto_filter_ref": worksheet.auto_filter.ref,
                "chart_count": len(worksheet._charts),
                "merged_ranges": [
                    str(item)
                    for item in worksheet.merged_cells.ranges
                ],
                "preview_records": preview,
            }
        )

    return {
        "success": True,
        "file_path": str(path),
        "sheet_names": workbook.sheetnames,
        "sheet_count": len(workbook.sheetnames),
        "chart_count": sum(
            len(worksheet._charts)
            for worksheet in workbook.worksheets
        ),
        "sheets": sheets,
    }


if __name__ == "__main__":
    print("excel_report_tools.py 已加载成功。")
