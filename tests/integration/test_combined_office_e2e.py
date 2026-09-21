from __future__ import annotations

import json
import re
import tempfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import pandas as pd
from docx import Document
from openpyxl import load_workbook

import core.agent as agent_module
from core.agent import DataPilotAgent
from core.task_planner import TaskPlan


class FakeTaskPlanner:
    """
    不调用外部 LLM 的确定性 TaskPlanner 替身。

    本测试验证真实：
    agent.py -> Workspace -> TaskPlan -> AgentLoop -> SkillSelector
    -> ToolExecutor -> ToolPreflight -> ToolRegistry
    -> Professional Excel + Professional Word
    -> Cross-deliverable Verification -> Completion Gate。
    """

    def __init__(self, *args, **kwargs):
        pass

    def create_plan(self, task, context=None):
        return TaskPlan(
            task_goal=(
                "读取销售 Excel，按城市统计销售额合计，"
                "基于同一份分析结果同时生成专业 Excel 报告和正式 Word 汇报。"
            ),
            evidence_requirements=[],
            source_requirements=[
                "所有业务数字必须来自真实源 Excel。",
            ],
            deliverable_requirements=[
                "生成最终专业 Excel 交付物。",
                "生成最终正式 Word 交付物。",
            ],
            execution_requirements=[
                "读取真实销售 Excel。",
                "按城市统计销售额合计并确定销售冠军。",
                "复用同一份分析结果生成专业 Excel 和正式 Word。",
            ],
            verification_requirements=[
                "生成后重新读取并检查最终 Excel。",
                "生成后重新读取并检查最终 Word。",
                (
                    "核对最终 Excel 和 Word 中的城市销售汇总、销售总额与"
                    "销售冠军彼此一致，并与源数据一致。"
                ),
            ],
            safety_requirements=[
                "不得覆盖源 Excel。",
                "最终 Excel 和 Word 必须位于 deliverables_dir。",
            ],
            assumptions=[],
        )


class FakeCompletions:
    """
    只替代 Agent 的“下一步决策”层。

    数据读取、分组、排序、两个 Office 文件生成、两个最终文件回读、
    Verification Engine 和 Completion Gate 全部使用 DataPilot 真实实现。
    """

    def __init__(self, source_path: Path):
        self.source_path = str(source_path.resolve())
        self.calls = []
        self.decision_index = 0
        self.excel_path = None
        self.word_path = None

    @staticmethod
    def _extract_deliverables_dir(messages):
        user_content = next(
            (
                item.get("content", "")
                for item in messages
                if item.get("role") == "user"
            ),
            "",
        )

        match = re.search(
            r'"deliverables_dir"\s*:\s*"((?:\\.|[^"])*)"',
            user_content,
        )
        if not match:
            raise AssertionError(
                "Fake Client 无法从真实 Agent 状态读取 deliverables_dir。"
            )

        return json.loads('"' + match.group(1) + '"')

    def create(self, **kwargs):
        self.calls.append(kwargs)
        messages = kwargs.get("messages", [])

        if self.excel_path is None:
            deliverables_dir = self._extract_deliverables_dir(messages)
            self.excel_path = str(
                (
                    Path(deliverables_dir)
                    / "城市销售分析报告.xlsx"
                ).resolve()
            )
            self.word_path = str(
                (
                    Path(deliverables_dir)
                    / "城市销售分析汇报.docx"
                ).resolve()
            )

        decisions = [
            {
                "action_type": "tool",
                "tool": "read_office_data",
                "arguments": {
                    "file_path": self.source_path,
                },
                "purpose": "读取真实销售源数据。",
            },
            {
                "action_type": "tool",
                "tool": "group_statistics",
                "arguments": {
                    "df": {"$ref": "step_1.output"},
                    "group_by": "城市",
                    "target_column": "销售额",
                    "operation": "sum",
                },
                "purpose": "按城市统计销售额合计，形成两个交付物共同复用的分析结果。",
            },
            {
                "action_type": "tool",
                "tool": "sort_data",
                "arguments": {
                    "df": {"$ref": "step_2.output"},
                    "column": "销售额_合计",
                    "ascending": False,
                },
                "purpose": "按销售额合计降序排列并确定销售冠军。",
            },
            {
                "action_type": "tool",
                "tool": "create_professional_excel_report",
                "arguments": {
                    "output_path": self.excel_path,
                    "dataframe": {"$ref": "step_3.output"},
                    "default_sheet_name": "城市销售汇总",
                    "report_title": "城市销售分析报告",
                    "subtitle": "Excel 与 Word 联合交付 · 同源分析结果",
                    "kpis": [
                        {"label": "销售冠军", "value": "澳门"},
                        {
                            "label": "最高销售额",
                            "value": 186,
                            "number_format": "#,##0",
                        },
                        {
                            "label": "销售总额",
                            "value": 468,
                            "number_format": "#,##0",
                        },
                    ],
                    "charts": [
                        {
                            "type": "bar",
                            "sheet_name": "城市销售汇总",
                            "category_column": "城市",
                            "value_column": "销售额_合计",
                            "title": "城市销售额对比",
                            "anchor": "E7",
                        }
                    ],
                },
                "purpose": "使用统一分析结果生成最终专业 Excel。",
            },
            {
                "action_type": "tool",
                "tool": "inspect_professional_excel_report",
                "arguments": {
                    "file_path": self.excel_path,
                },
                "purpose": "重新打开最终 Excel，取得跨交付物核验所需真实证据。",
            },
            {
                "action_type": "tool",
                "tool": "create_professional_word_report",
                "arguments": {
                    "output_path": self.word_path,
                    "report_title": "城市销售分析汇报",
                    "subtitle": "Excel 与 Word 联合交付 · 同源分析结果",
                    "metadata": {
                        "报告对象": "管理层",
                        "统计口径": "按城市汇总销售额",
                    },
                    "executive_summary": (
                        "本次城市销售汇总显示，澳门销售额最高，为 186；"
                        "横琴为 154，珠海为 128。三个城市销售总额为 468。"
                    ),
                    "kpis": [
                        {"label": "销售冠军", "value": "澳门"},
                        {"label": "最高销售额", "value": 186},
                        {"label": "销售总额", "value": 468},
                    ],
                    "sections": [
                        {
                            "title": "关键结论",
                            "type": "bullets",
                            "items": [
                                "澳门销售额为 186，排名第一。",
                                "横琴销售额为 154。",
                                "珠海销售额为 128。",
                            ],
                        },
                        {
                            "title": "城市销售汇总",
                            "type": "table",
                            "columns": ["城市", "销售额合计"],
                            "rows": [
                                {"城市": "澳门", "销售额合计": 186},
                                {"城市": "横琴", "销售额合计": 154},
                                {"城市": "珠海", "销售额合计": 128},
                            ],
                        },
                    ],
                    "source_note": (
                        "数据来源：销售数据.xlsx；"
                        "Excel 与 Word 使用同一份已验证城市销售汇总结果。"
                    ),
                },
                "purpose": "复用与 Excel 相同的业务结果生成最终正式 Word。",
            },
            {
                "action_type": "tool",
                "tool": "inspect_professional_word_report",
                "arguments": {
                    "file_path": self.word_path,
                },
                "purpose": "重新打开最终 Word，取得跨交付物核验所需真实证据。",
            },
            {
                "action_type": "finish",
                "final_answer": (
                    "已基于同一份销售分析结果生成专业 Excel 和正式 Word，"
                    "两个最终文件均已重新读取并完成一致性核验。"
                ),
            },
        ]

        decision = (
            decisions[self.decision_index]
            if self.decision_index < len(decisions)
            else decisions[-1]
        )
        self.decision_index += 1

        return SimpleNamespace(
            choices=[
                SimpleNamespace(
                    message=SimpleNamespace(
                        content=json.dumps(
                            decision,
                            ensure_ascii=False,
                        )
                    )
                )
            ]
        )


class FakeClient:
    def __init__(self, source_path: Path):
        self.chat = SimpleNamespace(
            completions=FakeCompletions(source_path)
        )


def assert_true(value, message):
    if not value:
        raise AssertionError(message)


def find_word_table(document, *needles):
    for table in document.tables:
        text = "\n".join(
            cell.text
            for row in table.rows
            for cell in row.cells
        )
        if all(str(item) in text for item in needles):
            return table
    return None


def main():
    print("=" * 72)
    print("DataPilot v5.0 Excel + Word Combined Agent E2E 确定性测试")
    print("=" * 72)

    with tempfile.TemporaryDirectory(
        prefix="datapilot_combined_office_e2e_"
    ) as temp_dir:
        root = Path(temp_dir)
        source_path = root / "销售数据.xlsx"
        output_root = root / "outputs"

        source_df = pd.DataFrame(
            {
                "城市": ["珠海", "澳门", "横琴", "澳门"],
                "销售额": [128, 100, 154, 86],
            }
        )
        source_df.to_excel(source_path, index=False)

        fake_client = FakeClient(source_path)
        datapilot = DataPilotAgent.__new__(DataPilotAgent)
        datapilot.api_key = "test-key"
        datapilot.base_url = "https://example.invalid"
        datapilot.model = "test-model"
        datapilot.progress_callback = None
        datapilot.client = fake_client

        user_task = (
            "分析这份销售数据，按城市汇总销售额，同时生成一份专业 Excel "
            "分析报告和一份可以直接发给领导的正式 Word 汇报。"
            "两个文件中的销售总额、城市汇总和销售冠军必须一致。"
            "不要覆盖原文件。生成后重新读取 Excel 和 Word，"
            "完成跨交付物核验，确认无误后再完成任务。"
        )

        print("\n测试 1：真实 agent.py 建立多交付物 TaskPlan")
        with patch.object(agent_module, "TaskPlanner", FakeTaskPlanner):
            result = datapilot.execute_v31_agent_task(
                user_task=user_task,
                input_paths=[str(source_path)],
                output_dir=str(output_root),
                max_iterations=10,
            )

        task_plan = result.get("task_plan", {})
        assert_true(isinstance(task_plan, dict), "没有返回 TaskPlan。")
        assert_true(
            len(task_plan.get("deliverable_requirements", [])) == 2,
            f"TaskPlan 没有两个交付物要求：{task_plan}",
        )
        print("PASS")

        print("\n测试 2：真实 Skill Guidance 同时包含 Excel + Word + 数据分析")
        calls = fake_client.chat.completions.calls
        assert_true(calls, "AgentLoop 没有调用决策模型接口。")
        first_messages = calls[0].get("messages", [])
        first_user_prompt = next(
            (
                item.get("content", "")
                for item in first_messages
                if item.get("role") == "user"
            ),
            "",
        )
        assert_true(
            '"excel_data_analysis"' in first_user_prompt,
            "没有选中 excel_data_analysis。",
        )
        assert_true(
            '"excel_report_delivery"' in first_user_prompt,
            "没有选中 excel_report_delivery。",
        )
        assert_true(
            '"professional_word_delivery"' in first_user_prompt,
            "没有选中 professional_word_delivery。",
        )
        print("PASS")

        history = result.get("tool_results", [])
        tool_names = [
            item.get("tool_name")
            for item in history
            if isinstance(item, dict)
        ]

        print("\n测试 3：源 Excel 只读取一次并形成统一分析链")
        assert_true(
            tool_names.count("read_office_data") == 1,
            f"源数据读取次数异常：{tool_names}",
        )
        assert_true(
            "group_statistics" in tool_names
            and "sort_data" in tool_names,
            f"缺少统一分析步骤：{tool_names}",
        )
        print("PASS")

        print("\n测试 4：真实生成 Professional Excel")
        assert_true(
            "create_professional_excel_report" in tool_names,
            f"没有调用 Professional Excel Tool：{tool_names}",
        )
        print("PASS")

        print("\n测试 5：最终 Excel 在写入后被重新检查")
        excel_create = tool_names.index(
            "create_professional_excel_report"
        )
        excel_inspect = tool_names.index(
            "inspect_professional_excel_report"
        )
        assert_true(
            excel_inspect > excel_create,
            "Excel inspect 没有发生在最后写入之后。",
        )
        print("PASS")

        print("\n测试 6：真实生成 Professional Word")
        assert_true(
            "create_professional_word_report" in tool_names,
            f"没有调用 Professional Word Tool：{tool_names}",
        )
        print("PASS")

        print("\n测试 7：最终 Word 在写入后被重新检查")
        word_create = tool_names.index(
            "create_professional_word_report"
        )
        word_inspect = tool_names.index(
            "inspect_professional_word_report"
        )
        assert_true(
            word_inspect > word_create,
            "Word inspect 没有发生在最后写入之后。",
        )
        print("PASS")

        excel_path = Path(fake_client.chat.completions.excel_path)
        word_path = Path(fake_client.chat.completions.word_path)

        print("\n测试 8：两个最终交付物都位于同一 Workspace deliverables_dir")
        assert_true(excel_path.exists(), "最终 Excel 不存在。")
        assert_true(word_path.exists(), "最终 Word 不存在。")
        assert_true(
            excel_path.parent == word_path.parent,
            "Excel 和 Word 不在同一 deliverables_dir。",
        )
        print("PASS")

        print("\n测试 9：Excel 城市汇总业务数据正确")
        workbook = load_workbook(excel_path, data_only=False)
        worksheet = workbook["城市销售汇总"]
        header_row = None
        for row_index in range(1, worksheet.max_row + 1):
            values = [
                worksheet.cell(row_index, col).value
                for col in range(1, min(worksheet.max_column, 4) + 1)
            ]
            if "城市" in values and "销售额_合计" in values:
                header_row = row_index
                break
        assert_true(header_row is not None, "Excel 没找到城市汇总表头。")
        excel_data = {}
        for row_index in range(header_row + 1, worksheet.max_row + 1):
            city = worksheet.cell(row_index, 1).value
            amount = worksheet.cell(row_index, 2).value
            if city in {"澳门", "横琴", "珠海"}:
                excel_data[str(city)] = amount
        expected = {"澳门": 186, "横琴": 154, "珠海": 128}
        assert_true(
            excel_data == expected,
            f"Excel 数据错误：{excel_data}",
        )
        print("PASS")

        print("\n测试 10：最终 Excel 回读包含真实 KPI 证据")
        excel_inspection = next(
            item.get("output", {})
            for item in history
            if isinstance(item, dict)
            and item.get("tool_name")
            == "inspect_professional_excel_report"
        )
        excel_kpis = excel_inspection.get("kpi_map", {})
        assert_true(
            excel_kpis.get("销售冠军") == "澳门",
            f"Excel 回读缺少销售冠军 KPI：{excel_kpis}",
        )
        assert_true(
            excel_kpis.get("最高销售额") == 186,
            f"Excel 回读最高销售额 KPI 错误：{excel_kpis}",
        )
        assert_true(
            excel_kpis.get("销售总额") == 468,
            f"Excel 回读销售总额 KPI 错误：{excel_kpis}",
        )
        print("PASS")

        print("\n测试 11：Word 城市汇总业务数据正确")
        document = Document(word_path)
        business_table = find_word_table(
            document,
            "城市",
            "销售额合计",
            "澳门",
            "横琴",
            "珠海",
        )
        assert_true(
            business_table is not None,
            "Word 没有找到城市销售汇总表。",
        )
        word_text = "\n".join(
            cell.text
            for row in business_table.rows
            for cell in row.cells
        )
        for city, amount in expected.items():
            assert_true(
                city in word_text and str(amount) in word_text,
                f"Word 缺少 {city}={amount}。",
            )
        print("PASS")

        print("\n测试 12：Excel 与 Word 共同业务结果一致")
        for city, amount in expected.items():
            assert_true(
                excel_data.get(city) == amount
                and city in word_text
                and str(amount) in word_text,
                f"跨交付物城市数据不一致：{city}",
            )
        assert_true(
            "澳门" in word_text
            and excel_data["澳门"] == 186,
            "跨交付物冠军证据不一致。",
        )
        print("PASS")

        print("\n测试 13：Completion Gate 最终 PASS")
        verification_report = result.get("verification_report")
        assert_true(
            isinstance(verification_report, dict),
            "没有返回 VerificationReport。",
        )
        assert_true(
            verification_report.get("verified") is True,
            f"Completion Gate 未 PASS：{verification_report}",
        )
        assert_true(
            result.get("success") is True
            and result.get("stop_reason") == "completed",
            f"任务没有 completed：{result.get('stop_reason')}",
        )
        print("PASS")

        print("\n测试 14：Cross-deliverable requirement 不再 pending")
        pending = verification_report.get(
            "pending_requirements",
            [],
        )
        assert_true(
            not any(
                "Excel" in str(item)
                and "Word" in str(item)
                and "一致" in str(item)
                for item in pending
            ),
            f"跨交付物一致性仍 pending：{pending}",
        )
        print("PASS")

        print("\n测试 15：公开 VerificationReport 协议正确表达跨交付物验收结果")
        # VerificationReport.to_dict() 的公开协议并不包含 semantic_evidence。
        # cross_deliverable_chain / cross_deliverable_source_chain 是
        # VerificationEngine 内部 semantic resolver 的审计证据，不应由
        # E2E 测试虚构一个不存在的公开字段。
        #
        # 对公开 E2E 来说，正确的可观察合同是：
        # 1. Completion Gate verified=True；
        # 2. 该跨交付物 requirement 不再 pending；
        # 3. 没有 verification failure。
        assert_true(
            "semantic_evidence" not in verification_report,
            "公开 VerificationReport 协议意外暴露了未定义 semantic_evidence 字段。",
        )
        assert_true(
            verification_report.get("verified") is True,
            f"跨交付物验收未通过：{verification_report}",
        )
        assert_true(
            not verification_report.get("pending_requirements", []),
            (
                "跨交付物 requirement 仍未被 Verification Engine 消解："
                f"{verification_report.get('pending_requirements')}"
            ),
        )
        assert_true(
            not verification_report.get("failures", []),
            (
                "VerificationReport 仍存在 failure："
                f"{verification_report.get('failures')}"
            ),
        )
        print("PASS")

        print("\n测试 16：源 Excel 未被覆盖")
        source_after = pd.read_excel(source_path)
        pd.testing.assert_frame_equal(
            source_after,
            source_df,
            check_dtype=False,
        )
        print("PASS")

        print("\n测试 17：最终结果同时返回两个真实交付物")
        output_files = {
            str(Path(item).resolve())
            for item in result.get("output_files", [])
        }
        assert_true(
            str(excel_path.resolve()) in output_files,
            f"output_files 缺少 Excel：{output_files}",
        )
        assert_true(
            str(word_path.resolve()) in output_files,
            f"output_files 缺少 Word：{output_files}",
        )
        print("PASS")

        print("\n测试 18：v4.0 Planned Workspace 返回协议继续兼容")
        assert_true(
            result.get("plan", {}).get("task_type")
            == "v4_0_planned_workspace_agent_loop",
            "破坏了 v4.0 plan 元数据。",
        )
        assert_true(
            result.get("runtime_context", {}).get("task_plan")
            == result.get("task_plan"),
            "runtime_context TaskPlan 与最终返回不一致。",
        )
        print("PASS")

    print("\n" + "=" * 72)
    print("Excel + Word Combined Agent E2E：18/18 PASS")
    print("=" * 72)


if __name__ == "__main__":
    main()
