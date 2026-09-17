from pathlib import Path
from datetime import datetime

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_cell_text(cell, text, bold=False):
    """
    设置 Word 表格单元格文字。
    """
    cell.text = ""

    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(9)
    run.font.name = "Microsoft YaHei"

    # 设置中文字体
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_shading(cell, fill):
    """
    设置表格单元格背景颜色。
    """
    tc_pr = cell._tc.get_or_add_tcPr()

    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)

    shd.set(qn("w:fill"), fill)


def set_table_borders(table):
    """
    设置 Word 表格边框。
    """
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    borders = tbl_pr.first_child_found_in("w:tblBorders")

    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))

        if element is None:
            element = OxmlElement(tag)
            borders.append(element)

        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "B7B7B7")


def set_document_default_font(document):
    """
    设置 Word 文档默认字体。
    """
    styles = document.styles

    normal_style = styles["Normal"]
    normal_style.font.name = "Microsoft YaHei"
    normal_style.font.size = Pt(10)

    normal_style._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        "Microsoft YaHei"
    )


def add_heading(document, text, level=1):
    """
    添加标题并设置中文字体。
    """
    paragraph = document.add_heading(level=level)
    run = paragraph.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    return paragraph


def add_body_paragraph(document, text, bold=False):
    """
    添加普通正文段落。
    """
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)

    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(10)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    return paragraph


def add_key_value_table(document, rows):
    """
    添加两列表格。
    rows 格式：
    [
        ("项目", "内容"),
        ...
    ]
    """
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    header_cells = table.rows[0].cells
    set_cell_text(header_cells[0], "项目", bold=True)
    set_cell_text(header_cells[1], "内容", bold=True)

    set_cell_shading(header_cells[0], "D9EAF7")
    set_cell_shading(header_cells[1], "D9EAF7")

    for key, value in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], key)
        set_cell_text(cells[1], value)

    set_table_borders(table)
    document.add_paragraph()

    return table


def add_list_table(document, title, items):
    """
    添加列表表格。
    """
    add_heading(document, title, level=2)

    if not items:
        add_body_paragraph(document, "暂无记录。")
        return

    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    headers = table.rows[0].cells
    set_cell_text(headers[0], "序号", bold=True)
    set_cell_text(headers[1], "内容", bold=True)

    set_cell_shading(headers[0], "D9EAF7")
    set_cell_shading(headers[1], "D9EAF7")

    for index, item in enumerate(items, start=1):
        cells = table.add_row().cells
        set_cell_text(cells[0], index)
        set_cell_text(cells[1], item)

    set_table_borders(table)
    document.add_paragraph()


def add_quality_table(document, title, quality_result):
    """
    添加数据质量检查结果表格。
    """
    add_heading(document, title, level=2)

    if not isinstance(quality_result, dict) or not quality_result:
        add_body_paragraph(document, "暂无质量检查结果。")
        return

    rows = []

    for key, value in quality_result.items():
        if isinstance(value, (dict, list, tuple)):
            value = str(value)

        rows.append((key, value))

    add_key_value_table(document, rows)


def add_statistics_table(document, statistics):
    """
    添加统计分析结果。
    支持嵌套字典。
    """
    add_heading(document, "统计分析结果", level=2)

    if not isinstance(statistics, dict) or not statistics:
        add_body_paragraph(document, "暂无统计分析结果。")
        return

    table = document.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    headers = table.rows[0].cells
    set_cell_text(headers[0], "字段", bold=True)
    set_cell_text(headers[1], "统计项目", bold=True)
    set_cell_text(headers[2], "结果", bold=True)

    for cell in headers:
        set_cell_shading(cell, "D9EAF7")

    for field_name, field_result in statistics.items():
        if isinstance(field_result, dict):
            for stat_name, stat_value in field_result.items():
                cells = table.add_row().cells
                set_cell_text(cells[0], field_name)
                set_cell_text(cells[1], stat_name)
                set_cell_text(cells[2], stat_value)
        else:
            cells = table.add_row().cells
            set_cell_text(cells[0], field_name)
            set_cell_text(cells[1], "结果")
            set_cell_text(cells[2], field_result)

    set_table_borders(table)
    document.add_paragraph()


def add_file_list_table(document, file_paths):
    """
    添加处理文件列表。
    """
    add_heading(document, "本次处理的文件", level=2)

    if not file_paths:
        add_body_paragraph(document, "没有记录到输入文件。")
        return

    table = document.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    headers = table.rows[0].cells
    set_cell_text(headers[0], "序号", bold=True)
    set_cell_text(headers[1], "文件名", bold=True)
    set_cell_text(headers[2], "完整路径", bold=True)

    for cell in headers:
        set_cell_shading(cell, "D9EAF7")

    for index, file_path in enumerate(file_paths, start=1):
        path = Path(str(file_path))

        cells = table.add_row().cells
        set_cell_text(cells[0], index)
        set_cell_text(cells[1], path.name)
        set_cell_text(cells[2], str(path))

    set_table_borders(table)
    document.add_paragraph()


def add_output_files_table(document, output_files):
    """
    添加输出文件清单。
    """
    add_heading(document, "输出文件清单", level=2)

    if not output_files:
        add_body_paragraph(document, "暂无输出文件记录。")
        return

    table = document.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    headers = table.rows[0].cells
    set_cell_text(headers[0], "序号", bold=True)
    set_cell_text(headers[1], "文件类型", bold=True)
    set_cell_text(headers[2], "文件路径", bold=True)

    for cell in headers:
        set_cell_shading(cell, "D9EAF7")

    for index, item in enumerate(output_files, start=1):
        if isinstance(item, tuple) and len(item) == 2:
            file_type, file_path = item
        else:
            file_type = "输出文件"
            file_path = item

        cells = table.add_row().cells
        set_cell_text(cells[0], index)
        set_cell_text(cells[1], file_type)
        set_cell_text(cells[2], file_path)

    set_table_borders(table)
    document.add_paragraph()


def normalize_file_paths(file_paths):
    """
    统一处理文件路径列表。
    """
    if file_paths is None:
        return []

    if isinstance(file_paths, (str, Path)):
        return [str(file_paths)]

    return [str(item) for item in file_paths]


def generate_batch_word_report(
    task,
    result,
    output_path="outputs/batch_report.docx"
):
    """
    根据批量数据处理结果生成 Word 报告。

    参数：
        task:
            用户输入的自然语言任务。

        result:
            run_batch_pipeline() 返回的结果字典。

        output_path:
            Word 报告保存路径。

    返回：
        Word 报告的完整路径字符串。
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    set_document_default_font(document)

    # 设置页面边距
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    # 标题
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    title_run = title.add_run("DataPilot 批量数据分析报告")
    title_run.bold = True
    title_run.font.name = "Microsoft YaHei"
    title_run.font.size = Pt(20)
    title_run._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        "Microsoft YaHei"
    )

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle_run = subtitle.add_run(
        "智能数据办公 Agent 自动生成"
    )
    subtitle_run.font.name = "Microsoft YaHei"
    subtitle_run.font.size = Pt(10)
    subtitle_run._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        "Microsoft YaHei"
    )

    document.add_paragraph()

    # 基本信息
    add_heading(document, "一、任务概况", level=1)

    task_text = task if task else "未提供任务说明"

    input_paths = normalize_file_paths(
        result.get("input_paths")
        or result.get("file_paths")
        or result.get("processed_files")
        or []
    )

    output_files = []

    excel_path = result.get("excel_path")
    statistics_path = result.get("statistics_path")
    chart_path = result.get("chart_path") or result.get("plot_path")
    word_path = result.get("word_path")

    if excel_path:
        output_files.append(("清洗后 Excel", str(excel_path)))

    if statistics_path:
        output_files.append(("统计结果 Excel", str(statistics_path)))

    if chart_path:
        output_files.append(("数据图表", str(chart_path)))

    if word_path:
        output_files.append(("Word 报告", str(word_path)))

    original_df = result.get("original_df")
    cleaned_df = result.get("cleaned_df")

    original_rows = ""
    cleaned_rows = ""
    original_columns = ""
    cleaned_columns = ""

    if original_df is not None:
        try:
            original_rows = len(original_df)
            original_columns = len(original_df.columns)
        except Exception:
            original_rows = "未知"
            original_columns = "未知"

    if cleaned_df is not None:
        try:
            cleaned_rows = len(cleaned_df)
            cleaned_columns = len(cleaned_df.columns)
        except Exception:
            cleaned_rows = "未知"
            cleaned_columns = "未知"

    overview_rows = [
        ("生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("用户任务", task_text),
        ("处理文件数量", len(input_paths)),
        ("清洗前数据行数", original_rows if original_rows != "" else "未知"),
        ("清洗后数据行数", cleaned_rows if cleaned_rows != "" else "未知"),
        ("清洗前字段数量", original_columns if original_columns != "" else "未知"),
        ("清洗后字段数量", cleaned_columns if cleaned_columns != "" else "未知"),
    ]

    add_key_value_table(document, overview_rows)

    # 文件列表
    add_heading(document, "二、输入文件", level=1)
    add_file_list_table(document, input_paths)

    # 清洗前质量
    add_heading(document, "三、数据质量检查", level=1)

    before_quality = result.get("before_quality")
    after_quality = result.get("after_quality")

    add_quality_table(
        document,
        "清洗前数据质量",
        before_quality
    )

    add_quality_table(
        document,
        "清洗后数据质量",
        after_quality
    )

    # 清洗记录
    cleaning_log = result.get("cleaning_log")

    if cleaning_log:
        if isinstance(cleaning_log, dict):
            cleaning_items = [
                f"{key}: {value}"
                for key, value in cleaning_log.items()
            ]
        elif isinstance(cleaning_log, list):
            cleaning_items = [str(item) for item in cleaning_log]
        else:
            cleaning_items = [str(cleaning_log)]

        add_list_table(
            document,
            "数据清洗记录",
            cleaning_items
        )

    # 统计结果
    statistics = result.get("statistics")
    add_heading(document, "四、统计分析", level=1)
    add_statistics_table(document, statistics)

    # 图表
    if chart_path:
        chart_file = Path(str(chart_path))

        if chart_file.exists():
            add_heading(document, "五、数据可视化", level=1)

            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

            run = paragraph.add_run()
            run.add_picture(
                str(chart_file),
                width=Inches(6.2)
            )

            caption = document.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

            caption_run = caption.add_run(
                f"图：{chart_file.name}"
            )
            caption_run.font.name = "Microsoft YaHei"
            caption_run.font.size = Pt(9)

    # 输出文件
    add_heading(document, "六、输出文件", level=1)
    add_output_files_table(document, output_files)

    # 结论
    add_heading(document, "七、处理说明", level=1)

    add_body_paragraph(
        document,
        "本报告由 DataPilot 智能数据办公 Agent 自动生成。"
        "系统根据用户输入的自然语言任务，完成了数据读取、"
        "质量检查、数据清洗、统计分析、图表生成及文件导出。"
    )

    add_body_paragraph(
        document,
        "如需进一步分析，可以继续向 DataPilot 提出新的自然语言任务，"
        "例如要求筛选指定日期、比较不同文件、分析异常值或生成专项报告。"
    )

    document.save(str(output_path))

    return str(output_path.resolve())


if __name__ == "__main__":
    print("batch_report_generator.py 已准备完成。")
    print("请通过 DataPilot 批量处理流程调用 generate_batch_word_report()。")