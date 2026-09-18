from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


FONT_NAME = "Microsoft YaHei"
ACCENT = "2F5D7C"
ACCENT_LIGHT = "EAF0F6"
HEADER_FILL = "D9E2F3"
ALT_FILL = "F7F9FC"
BORDER = "B8C4D1"


def _normalize_docx_path(path_value: Any) -> Path:
    path = Path(str(path_value or "")).expanduser()
    if path.suffix.lower() != ".docx":
        raise ValueError("Professional Word 报告只支持 .docx 输出。")
    try:
        return path.resolve()
    except Exception:
        return path.absolute()


def _set_run_font(run, *, size: float = 10.5, bold: bool = False):
    run.font.name = FONT_NAME
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run.font.size = Pt(size)
    run.bold = bold


def _set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)

    for margin_name, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_cell_borders(cell, color: str = BORDER, size: str = "6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def _set_cell_width(cell, width_cm: float):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(Cm(width_cm).emu / 635)))
    tc_w.set(qn("w:type"), "dxa")


def _set_cell_text(
    cell,
    value: Any,
    *,
    bold: bool = False,
    size: float = 9.3,
    align=WD_ALIGN_PARAGRAPH.LEFT,
):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run("" if value is None else str(value))
    _set_run_font(run, size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_cell_margins(cell)


def _set_keep_with_next(paragraph, enabled: bool = True):
    p_pr = paragraph._p.get_or_add_pPr()
    node = p_pr.find(qn("w:keepNext"))
    if enabled and node is None:
        node = OxmlElement("w:keepNext")
        p_pr.append(node)
    elif not enabled and node is not None:
        p_pr.remove(node)


def _set_keep_lines(paragraph, enabled: bool = True):
    p_pr = paragraph._p.get_or_add_pPr()
    node = p_pr.find(qn("w:keepLines"))
    if enabled and node is None:
        node = OxmlElement("w:keepLines")
        p_pr.append(node)


def _set_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = tr_pr.find(qn("w:cantSplit"))
    if node is None:
        node = OxmlElement("w:cantSplit")
        tr_pr.append(node)


def _set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def _add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    _set_run_font(run, size=8.5)


def _configure_document(document: Document):
    styles = document.styles

    normal = styles["Normal"]
    normal.font.name = FONT_NAME
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.font.size = Pt(10.0)
    normal.paragraph_format.space_after = Pt(4.5)
    normal.paragraph_format.line_spacing = 1.18

    heading_specs = {
        "Title": (20, True),
        "Heading 1": (13.5, True),
        "Heading 2": (11.5, True),
    }
    for style_name, (size, bold) in heading_specs.items():
        style = styles[style_name]
        style.font.name = FONT_NAME
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_NAME)
        style.font.size = Pt(size)
        style.font.bold = bold

    for section in document.sections:
        section.top_margin = Cm(1.75)
        section.bottom_margin = Cm(1.65)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header_run = header.add_run("DataPilot")
        _set_run_font(header_run, size=8.0)

        footer = section.footer.paragraphs[0]
        _add_page_number(footer)


def _add_title_block(
    document: Document,
    title: str,
    subtitle: Optional[str],
    metadata: Optional[Dict[str, Any]],
):
    title_p = document.add_paragraph(style="Title")
    title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_p.paragraph_format.space_after = Pt(3)
    _set_keep_with_next(title_p)
    run = title_p.add_run(str(title))
    _set_run_font(run, size=20, bold=True)

    # restrained accent line
    p_pr = title_p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), ACCENT)
    borders.append(bottom)
    p_pr.append(borders)

    if subtitle:
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        _set_keep_with_next(p)
        run = p.add_run(str(subtitle))
        _set_run_font(run, size=9.5)

    if metadata:
        parts = [
            f"{key}：{value}"
            for key, value in metadata.items()
            if value not in (None, "")
        ]
        if parts:
            # Metadata becomes a compact two-column grid rather than one long line.
            cols = 2 if len(parts) > 1 else 1
            rows = (len(parts) + cols - 1) // cols
            table = document.add_table(rows=rows, cols=cols)
            table.alignment = WD_TABLE_ALIGNMENT.LEFT
            table.autofit = True
            index = 0
            for row in table.rows:
                _set_cant_split(row)
                for cell in row.cells:
                    if index < len(parts):
                        _set_cell_shading(cell, "F5F7FA")
                        _set_cell_borders(cell, color="D7DEE7", size="4")
                        _set_cell_text(cell, parts[index], size=8.2)
                    else:
                        _set_cell_shading(cell, "FFFFFF")
                        _set_cell_borders(cell, color="FFFFFF", size="0")
                        _set_cell_text(cell, "", size=8.2)
                    index += 1
            spacer = document.add_paragraph()
            spacer.paragraph_format.space_after = Pt(2)


def _add_heading(document: Document, text: str, level: int = 1):
    paragraph = document.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.space_before = Pt(8 if level == 1 else 5)
    paragraph.paragraph_format.space_after = Pt(4)
    _set_keep_with_next(paragraph)
    run = paragraph.add_run(str(text))
    _set_run_font(run, size=13.5 if level == 1 else 11.5, bold=True)
    return paragraph


def _add_body_paragraph(document: Document, text: Any, *, bold: bool = False):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4.5)
    paragraph.paragraph_format.line_spacing = 1.18
    _set_keep_lines(paragraph)
    run = paragraph.add_run("" if text is None else str(text))
    _set_run_font(run, size=10.0, bold=bold)
    return paragraph


def _add_bullets(document: Document, items: Iterable[Any]):
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2.5)
        paragraph.paragraph_format.line_spacing = 1.12
        _set_keep_lines(paragraph)
        run = paragraph.add_run(str(item))
        _set_run_font(run, size=9.8)


def _kpi_rows(kpis: List[Dict[str, Any]], max_columns: int = 4):
    if not kpis:
        return []
    rows = []
    for start in range(0, len(kpis), max_columns):
        rows.append(kpis[start:start + max_columns])
    return rows


def _add_kpi_table(document: Document, kpis: List[Dict[str, Any]]):
    if not kpis:
        return

    rows_data = _kpi_rows(kpis, max_columns=4)
    table = document.add_table(
        rows=len(rows_data),
        cols=4,
    )
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    usable_width = 17.0
    cell_width = usable_width / 4.0

    for row_index, row_items in enumerate(rows_data):
        row = table.rows[row_index]
        _set_cant_split(row)

        for col_index in range(4):
            cell = row.cells[col_index]
            _set_cell_width(cell, cell_width)
            _set_cell_margins(cell, top=105, bottom=105, start=90, end=90)

            if col_index >= len(row_items):
                _set_cell_shading(cell, "FFFFFF")
                _set_cell_borders(cell, color="FFFFFF", size="0")
                continue

            item = row_items[col_index]
            _set_cell_shading(cell, ACCENT_LIGHT)
            _set_cell_borders(cell, color="D2DCE6", size="5")
            cell.text = ""

            label_p = cell.paragraphs[0]
            label_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            label_p.paragraph_format.space_after = Pt(2)
            label_p.paragraph_format.line_spacing = 1.0
            label_run = label_p.add_run(str(item.get("label", "")))
            _set_run_font(label_run, size=8.0)

            value_p = cell.add_paragraph()
            value_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            value_p.paragraph_format.space_after = Pt(0)
            value_p.paragraph_format.line_spacing = 1.0
            value_run = value_p.add_run(str(item.get("value", "")))
            _set_run_font(value_run, size=13.0, bold=True)

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)


def _add_business_table(
    document: Document,
    columns: List[str],
    rows: List[Dict[str, Any]],
):
    if not columns:
        return

    table = document.add_table(rows=1, cols=len(columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True

    header_row = table.rows[0]
    _set_repeat_table_header(header_row)
    _set_cant_split(header_row)

    for col_index, column in enumerate(columns):
        cell = header_row.cells[col_index]
        _set_cell_shading(cell, HEADER_FILL)
        _set_cell_borders(cell)
        _set_cell_text(
            cell,
            column,
            bold=True,
            size=9.0,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

    for row_index, row_data in enumerate(rows):
        row = table.add_row()
        _set_cant_split(row)
        cells = row.cells
        for col_index, column in enumerate(columns):
            cell = cells[col_index]
            _set_cell_borders(cell)
            if row_index % 2 == 1:
                _set_cell_shading(cell, ALT_FILL)
            _set_cell_text(
                cell,
                row_data.get(column, ""),
                size=8.9,
                align=(
                    WD_ALIGN_PARAGRAPH.CENTER
                    if col_index == 0
                    else WD_ALIGN_PARAGRAPH.LEFT
                ),
            )

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)


def _normalize_sections(
    sections: Optional[List[Dict[str, Any]]],
) -> List[Dict[str, Any]]:
    normalized = []
    for item in sections or []:
        if not isinstance(item, dict):
            raise TypeError("sections 中每一项都必须是字典。")
        title = str(item.get("title") or "").strip()
        if not title:
            raise ValueError("sections 中的 section title 不能为空。")
        section_type = str(
            item.get("type") or "paragraphs"
        ).strip().lower()
        if section_type not in {"paragraphs", "bullets", "table"}:
            raise ValueError(
                f"不支持的 Word section type：{section_type}"
            )
        normalized.append(dict(item))
    return normalized


def create_professional_word_report(
    output_path: str,
    report_title: str,
    *,
    subtitle: Optional[str] = None,
    metadata: Optional[Dict[str, Any]] = None,
    executive_summary: Optional[str] = None,
    kpis: Optional[List[Dict[str, Any]]] = None,
    sections: Optional[List[Dict[str, Any]]] = None,
    source_note: Optional[str] = None,
) -> Dict[str, Any]:
    """
    创建可直接交付的企业风格 Word 报告。

    v5.0 Visual Baseline：
    - KPI 最多 4 个/行，自动换为多行卡片；
    - 标题/正文采用更紧凑的企业汇报节奏；
    - 元数据使用紧凑网格，避免超长单行；
    - 表格表头可跨页重复；
    - 表格行禁止跨页拆分；
    - 标题与后续内容尽量保持在同页；
    - 保留页眉、页码、来源说明和确定性结构。

    sections 支持：
    - {"title": "...", "type": "paragraphs", "content": ["...", "..."]}
    - {"title": "...", "type": "bullets", "items": ["...", "..."]}
    - {"title": "...", "type": "table", "columns": [...], "rows": [{...}]}
    """
    target = _normalize_docx_path(output_path)
    if not str(report_title or "").strip():
        raise ValueError("report_title 不能为空。")

    normalized_sections = _normalize_sections(sections)
    target.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    _configure_document(document)
    _add_title_block(
        document,
        str(report_title).strip(),
        subtitle,
        metadata,
    )

    if executive_summary:
        _add_heading(document, "执行摘要", level=1)
        _add_body_paragraph(document, executive_summary)

    if kpis:
        _add_heading(document, "核心指标", level=1)
        _add_kpi_table(document, list(kpis))

    for section in normalized_sections:
        _add_heading(document, section["title"], level=1)
        section_type = str(
            section.get("type") or "paragraphs"
        ).lower()

        if section_type == "paragraphs":
            content = section.get("content", [])
            if isinstance(content, str):
                content = [content]
            for paragraph_text in content or []:
                _add_body_paragraph(document, paragraph_text)

        elif section_type == "bullets":
            _add_bullets(
                document,
                section.get("items", []) or [],
            )

        elif section_type == "table":
            columns = [
                str(item)
                for item in section.get("columns", []) or []
            ]
            rows = section.get("rows", []) or []
            if not all(isinstance(row, dict) for row in rows):
                raise TypeError(
                    "table section 的 rows 必须是字典列表。"
                )
            _add_business_table(
                document,
                columns,
                rows,
            )

    if source_note:
        _add_heading(document, "数据与来源说明", level=1)
        _add_body_paragraph(document, source_note)

    document.save(str(target))

    return {
        "success": True,
        "output_path": str(target),
        "report_title": str(report_title).strip(),
        "section_count": len(normalized_sections),
        "kpi_count": len(kpis or []),
        "has_executive_summary": bool(executive_summary),
        "has_source_note": bool(source_note),
        "visual_profile": "v5_0_professional_word_baseline",
        "kpi_columns_per_row": 4,
    }


def inspect_professional_word_report(
    file_path: str,
) -> Dict[str, Any]:
    """
    重新打开最终 Word，形成 Completion Gate 可使用的结构证据。
    """
    path = _normalize_docx_path(file_path)
    if not path.exists():
        raise FileNotFoundError(
            f"Word 报告不存在：{path}"
        )

    document = Document(str(path))
    paragraphs = [
        p.text.strip()
        for p in document.paragraphs
        if p.text.strip()
    ]
    headings = []

    for paragraph in document.paragraphs:
        style_name = getattr(
            paragraph.style,
            "name",
            "",
        ) or ""
        if paragraph.text.strip() and (
            style_name == "Title"
            or style_name.startswith("Heading")
        ):
            headings.append(
                paragraph.text.strip()
            )

    tables = []
    for table_index, table in enumerate(document.tables):
        rows = []
        for row in table.rows[:8]:
            rows.append(
                [
                    cell.text.strip()
                    for cell in row.cells
                ]
            )
        tables.append({
            "table_index": table_index,
            "row_count": len(table.rows),
            "column_count": len(table.columns),
            "preview_rows": rows,
        })

    section_info = []
    for index, section in enumerate(document.sections):
        header_text = " ".join(
            p.text.strip()
            for p in section.header.paragraphs
            if p.text.strip()
        )
        footer_text = " ".join(
            p.text.strip()
            for p in section.footer.paragraphs
            if p.text.strip()
        )
        section_info.append({
            "section_index": index,
            "orientation": (
                "landscape"
                if section.orientation == WD_ORIENT.LANDSCAPE
                else "portrait"
            ),
            "header_text": header_text,
            "footer_text": footer_text,
        })

    return {
        "success": True,
        "file_path": str(path),
        "paragraph_count": len(document.paragraphs),
        "nonempty_paragraph_count": len(paragraphs),
        "table_count": len(document.tables),
        "headings": headings,
        "paragraph_preview": paragraphs[:20],
        "tables": tables,
        "sections": section_info,
    }
