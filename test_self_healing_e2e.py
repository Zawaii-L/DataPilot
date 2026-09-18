from __future__ import annotations

import json
import re
import tempfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

import agent as agent_module
from agent import DataPilotAgent
from task_planner import TaskPlan


class FakeTaskPlanner:
    """
    不调用外部 LLM。

    本测试只验证真实：
    agent.py -> Workspace -> AgentLoop -> ToolExecutor -> ToolPreflight
    -> ToolFailureRecovery -> 下一轮 Agent state -> 自纠 -> 最终交付物
    -> Verification Engine -> Completion Gate。
    """

    def __init__(self, *args, **kwargs):
        pass

    def create_plan(self, task, context=None):
        return TaskPlan(
            task_goal=(
                "读取真实销售 Excel，按城市统计销售额合计，"
                "生成最终 Excel，并在工具失败后依据 RecoveryHint 自纠。"
            ),
            evidence_requirements=[
                "城市销售汇总必须来自真实源 Excel。",
            ],
            source_requirements=[
                "必须读取用户提供的真实 Excel。",
            ],
            deliverable_requirements=[
                "生成最终 Excel 交付物。",
            ],
            execution_requirements=[
                "读取真实销售 Excel。",
                "按城市统计销售额合计。",
                "如果列名错误导致工具失败，必须根据真实 Observation 修正后继续。",
            ],
            verification_requirements=[
                "生成后重新读取并检查最终 Excel。",
                "核对澳门、横琴、珠海的销售额合计和销售冠军。",
            ],
            safety_requirements=[
                "不得覆盖源 Excel。",
                "最终 Excel 必须位于 deliverables_dir。",
            ],
            assumptions=[],
        )


class SelfHealingCompletions:
    """
    只替代 LLM 的下一步决策。

    关键点：
    第 2 轮故意把 target_column 写成不存在的“销售金额”。
    第 3 轮不盲目继续，而是先检查真实 Agent state：
    - step_2 必须失败；
    - recovery.category 必须为 missing_column；
    - recommended_actions / avoid_actions 必须已进入 Observation；
    - step_1 的真实 columns 必须仍然可见。
    满足后才返回修正后的 target_column='销售额'。

    因此测试的不是“预先写死第二次一定成功”，而是验证 RecoveryHint
    确实穿过真实 AgentLoop state 后，决策层才能执行纠错分支。
    """

    def __init__(self, source_path: Path):
        self.source_path = str(source_path.resolve())
        self.calls = []
        self.decision_index = 0
        self.excel_path = None
        self.recovery_seen = False
        self.real_columns_seen_after_failure = False

    @staticmethod
    def _extract_user_content(messages):
        return next(
            (
                item.get("content", "")
                for item in messages
                if item.get("role") == "user"
            ),
            "",
        )

    @classmethod
    def _extract_state(cls, messages):
        content = cls._extract_user_content(messages)

        marker = "当前真实执行状态："
        if marker not in content:
            raise AssertionError("没有找到 AgentLoop 当前执行状态。")

        state_text = content.split(marker, 1)[1].strip()

        # Prompt 在 state JSON 后还可能有额外说明。
        decoder = json.JSONDecoder()
        state, _ = decoder.raw_decode(state_text)
        return state

    @staticmethod
    def _extract_deliverables_dir_from_state(state):
        runtime = state.get("runtime_context", {})
        workspace = runtime.get("workspace", {})
        value = workspace.get("deliverables_dir")
        if not value:
            raise AssertionError(
                f"真实 Agent state 缺少 deliverables_dir：{workspace}"
            )
        return value

    def create(self, **kwargs):
        self.calls.append(kwargs)
        messages = kwargs.get("messages", [])
        state = self._extract_state(messages)

        if self.excel_path is None:
            deliverables_dir = self._extract_deliverables_dir_from_state(state)
            self.excel_path = str(
                (
                    Path(deliverables_dir)
                    / "城市销售自纠分析.xlsx"
                ).resolve()
            )

        index = self.decision_index

        if index == 0:
            decision = {
                "action_type": "tool",
                "tool": "read_office_data",
                "arguments": {
                    "file_path": self.source_path,
                },
                "purpose": "读取真实销售源数据并取得真实列名。",
            }

        elif index == 1:
            completed = state.get("completed_tool_steps", [])
            if not completed or not completed[0].get("success"):
                raise AssertionError(
                    "第 2 轮前没有真实 read_office_data 成功 Observation。"
                )

            columns = (
                completed[0]
                .get("observation", {})
                .get("columns", [])
            )
            if "销售额" not in columns:
                raise AssertionError(
                    f"真实源数据列名异常：{columns}"
                )

            # 故意制造真实工具失败。
            decision = {
                "action_type": "tool",
                "tool": "group_statistics",
                "arguments": {
                    "df": {"$ref": "step_1.output"},
                    "group_by": "城市",
                    "target_column": "销售金额",
                    "operation": "sum",
                },
                "purpose": (
                    "故意使用不存在的销售金额列，"
                    "验证真实 Tool Failure Recovery。"
                ),
            }

        elif index == 2:
            completed = state.get("completed_tool_steps", [])
            if len(completed) < 2:
                raise AssertionError(
                    f"失败步骤没有进入 Agent state：{completed}"
                )

            failed_step = completed[1]
            if failed_step.get("success") is not False:
                raise AssertionError(
                    f"step_2 本应真实失败：{failed_step}"
                )

            observation = failed_step.get("observation", {})
            recovery = observation.get("recovery", {})

            if recovery.get("category") != "missing_column":
                raise AssertionError(
                    f"没有收到 missing_column RecoveryHint：{observation}"
                )

            if recovery.get("recoverable") is not True:
                raise AssertionError(
                    f"missing_column 应可恢复：{recovery}"
                )

            if not recovery.get("recommended_actions"):
                raise AssertionError(
                    f"RecoveryHint 缺少 recommended_actions：{recovery}"
                )

            if not recovery.get("avoid_actions"):
                raise AssertionError(
                    f"RecoveryHint 缺少 avoid_actions：{recovery}"
                )

            first_columns = (
                completed[0]
                .get("observation", {})
                .get("columns", [])
            )
            if "销售额" not in first_columns:
                raise AssertionError(
                    f"失败后真实 columns 丢失：{first_columns}"
                )

            failed_counts = state.get("failed_tool_counts", {})
            if failed_counts.get("group_statistics") != 1:
                raise AssertionError(
                    f"failed_tool_counts 异常：{failed_counts}"
                )

            self.recovery_seen = True
            self.real_columns_seen_after_failure = True

            # 根据真实 columns + RecoveryHint 修正。
            decision = {
                "action_type": "tool",
                "tool": "group_statistics",
                "arguments": {
                    "df": {"$ref": "step_1.output"},
                    "group_by": "城市",
                    "target_column": "销售额",
                    "operation": "sum",
                },
                "purpose": (
                    "根据 missing_column RecoveryHint 和真实 columns，"
                    "把错误列名销售金额修正为销售额。"
                ),
            }

        elif index == 3:
            decision = {
                "action_type": "tool",
                "tool": "sort_data",
                "arguments": {
                    "df": {"$ref": "step_3.output"},
                    "column": "销售额_合计",
                    "ascending": False,
                },
                "purpose": "对自纠后的真实城市销售汇总降序排序。",
            }

        elif index == 4:
            decision = {
                "action_type": "tool",
                "tool": "create_professional_excel_report",
                "arguments": {
                    "output_path": self.excel_path,
                    "dataframe": {"$ref": "step_4.output"},
                    "default_sheet_name": "城市销售汇总",
                    "report_title": "城市销售自纠分析报告",
                    "subtitle": "DataPilot Self-Healing E2E",
                    "kpis": [
                        {"label": "销售冠军", "value": "澳门"},
                        {"label": "最高销售额", "value": 186},
                        {"label": "销售总额", "value": 468},
                    ],
                    "charts": [
                        {
                            "type": "bar",
                            "sheet_name": "城市销售汇总",
                            "category_column": "城市",
                            "value_column": "销售额_合计",
                            "title": "城市销售额对比",
                            "anchor": "E7",
                        }
                    ],
                },
                "purpose": "使用自纠后的真实分析结果生成最终 Excel。",
            }

        elif index == 5:
            decision = {
                "action_type": "tool",
                "tool": "inspect_professional_excel_report",
                "arguments": {
                    "file_path": self.excel_path,
                },
                "purpose": "重新打开最终 Excel，取得最终验收证据。",
            }

        else:
            decision = {
                "action_type": "finish",
                "final_answer": (
                    "首次分组因错误列名失败后，已依据 RecoveryHint "
                    "和真实列名完成自纠；最终 Excel 已生成并重新读取核验。"
                ),
            }

        self.decision_index += 1

        return SimpleNamespace(
            choices=[
                SimpleNamespace(
                    message=SimpleNamespace(
                        content=json.dumps(
                            decision,
                            ensure_ascii=False,
                        )
                    )
                )
            ]
        )


class FakeClient:
    def __init__(self, source_path: Path):
        self.chat = SimpleNamespace(
            completions=SelfHealingCompletions(source_path)
        )


def assert_true(value, message):
    if not value:
        raise AssertionError(message)


def main():
    print("=" * 72)
    print("DataPilot v5.0 Self-Healing Agent E2E 确定性测试")
    print("=" * 72)

    with tempfile.TemporaryDirectory(
        prefix="datapilot_self_healing_e2e_"
    ) as temp_dir:
        root = Path(temp_dir)
        source_path = root / "销售数据.xlsx"
        output_root = root / "outputs"

        source_df = pd.DataFrame(
            {
                "城市": ["珠海", "澳门", "横琴", "澳门"],
                "销售额": [128, 100, 154, 86],
            }
        )
        source_df.to_excel(source_path, index=False)
        original_bytes = source_path.read_bytes()

        fake_client = FakeClient(source_path)

        datapilot = DataPilotAgent.__new__(DataPilotAgent)
        datapilot.api_key = "test-key"
        datapilot.base_url = "https://example.invalid"
        datapilot.model = "test-model"
        datapilot.progress_callback = None
        datapilot.client = fake_client

        user_task = (
            "读取销售数据，按城市汇总销售额并生成专业 Excel。"
            "如果执行过程中工具失败，必须根据真实错误和 Observation 自行修正，"
            "不要覆盖原文件。生成后重新读取最终 Excel，"
            "核对城市汇总和销售冠军后再完成任务。"
        )

        print("\n测试 1：真实 agent.py 建立 Self-Healing TaskPlan")
        with patch.object(agent_module, "TaskPlanner", FakeTaskPlanner):
            result = datapilot.execute_v31_agent_task(
                user_task=user_task,
                input_paths=[str(source_path)],
                output_dir=str(output_root),
                max_iterations=9,
            )

        task_plan = result.get("task_plan", {})
        assert_true(isinstance(task_plan, dict), "没有返回 TaskPlan。")
        assert_true(
            any(
                "失败" in str(item) or "修正" in str(item)
                for item in task_plan.get("execution_requirements", [])
            ),
            f"TaskPlan 没有自纠要求：{task_plan}",
        )
        print("PASS")

        history = result.get("tool_results", [])
        assert_true(isinstance(history, list), history)

        print("\n测试 2：第一次 group_statistics 真实失败")
        group_results = [
            item
            for item in history
            if isinstance(item, dict)
            and item.get("tool_name") == "group_statistics"
        ]
        assert_true(
            len(group_results) == 2,
            f"group_statistics 调用次数不是 2：{group_results}",
        )
        assert_true(
            group_results[0].get("success") is False,
            f"第一次分组没有真实失败：{group_results[0]}",
        )
        assert_true(
            "销售金额" in str(group_results[0].get("error_message", "")),
            group_results[0],
        )
        print("PASS")

        print("\n测试 3：失败后的 Agent state 收到 missing_column RecoveryHint")
        completions = fake_client.chat.completions
        assert_true(
            completions.recovery_seen,
            "决策层没有真实看到 missing_column RecoveryHint。",
        )
        print("PASS")

        print("\n测试 4：失败后仍保留真实源数据 columns")
        assert_true(
            completions.real_columns_seen_after_failure,
            "失败后没有保留并复用真实 columns。",
        )
        print("PASS")

        print("\n测试 5：Agent 没有原样重复错误列名")
        assert_true(
            group_results[1].get("success") is True,
            f"第二次分组没有恢复成功：{group_results[1]}",
        )
        second_args = group_results[1].get("arguments", {})
        assert_true(
            second_args.get("target_column") == "销售额",
            f"没有修正为真实列名：{second_args}",
        )
        print("PASS")

        print("\n测试 6：自纠后的城市销售汇总正确")
        # DataPilotAgent 对外返回的 tool_results 使用 ToolExecutionResult.to_dict()。
        # 对 DataFrame output，公开结果不会携带 preview；真实业务 preview
        # 已经通过 AgentLoop state 进入下一轮 Observation。
        # 因此这里从第 4 轮真实 Agent state 读取 step_3 的成功 Observation，
        # 而不是错误地假设 result["tool_results"][].output 仍包含 preview。
        fourth_call_messages = completions.calls[3].get("messages", [])
        state_after_recovery = completions._extract_state(
            fourth_call_messages
        )
        recovered_steps = state_after_recovery.get(
            "completed_tool_steps",
            [],
        )
        assert_true(
            len(recovered_steps) >= 3,
            f"自纠成功步骤没有进入 Agent state：{recovered_steps}",
        )
        recovered_step = recovered_steps[2]
        assert_true(
            recovered_step.get("tool") == "group_statistics"
            and recovered_step.get("success") is True,
            f"step_3 不是成功的 group_statistics：{recovered_step}",
        )
        recovered_observation = recovered_step.get("observation", {})
        preview = recovered_observation.get("preview", [])
        recovered = {
            str(item["城市"]): item["销售额_合计"]
            for item in preview
        }
        expected = {"澳门": 186, "横琴": 154, "珠海": 128}
        assert_true(
            recovered == expected,
            f"自纠后的汇总错误：{recovered}",
        )
        print("PASS")

        print("\n测试 7：失败后继续完成排序与最终 Excel 生成")
        tool_names = [
            item.get("tool_name")
            for item in history
            if isinstance(item, dict)
        ]
        assert_true("sort_data" in tool_names, tool_names)
        assert_true(
            "create_professional_excel_report" in tool_names,
            tool_names,
        )
        assert_true(
            "inspect_professional_excel_report" in tool_names,
            tool_names,
        )
        print("PASS")

        print("\n测试 8：最终 Excel 真实存在且业务数据正确")
        excel_path = Path(completions.excel_path)
        assert_true(excel_path.exists(), f"最终 Excel 不存在：{excel_path}")

        workbook = load_workbook(excel_path, data_only=False)
        worksheet = workbook["城市销售汇总"]

        header_row = None
        for row_index in range(1, worksheet.max_row + 1):
            values = [
                worksheet.cell(row_index, col).value
                for col in range(
                    1,
                    min(worksheet.max_column, 4) + 1,
                )
            ]
            if "城市" in values and "销售额_合计" in values:
                header_row = row_index
                break

        assert_true(header_row is not None, "最终 Excel 没找到业务表头。")

        actual = {}
        for row_index in range(header_row + 1, worksheet.max_row + 1):
            city = worksheet.cell(row_index, 1).value
            amount = worksheet.cell(row_index, 2).value
            if city in expected:
                actual[str(city)] = amount

        assert_true(actual == expected, f"最终 Excel 数据错误：{actual}")
        print("PASS")

        print("\n测试 9：最终 Excel 在最后写入后被重新读取")
        create_index = tool_names.index("create_professional_excel_report")
        inspect_index = tool_names.index("inspect_professional_excel_report")
        assert_true(
            inspect_index > create_index,
            f"inspect 顺序错误：{tool_names}",
        )

        inspection = next(
            item.get("output", {})
            for item in history
            if isinstance(item, dict)
            and item.get("tool_name")
            == "inspect_professional_excel_report"
        )
        kpis = inspection.get("kpi_map", {})
        assert_true(kpis.get("销售冠军") == "澳门", kpis)
        assert_true(kpis.get("最高销售额") == 186, kpis)
        assert_true(kpis.get("销售总额") == 468, kpis)
        print("PASS")

        print("\n测试 10：Completion Gate 最终 PASS")
        verification = result.get("verification_report", {})
        assert_true(
            verification.get("verified") is True,
            f"Completion Gate 未 PASS：{verification}",
        )
        assert_true(result.get("success") is True, result)
        assert_true(
            result.get("stop_reason") == "completed",
            result.get("stop_reason"),
        )
        print("PASS")

        print("\n测试 11：早期工具失败被后续成功恢复")
        checks = verification.get("checks", [])
        tool_failure_check = next(
            (
                item
                for item in checks
                if item.get("check_id")
                == "tool_failures_resolved"
            ),
            None,
        )
        assert_true(
            tool_failure_check is not None,
            f"缺少 tool_failures_resolved：{checks}",
        )
        assert_true(
            tool_failure_check.get("passed") is True,
            tool_failure_check,
        )
        print("PASS")

        print("\n测试 12：源 Excel 未被覆盖")
        assert_true(source_path.exists(), "源 Excel 被删除。")
        assert_true(
            source_path.read_bytes() == original_bytes,
            "源 Excel 内容被修改。",
        )
        assert_true(
            excel_path.resolve() != source_path.resolve(),
            "最终 Excel 覆盖了源 Excel。",
        )
        print("PASS")

        print("\n" + "=" * 72)
        print("Self-Healing Agent E2E：12/12 PASS")
        print("=" * 72)


if __name__ == "__main__":
    main()
