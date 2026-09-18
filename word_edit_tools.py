from copy import deepcopy
from pathlib import Path
from typing import Any, Dict, List, Optional

from docx import Document

from document_tools import normalize_document_path


def validate_word_file(file_path) -> Path:
    """
    校验 Word 文件。

    当前 v3.2 只编辑 .docx。
    """
    path = normalize_document_path(file_path)

    if not path.exists():
        raise FileNotFoundError(
            f"Word 文件不存在：{path}"
        )

    if not path.is_file():
        raise ValueError(
            f"路径不是文件：{path}"
        )

    if path.suffix.lower() != ".docx":
        raise ValueError(
            f"当前 Word 编辑工具只支持 .docx：{path}"
        )

    return path


def build_output_path(
    source_path,
    output_path=None,
    suffix="_DataPilot编辑",
) -> Path:
    """
    生成输出路径。

    默认不覆盖原文件，而是在原文件名后追加后缀。
    """
    source = validate_word_file(source_path)

    if output_path:
        target = Path(output_path).expanduser()

        if target.suffix.lower() != ".docx":
            target = target.with_suffix(".docx")

        try:
            target = target.resolve()
        except Exception:
            target = target.absolute()
    else:
        target = source.with_name(
            f"{source.stem}{suffix}.docx"
        )

    if target == source:
        raise ValueError(
            "为保护原文件，默认禁止直接覆盖源 Word。"
        )

    target.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    return target


def _replace_text_in_paragraph(
    paragraph,
    old_text: str,
    new_text: str,
) -> int:
    """
    在一个段落中替换文本。

    优先在单个 run 内替换，以尽量保留格式。
    如果目标文本跨 run，则回退为重建该段落文本。
    """
    if not old_text:
        raise ValueError(
            "old_text 不能为空。"
        )

    replacement_count = 0

    for run in paragraph.runs:
        if old_text in run.text:
            count = run.text.count(old_text)
            run.text = run.text.replace(
                old_text,
                new_text,
            )
            replacement_count += count

    if replacement_count > 0:
        return replacement_count

    full_text = paragraph.text

    if old_text not in full_text:
        return 0

    count = full_text.count(old_text)
    replaced_text = full_text.replace(
        old_text,
        new_text,
    )

    if paragraph.runs:
        first_run = paragraph.runs[0]
        first_run.text = replaced_text

        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(
            replaced_text
        )

    return count


def replace_text(
    file_path,
    old_text: str,
    new_text: str,
    output_path=None,
    include_tables: bool = True,
) -> Dict[str, Any]:
    """
    全文替换 Word 中的指定文本。

    支持正文和表格。
    默认另存新文件，不覆盖原文件。
    """
    source = validate_word_file(
        file_path
    )
    target = build_output_path(
        source,
        output_path=output_path,
    )

    document = Document(
        str(source)
    )

    paragraph_replacements = 0
    table_replacements = 0

    for paragraph in document.paragraphs:
        paragraph_replacements += (
            _replace_text_in_paragraph(
                paragraph,
                old_text,
                new_text,
            )
        )

    if include_tables:
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        table_replacements += (
                            _replace_text_in_paragraph(
                                paragraph,
                                old_text,
                                new_text,
                            )
                        )

    document.save(
        str(target)
    )

    return {
        "success": True,
        "source_path": str(source),
        "output_path": str(target),
        "old_text": old_text,
        "new_text": new_text,
        "paragraph_replacements": (
            paragraph_replacements
        ),
        "table_replacements": (
            table_replacements
        ),
        "total_replacements": (
            paragraph_replacements
            + table_replacements
        ),
    }


def replace_paragraph(
    file_path,
    search_text: str,
    new_text: str,
    output_path=None,
    match_mode: str = "contains",
    replace_all: bool = False,
) -> Dict[str, Any]:
    """
    按段落定位并替换整段内容。

    match_mode:
    - contains：段落包含 search_text
    - exact：段落内容完全等于 search_text
    """
    source = validate_word_file(
        file_path
    )
    target = build_output_path(
        source,
        output_path=output_path,
    )

    if match_mode not in {
        "contains",
        "exact",
    }:
        raise ValueError(
            "match_mode 只支持 contains 或 exact。"
        )

    document = Document(
        str(source)
    )

    changed = 0

    for paragraph in document.paragraphs:
        current = paragraph.text

        matched = (
            current == search_text
            if match_mode == "exact"
            else search_text in current
        )

        if not matched:
            continue

        if paragraph.runs:
            paragraph.runs[0].text = new_text

            for run in paragraph.runs[1:]:
                run.text = ""
        else:
            paragraph.add_run(
                new_text
            )

        changed += 1

        if not replace_all:
            break

    document.save(
        str(target)
    )

    return {
        "success": True,
        "source_path": str(source),
        "output_path": str(target),
        "matched_paragraphs": changed,
        "search_text": search_text,
        "new_text": new_text,
    }


def append_paragraph(
    file_path,
    text: str,
    output_path=None,
    style: Optional[str] = None,
) -> Dict[str, Any]:
    """
    在 Word 文档末尾追加一个段落。
    """
    source = validate_word_file(
        file_path
    )
    target = build_output_path(
        source,
        output_path=output_path,
    )

    document = Document(
        str(source)
    )

    if style:
        document.add_paragraph(
            text,
            style=style,
        )
    else:
        document.add_paragraph(
            text
        )

    document.save(
        str(target)
    )

    return {
        "success": True,
        "source_path": str(source),
        "output_path": str(target),
        "appended_text": text,
    }


def delete_paragraphs(
    file_path,
    search_text: str,
    output_path=None,
    match_mode: str = "contains",
    delete_all: bool = True,
) -> Dict[str, Any]:
    """
    删除匹配的正文段落。
    """
    source = validate_word_file(
        file_path
    )
    target = build_output_path(
        source,
        output_path=output_path,
    )

    if match_mode not in {
        "contains",
        "exact",
    }:
        raise ValueError(
            "match_mode 只支持 contains 或 exact。"
        )

    document = Document(
        str(source)
    )

    deleted = 0

    for paragraph in list(
        document.paragraphs
    ):
        current = paragraph.text

        matched = (
            current == search_text
            if match_mode == "exact"
            else search_text in current
        )

        if not matched:
            continue

        element = paragraph._element
        parent = element.getparent()

        if parent is not None:
            parent.remove(element)
            deleted += 1

        if not delete_all:
            break

    document.save(
        str(target)
    )

    return {
        "success": True,
        "source_path": str(source),
        "output_path": str(target),
        "deleted_paragraphs": deleted,
        "search_text": search_text,
    }


def update_table_cell(
    file_path,
    table_index: int,
    row_index: int,
    column_index: int,
    new_text: str,
    output_path=None,
) -> Dict[str, Any]:
    """
    修改指定 Word 表格单元格。

    table_index / row_index / column_index 均从 0 开始。
    """
    source = validate_word_file(
        file_path
    )
    target = build_output_path(
        source,
        output_path=output_path,
    )

    document = Document(
        str(source)
    )

    if table_index < 0 or table_index >= len(
        document.tables
    ):
        raise IndexError(
            f"table_index 超出范围：{table_index}"
        )

    table = document.tables[
        table_index
    ]

    if row_index < 0 or row_index >= len(
        table.rows
    ):
        raise IndexError(
            f"row_index 超出范围：{row_index}"
        )

    row = table.rows[
        row_index
    ]

    if (
        column_index < 0
        or column_index >= len(row.cells)
    ):
        raise IndexError(
            f"column_index 超出范围：{column_index}"
        )

    cell = row.cells[
        column_index
    ]
    old_text = cell.text
    cell.text = new_text

    document.save(
        str(target)
    )

    return {
        "success": True,
        "source_path": str(source),
        "output_path": str(target),
        "table_index": table_index,
        "row_index": row_index,
        "column_index": column_index,
        "old_text": old_text,
        "new_text": new_text,
    }


def apply_word_edits(
    file_path,
    operations: List[Dict[str, Any]],
    output_path=None,
) -> Dict[str, Any]:
    """
    一次性对 Word 执行多个编辑操作，只保存一次。

    支持 action：
    - replace_text
    - replace_paragraph
    - append_paragraph
    - delete_paragraphs
    - update_table_cell
    """
    source = validate_word_file(
        file_path
    )
    target = build_output_path(
        source,
        output_path=output_path,
    )

    if not isinstance(
        operations,
        list,
    ):
        raise TypeError(
            "operations 必须是列表。"
        )

    document = Document(
        str(source)
    )
    logs = []

    for index, operation in enumerate(
        operations,
        start=1,
    ):
        if not isinstance(
            operation,
            dict,
        ):
            raise TypeError(
                f"第 {index} 个操作必须是字典。"
            )

        action = str(
            operation.get("action")
            or ""
        ).strip()

        if action == "replace_text":
            old_text = str(
                operation.get("old_text")
                or ""
            )
            new_text = str(
                operation.get("new_text")
                or ""
            )

            if not old_text:
                raise ValueError(
                    f"第 {index} 个 replace_text 缺少 old_text。"
                )

            paragraph_count = 0
            table_count = 0

            for paragraph in document.paragraphs:
                paragraph_count += (
                    _replace_text_in_paragraph(
                        paragraph,
                        old_text,
                        new_text,
                    )
                )

            if operation.get(
                "include_tables",
                True,
            ):
                for table in document.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            for paragraph in cell.paragraphs:
                                table_count += (
                                    _replace_text_in_paragraph(
                                        paragraph,
                                        old_text,
                                        new_text,
                                    )
                                )

            logs.append(
                {
                    "action": action,
                    "total_replacements": (
                        paragraph_count
                        + table_count
                    ),
                }
            )

        elif action == "replace_paragraph":
            search_text = str(
                operation.get("search_text")
                or ""
            )
            new_text = str(
                operation.get("new_text")
                or ""
            )
            match_mode = str(
                operation.get("match_mode")
                or "contains"
            )
            replace_all = bool(
                operation.get(
                    "replace_all",
                    False,
                )
            )

            changed = 0

            for paragraph in document.paragraphs:
                current = paragraph.text

                matched = (
                    current == search_text
                    if match_mode == "exact"
                    else search_text in current
                )

                if not matched:
                    continue

                if paragraph.runs:
                    paragraph.runs[0].text = new_text

                    for run in paragraph.runs[1:]:
                        run.text = ""
                else:
                    paragraph.add_run(
                        new_text
                    )

                changed += 1

                if not replace_all:
                    break

            logs.append(
                {
                    "action": action,
                    "matched_paragraphs": changed,
                }
            )

        elif action == "append_paragraph":
            text = str(
                operation.get("text")
                or ""
            )
            style = operation.get(
                "style"
            )

            if style:
                document.add_paragraph(
                    text,
                    style=style,
                )
            else:
                document.add_paragraph(
                    text
                )

            logs.append(
                {
                    "action": action,
                    "appended": True,
                }
            )

        elif action == "delete_paragraphs":
            search_text = str(
                operation.get("search_text")
                or ""
            )
            match_mode = str(
                operation.get("match_mode")
                or "contains"
            )
            delete_all = bool(
                operation.get(
                    "delete_all",
                    True,
                )
            )

            deleted = 0

            for paragraph in list(
                document.paragraphs
            ):
                current = paragraph.text

                matched = (
                    current == search_text
                    if match_mode == "exact"
                    else search_text in current
                )

                if not matched:
                    continue

                element = paragraph._element
                parent = element.getparent()

                if parent is not None:
                    parent.remove(element)
                    deleted += 1

                if not delete_all:
                    break

            logs.append(
                {
                    "action": action,
                    "deleted_paragraphs": deleted,
                }
            )

        elif action == "update_table_cell":
            table_index = int(
                operation.get(
                    "table_index",
                    0,
                )
            )
            row_index = int(
                operation.get(
                    "row_index",
                    0,
                )
            )
            column_index = int(
                operation.get(
                    "column_index",
                    0,
                )
            )
            new_text = str(
                operation.get("new_text")
                or ""
            )

            if (
                table_index < 0
                or table_index >= len(
                    document.tables
                )
            ):
                raise IndexError(
                    f"第 {index} 个操作 table_index 超出范围。"
                )

            table = document.tables[
                table_index
            ]

            if (
                row_index < 0
                or row_index >= len(
                    table.rows
                )
            ):
                raise IndexError(
                    f"第 {index} 个操作 row_index 超出范围。"
                )

            row = table.rows[
                row_index
            ]

            if (
                column_index < 0
                or column_index >= len(
                    row.cells
                )
            ):
                raise IndexError(
                    f"第 {index} 个操作 column_index 超出范围。"
                )

            old_text = row.cells[
                column_index
            ].text

            row.cells[
                column_index
            ].text = new_text

            logs.append(
                {
                    "action": action,
                    "old_text": old_text,
                    "new_text": new_text,
                }
            )

        else:
            raise ValueError(
                f"暂不支持的 Word 编辑 action：{action}"
            )

    document.save(
        str(target)
    )

    return {
        "success": True,
        "source_path": str(source),
        "output_path": str(target),
        "operation_count": len(operations),
        "operations": logs,
    }


def main():
    print("=" * 70)
    print("DataPilot v3.2 Word 编辑工具")
    print("=" * 70)
    print(
        "支持：全文替换、整段替换、追加段落、删除段落、"
        "表格单元格修改、多操作一次执行。"
    )
    print("默认另存新文件，不覆盖源 Word。")
    print("=" * 70)


if __name__ == "__main__":
    main()
