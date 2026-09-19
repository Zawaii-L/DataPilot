from pathlib import Path

from docx import Document

from core.agent import DataPilotAgent
from document_tools import read_word_file


def create_test_document(file_path: Path):
    """
    创建 v3.2 Agent Loop 测试 Word。
    """
    document = Document()

    document.add_heading(
        "2025年度工作总结",
        level=1,
    )

    document.add_paragraph(
        "本报告总结2025年度重点工作。"
    )

    document.add_paragraph(
        "临时说明：这一段需要删除。"
    )

    document.add_paragraph(
        "DataPilot 项目已完成基础办公自动化能力。"
    )

    table = document.add_table(
        rows=2,
        cols=2,
    )

    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "年度"

    table.cell(1, 0).text = "DataPilot"
    table.cell(1, 1).text = "2025年度"

    document.save(
        str(file_path)
    )


def main():
    print("=" * 70)
    print("DataPilot v3.2 真实 Word Agent Loop 测试")
    print("=" * 70)

    project_dir = Path.cwd()

    test_dir = (
        project_dir
        / "outputs"
        / "v32_word_agent_test"
    )

    test_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    source_file = (
        test_dir
        / "v32_word_agent_source.docx"
    )

    create_test_document(
        source_file
    )

    print(
        f"\n测试 Word 已创建：{source_file.resolve()}"
    )

    print("\n原始文档内容：")
    print("-" * 70)
    print(
        read_word_file(
            source_file
        )
    )
    print("-" * 70)

    task = (
        "请修改我提供的 Word 文件。"
        "先读取并理解原文件，然后完成以下修改："
        "把文档中所有“2025年度”替换成“2026年度”，"
        "包括正文和表格；"
        "删除包含“临时说明”的段落；"
        "最后在文档末尾追加一段“审核状态：已完成。”。"
        "不要覆盖原文件，必须另存为新的 Word 文件。"
        "修改完成后，请再次读取新文件核对修改是否全部成功，"
        "确认无误后再结束任务。"
    )

    print("\n用户任务：")
    print(task)

    print("\n" + "=" * 70)
    print("开始运行真实 DataPilot Agent Loop")
    print("=" * 70)

    agent = DataPilotAgent()

    result = agent.execute_v31_agent_task(
        user_task=task,
        input_paths=[
            str(source_file.resolve())
        ],
        output_dir=str(
            test_dir.resolve()
        ),
        max_iterations=10,
    )

    print("\n" + "=" * 70)
    print("Agent Loop 返回结果")
    print("=" * 70)

    print(
        "SUCCESS =",
        result.get("success")
    )

    print(
        "STOP_REASON =",
        result.get("stop_reason")
    )

    print(
        "ITERATIONS =",
        result.get("iterations")
    )

    print(
        "TOOL_COUNT =",
        result.get("tool_count")
    )

    print("\n工具调用链：")

    tool_results = (
        result.get("tool_results")
        or []
    )

    for index, item in enumerate(
        tool_results,
        start=1,
    ):
        print(
            f"{index}. "
            f"{item.get('tool_name')} "
            f"SUCCESS={item.get('success')}"
        )

    print("\nAgent 最终回答：")
    print(
        result.get("final_answer")
        or result.get("answer")
        or ""
    )

    # ------------------------------------------------------------
    # 找出 Agent 真正生成的新 Word
    # ------------------------------------------------------------

    source_resolved = (
        source_file
        .resolve()
    )

    candidate_files = []

    for file_path in test_dir.glob(
        "*.docx"
    ):
        try:
            resolved = (
                file_path
                .resolve()
            )
        except Exception:
            continue

        if resolved == source_resolved:
            continue

        candidate_files.append(
            resolved
        )

    if not candidate_files:
        raise AssertionError(
            "Agent 没有生成新的 Word 文件。"
        )

    candidate_files.sort(
        key=lambda path: (
            path.stat().st_mtime
        ),
        reverse=True,
    )

    edited_file = (
        candidate_files[0]
    )

    print(
        f"\n检测到 Agent 生成的新 Word：{edited_file}"
    )

    edited_text = read_word_file(
        edited_file
    )

    print("\n修改后文档内容：")
    print("-" * 70)
    print(edited_text)
    print("-" * 70)

    # ------------------------------------------------------------
    # 确定性验证
    # ------------------------------------------------------------

    print("\n开始确定性验证……")

    assert (
        result.get("success") is True
    ), "Agent Loop 没有成功结束。"

    assert (
        "2025年度"
        not in edited_text
    ), "新文件中仍然存在 2025年度。"

    assert (
        edited_text.count(
            "2026年度"
        )
        >= 3
    ), "2026年度替换数量不足。"

    assert (
        "临时说明"
        not in edited_text
    ), "临时说明段落没有删除。"

    assert (
        "审核状态：已完成。"
        in edited_text
    ), "没有成功追加审核状态。"

    # ------------------------------------------------------------
    # 验证源文件没有被覆盖
    # ------------------------------------------------------------

    original_text = read_word_file(
        source_file
    )

    assert (
        "2025年度"
        in original_text
    ), "源 Word 被意外修改。"

    assert (
        "临时说明"
        in original_text
    ), "源 Word 被意外修改。"

    assert (
        "审核状态：已完成。"
        not in original_text
    ), "源 Word 被意外修改。"

    # ------------------------------------------------------------
    # 验证 Agent 确实调用了 Word 编辑工具
    # ------------------------------------------------------------

    tool_names = [
        item.get("tool_name")
        for item in tool_results
    ]

    assert (
        "apply_word_edits"
        in tool_names
    ), (
        "Agent 没有调用 apply_word_edits，"
        "说明没有真正走 v3.2 Word 编辑工具。"
    )

    print("\n" + "=" * 70)
    print("DataPilot v3.2 真实 Word Agent Loop 测试通过！")
    print("=" * 70)

    print()
    print("已验证：")
    print("1. Agent 自主读取原 Word")
    print("2. Agent 自主选择 apply_word_edits")
    print("3. 正文和表格中的年度全部替换")
    print("4. 指定段落删除")
    print("5. 新段落追加")
    print("6. 新 Word 文件成功生成")
    print("7. 原 Word 文件保持不变")
    print("8. Agent 根据 Observation 完成任务")
    print("=" * 70)


if __name__ == "__main__":
    main()