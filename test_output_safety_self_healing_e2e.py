from __future__ import annotations

import json
import tempfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

import agent as agent_module
from agent import DataPilotAgent
from task_planner import TaskPlan


class FakeTaskPlanner:
    """不调用外部 LLM，专门建立 output_safety 自恢复验收合同。"""

    def __init__(self, *args, **kwargs):
        pass

    def create_plan(self, task, context=None):
        return TaskPlan(
            task_goal=(
                "读取真实销售 Excel，按城市汇总销售额，生成新的最终 Excel；"
                "如果首次输出路径试图覆盖受保护源文件，必须依据 RecoveryHint "
                "改用 Workspace deliverables_dir 后继续。"
            ),
            evidence_requirements=[
                "城市销售汇总必须来自真实源 Excel。",
            ],
            source_requirements=[
                "必须读取用户提供的真实 Excel。",
            ],
            deliverable_requirements=[
                "生成新的最终 Excel 交付物。",
            ],
            execution_requirements=[
                "读取真实源 Excel。",
                "按城市统计销售额合计。",
                "输出安全失败后必须修正输出路径并继续，不得覆盖源文件。",
            ],
            verification_requirements=[
                "最终 Excel 必须在最后一次成功写入后重新读取检查。",
                "核对澳门、横琴、珠海的销售额合计。",
            ],
            safety_requirements=[
                "源 Excel 属于 protected input，绝对不得覆盖。",
                "最终交付物必须位于 deliverables_dir。",
            ],
            assumptions=[],
        )


class OutputSafetyCompletions:
    """
    第 3 轮故意把 create_professional_excel_report.output_path
    指向 protected source。

    下一轮只有在真实 Agent state 中确认：
    - 写入失败；
    - recovery.category == output_safety；
    - recoverable == True；
    - protected source 仍在 Workspace；
    - deliverables_dir 可用；
    才会改用 deliverables_dir 重新生成。
    """

    def __init__(self, source_path: Path):
        self.source_path = str(source_path.resolve())
        self.calls = []
        self.decision_index = 0
        self.safe_output_path = None
        self.output_safety_seen = False
        self.protected_path_seen = False
        self.deliverables_dir_seen = False

    @staticmethod
    def _extract_user_content(messages):
        return next(
            (
                item.get("content", "")
                for item in messages
                if item.get("role") == "user"
            ),
            "",
        )

    @classmethod
    def _extract_state(cls, messages):
        content = cls._extract_user_content(messages)
        marker = "当前真实执行状态："
        if marker not in content:
            raise AssertionError("没有找到 AgentLoop 当前真实执行状态。")

        state_text = content.split(marker, 1)[1].strip()
        decoder = json.JSONDecoder()
        state, _ = decoder.raw_decode(state_text)
        return state

    @staticmethod
    def _workspace(state):
        runtime = state.get("runtime_context", {})
        workspace = runtime.get("workspace", {})
        if not isinstance(workspace, dict):
            raise AssertionError(f"Workspace 状态异常：{workspace}")
        return workspace

    def create(self, **kwargs):
        self.calls.append(kwargs)
        messages = kwargs.get("messages", [])
        state = self._extract_state(messages)
        workspace = self._workspace(state)

        if self.safe_output_path is None:
            deliverables_dir = workspace.get("deliverables_dir")
            if not deliverables_dir:
                raise AssertionError(
                    f"Workspace 缺少 deliverables_dir：{workspace}"
                )
            self.safe_output_path = str(
                (
                    Path(deliverables_dir)
                    / "城市销售安全恢复报告.xlsx"
                ).resolve()
            )

        index = self.decision_index

        if index == 0:
            decision = {
                "action_type": "tool",
                "tool": "read_office_data",
                "arguments": {
                    "file_path": self.source_path,
                },
                "purpose": "读取真实销售源数据。",
            }

        elif index == 1:
            completed = state.get("completed_tool_steps", [])
            columns = (
                completed[0]
                .get("observation", {})
                .get("columns", [])
                if completed
                else []
            )
            if "城市" not in columns or "销售额" not in columns:
                raise AssertionError(f"真实源列名异常：{columns}")

            decision = {
                "action_type": "tool",
                "tool": "group_statistics",
                "arguments": {
                    "df": {"$ref": "step_1.output"},
                    "group_by": "城市",
                    "target_column": "销售额",
                    "operation": "sum",
                },
                "purpose": "按城市汇总真实销售额。",
            }

        elif index == 2:
            # 故意违反 Workspace Artifact Governance：
            # 最终报告输出路径直接指向 protected source。
            decision = {
                "action_type": "tool",
                "tool": "create_professional_excel_report",
                "arguments": {
                    "output_path": self.source_path,
                    "dataframe": {"$ref": "step_2.output"},
                    "default_sheet_name": "城市销售汇总",
                    "report_title": "城市销售安全恢复报告",
                    "subtitle": "DataPilot Output Safety Self-Healing E2E",
                    "kpis": [
                        {"label": "销售冠军", "value": "澳门"},
                        {"label": "最高销售额", "value": 186},
                        {"label": "销售总额", "value": 468},
                    ],
                    "charts": [],
                },
                "purpose": (
                    "故意把输出路径指向源 Excel，"
                    "验证 ToolPreflight + output_safety Recovery。"
                ),
            }

        elif index == 3:
            completed = state.get("completed_tool_steps", [])
            if len(completed) < 3:
                raise AssertionError(
                    f"受保护路径失败没有进入 Agent state：{completed}"
                )

            failed_step = completed[2]
            if failed_step.get("success") is not False:
                raise AssertionError(
                    f"step_3 本应被 ToolPreflight 拒绝：{failed_step}"
                )

            observation = failed_step.get("observation", {})
            recovery = observation.get("recovery", {})

            if recovery.get("category") != "output_safety":
                raise AssertionError(
                    f"没有收到 output_safety RecoveryHint：{observation}"
                )
            if recovery.get("recoverable") is not True:
                raise AssertionError(
                    f"output_safety 应允许安全改路径后恢复：{recovery}"
                )
            if not recovery.get("recommended_actions"):
                raise AssertionError(
                    f"RecoveryHint 缺少 recommended_actions：{recovery}"
                )
            if not recovery.get("avoid_actions"):
                raise AssertionError(
                    f"RecoveryHint 缺少 avoid_actions：{recovery}"
                )

            protected = [
                str(Path(item).resolve())
                for item in workspace.get("protected_input_paths", [])
            ]
            if self.source_path not in protected:
                raise AssertionError(
                    f"真实 protected_input_paths 丢失：{protected}"
                )

            deliverables_dir = workspace.get("deliverables_dir")
            if not deliverables_dir:
                raise AssertionError(
                    f"Recovery 时没有 deliverables_dir：{workspace}"
                )

            failed_counts = state.get("failed_tool_counts", {})
            if failed_counts.get("create_professional_excel_report") != 1:
                raise AssertionError(
                    f"failed_tool_counts 异常：{failed_counts}"
                )

            self.output_safety_seen = True
            self.protected_path_seen = True
            self.deliverables_dir_seen = True

            decision = {
                "action_type": "tool",
                "tool": "create_professional_excel_report",
                "arguments": {
                    "output_path": self.safe_output_path,
                    "dataframe": {"$ref": "step_2.output"},
                    "default_sheet_name": "城市销售汇总",
                    "report_title": "城市销售安全恢复报告",
                    "subtitle": "DataPilot Output Safety Self-Healing E2E",
                    "kpis": [
                        {"label": "销售冠军", "value": "澳门"},
                        {"label": "最高销售额", "value": 186},
                        {"label": "销售总额", "value": 468},
                    ],
                    "charts": [],
                },
                "purpose": (
                    "依据 output_safety RecoveryHint，"
                    "保持源文件不变并改用 deliverables_dir。"
                ),
            }

        elif index == 4:
            decision = {
                "action_type": "tool",
                "tool": "inspect_professional_excel_report",
                "arguments": {
                    "file_path": self.safe_output_path,
                },
                "purpose": "重新读取安全路径中的最终 Excel。",
            }

        else:
            decision = {
                "action_type": "finish",
                "final_answer": (
                    "首次输出因试图覆盖受保护源文件被拒绝；"
                    "已依据 output_safety RecoveryHint 改用 deliverables_dir，"
                    "生成并重新读取最终 Excel。"
                ),
            }

        self.decision_index += 1

        return SimpleNamespace(
            choices=[
                SimpleNamespace(
                    message=SimpleNamespace(
                        content=json.dumps(decision, ensure_ascii=False)
                    )
                )
            ]
        )


class FakeClient:
    def __init__(self, source_path: Path):
        self.chat = SimpleNamespace(
            completions=OutputSafetyCompletions(source_path)
        )


def assert_true(value, message):
    if not value:
        raise AssertionError(message)


def main():
    print("=" * 72)
    print("DataPilot v5.0 Output Safety Self-Healing E2E 测试")
    print("=" * 72)

    with tempfile.TemporaryDirectory(
        prefix="datapilot_output_safety_e2e_"
    ) as temp_dir:
        root = Path(temp_dir)
        source_path = root / "销售数据.xlsx"
        output_root = root / "outputs"

        source_df = pd.DataFrame(
            {
                "城市": ["珠海", "澳门", "横琴", "澳门"],
                "销售额": [128, 100, 154, 86],
            }
        )
        source_df.to_excel(source_path, index=False)
        original_bytes = source_path.read_bytes()

        fake_client = FakeClient(source_path)

        datapilot = DataPilotAgent.__new__(DataPilotAgent)
        datapilot.api_key = "test-key"
        datapilot.base_url = "https://example.invalid"
        datapilot.model = "test-model"
        datapilot.progress_callback = None
        datapilot.client = fake_client

        user_task = (
            "读取销售数据，按城市汇总销售额并生成新的专业 Excel。"
            "绝对不要覆盖原文件。如果输出路径被安全规则拒绝，"
            "必须根据真实 RecoveryHint 改用 Workspace 安全交付目录，"
            "生成后重新读取最终 Excel 再完成任务。"
        )

        print("\n测试 1：真实 agent.py 建立 output_safety 自恢复 TaskPlan")
        with patch.object(agent_module, "TaskPlanner", FakeTaskPlanner):
            result = datapilot.execute_v31_agent_task(
                user_task=user_task,
                input_paths=[str(source_path)],
                output_dir=str(output_root),
                max_iterations=8,
            )

        task_plan = result.get("task_plan", {})
        assert_true(isinstance(task_plan, dict), "没有返回 TaskPlan。")
        assert_true(
            any(
                "不得覆盖" in str(item)
                or "protected" in str(item).lower()
                for item in task_plan.get("safety_requirements", [])
            ),
            f"TaskPlan 缺少源文件保护要求：{task_plan}",
        )
        print("PASS")

        history = result.get("tool_results", [])
        assert_true(isinstance(history, list), history)

        create_results = [
            item
            for item in history
            if isinstance(item, dict)
            and item.get("tool_name")
            == "create_professional_excel_report"
        ]

        print("\n测试 2：第一次最终 Excel 写入真实被 ToolPreflight 拒绝")
        assert_true(
            len(create_results) == 2,
            f"专业 Excel 创建调用次数不是 2：{create_results}",
        )
        assert_true(
            create_results[0].get("success") is False,
            f"第一次危险写入没有失败：{create_results[0]}",
        )
        error_text = str(create_results[0].get("error_message", ""))
        assert_true(
            "受保护输入文件" in error_text,
            f"不是 protected input 拒绝：{error_text}",
        )
        print("PASS")

        completions = fake_client.chat.completions

        print("\n测试 3：Agent state 收到 output_safety RecoveryHint")
        assert_true(
            completions.output_safety_seen,
            "决策层没有看到 output_safety RecoveryHint。",
        )
        print("PASS")

        print("\n测试 4：Recovery 时 protected_input_paths 仍然存在")
        assert_true(
            completions.protected_path_seen,
            "Recovery 时丢失 protected_input_paths。",
        )
        print("PASS")

        print("\n测试 5：Recovery 时真实 deliverables_dir 可用")
        assert_true(
            completions.deliverables_dir_seen,
            "Recovery 时没有真实 deliverables_dir。",
        )
        print("PASS")

        print("\n测试 6：第二次写入改用安全交付路径并成功")
        assert_true(
            create_results[1].get("success") is True,
            f"安全路径写入没有成功：{create_results[1]}",
        )
        safe_path = Path(completions.safe_output_path)
        assert_true(safe_path.exists(), f"安全交付物不存在：{safe_path}")
        assert_true(
            safe_path.resolve() != source_path.resolve(),
            "第二次仍然试图覆盖源文件。",
        )
        print("PASS")

        print("\n测试 7：最终交付物位于 Workspace deliverables_dir")
        workspace = result.get("workspace", {})
        deliverables_dir = Path(workspace["deliverables_dir"]).resolve()
        assert_true(
            safe_path.resolve().parent == deliverables_dir,
            f"最终文件不在 deliverables_dir：{safe_path}",
        )
        print("PASS")

        print("\n测试 8：最终 Excel 业务数据正确")
        workbook = load_workbook(safe_path, data_only=False)
        worksheet = workbook["城市销售汇总"]

        header_row = None
        for row_index in range(1, worksheet.max_row + 1):
            values = [
                worksheet.cell(row_index, col).value
                for col in range(
                    1,
                    min(worksheet.max_column, 4) + 1,
                )
            ]
            if "城市" in values and "销售额_合计" in values:
                header_row = row_index
                break

        assert_true(header_row is not None, "最终 Excel 没找到业务表头。")

        expected = {"澳门": 186, "横琴": 154, "珠海": 128}
        actual = {}
        for row_index in range(header_row + 1, worksheet.max_row + 1):
            city = worksheet.cell(row_index, 1).value
            amount = worksheet.cell(row_index, 2).value
            if city in expected:
                actual[str(city)] = amount

        assert_true(actual == expected, f"最终 Excel 数据错误：{actual}")
        print("PASS")

        print("\n测试 9：最终 Excel 在安全写入后被重新读取")
        tool_names = [
            item.get("tool_name")
            for item in history
            if isinstance(item, dict)
        ]
        successful_create_index = max(
            index
            for index, item in enumerate(history)
            if isinstance(item, dict)
            and item.get("tool_name")
            == "create_professional_excel_report"
            and item.get("success") is True
        )
        inspect_index = next(
            index
            for index, item in enumerate(history)
            if isinstance(item, dict)
            and item.get("tool_name")
            == "inspect_professional_excel_report"
            and item.get("success") is True
        )
        assert_true(
            inspect_index > successful_create_index,
            f"最终 inspect 没有发生在安全写入之后：{tool_names}",
        )
        print("PASS")

        print("\n测试 10：Completion Gate 最终 PASS")
        verification = result.get("verification_report", {})
        assert_true(
            verification.get("verified") is True,
            f"Completion Gate 未 PASS：{verification}",
        )
        assert_true(result.get("success") is True, result)
        assert_true(
            result.get("stop_reason") == "completed",
            result.get("stop_reason"),
        )
        print("PASS")

        print("\n测试 11：早期 output_safety 失败被后续成功恢复")
        checks = verification.get("checks", [])
        tool_failure_check = next(
            (
                item
                for item in checks
                if item.get("check_id") == "tool_failures_resolved"
            ),
            None,
        )
        assert_true(
            tool_failure_check is not None,
            f"缺少 tool_failures_resolved：{checks}",
        )
        assert_true(
            tool_failure_check.get("passed") is True,
            tool_failure_check,
        )
        print("PASS")

        print("\n测试 12：源 Excel 字节级保持不变")
        assert_true(source_path.exists(), "源 Excel 被删除。")
        assert_true(
            source_path.read_bytes() == original_bytes,
            "源 Excel 内容被修改。",
        )
        print("PASS")

        print("\n" + "=" * 72)
        print("Output Safety Self-Healing E2E：12/12 PASS")
        print("=" * 72)


if __name__ == "__main__":
    main()
