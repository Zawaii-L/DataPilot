from __future__ import annotations

import os
import shutil
from pathlib import Path
from typing import Any, Dict, List

from docx import Document
from openpyxl import Workbook, load_workbook

from core.agent import DataPilotAgent
from document_tools import read_document


TEST_ROOT = Path("outputs") / "cross_file_natural_task_test"
SOURCE_EXCEL = TEST_ROOT / "销售数据.xlsx"
SOURCE_WORD = TEST_ROOT / "销售月报.docx"


def reset_test_directory() -> None:
    if TEST_ROOT.exists():
        shutil.rmtree(TEST_ROOT)

    TEST_ROOT.mkdir(parents=True, exist_ok=True)


def create_test_excel() -> None:
    workbook = Workbook()

    sales_sheet = workbook.active
    sales_sheet.title = "销售数据"

    sales_sheet.append(["城市", "月份", "销售额", "订单数"])
    sales_sheet.append(["珠海", "2026-09", 128, 16])
    sales_sheet.append(["澳门", "2026-09", 186, 21])
    sales_sheet.append(["横琴", "2026-09", 154, 18])

    note_sheet = workbook.create_sheet("说明")
    note_sheet["A1"] = "单位：万元"
    note_sheet["A2"] = "本工作簿为 DataPilot 跨文件弱提示测试数据。"

    workbook.save(SOURCE_EXCEL)


def create_test_word() -> None:
    document = Document()

    document.add_heading("2026年8月销售月报", level=1)
    document.add_paragraph("本报告汇总2026年8月三个城市的销售情况。")
    document.add_paragraph("2026年8月总销售额为390万元。")
    document.add_paragraph("销售额最高的城市是珠海，销售额为150万元。")
    document.add_paragraph("城市销售额排名：珠海第一、澳门第二、横琴第三。")
    document.add_paragraph("报告状态：待更新。")

    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "项目"
    table.rows[0].cells[1].text = "内容"

    row = table.add_row().cells
    row[0].text = "报告月份"
    row[1].text = "2026年8月"

    row = table.add_row().cells
    row[0].text = "总销售额"
    row[1].text = "390万元"

    row = table.add_row().cells
    row[0].text = "销售冠军"
    row[1].text = "珠海"

    document.save(SOURCE_WORD)


def normalize_path(value: Any) -> str:
    return os.path.normcase(os.path.abspath(os.fspath(value)))


def collect_generated_word_files() -> List[Path]:
    source_normalized = normalize_path(SOURCE_WORD)
    candidates: List[Path] = []

    for path in TEST_ROOT.rglob("*.docx"):
        if normalize_path(path) == source_normalized:
            continue
        candidates.append(path)

    return sorted(
        candidates,
        key=lambda item: item.stat().st_mtime,
        reverse=True,
    )


def get_tool_results(result: Dict[str, Any]) -> List[Dict[str, Any]]:
    raw_results = result.get("tool_results", []) or []
    normalized: List[Dict[str, Any]] = []

    for item in raw_results:
        if isinstance(item, dict):
            normalized.append(item)
        elif hasattr(item, "to_dict"):
            normalized.append(item.to_dict())

    return normalized


def get_tool_name(item: Dict[str, Any]) -> str:
    return str(
        item.get("tool_name")
        or item.get("tool")
        or ""
    ).strip()


def get_tool_success(item: Dict[str, Any]) -> bool:
    return bool(item.get("success"))


def print_sources() -> None:
    workbook = load_workbook(SOURCE_EXCEL, data_only=False)
    sheet = workbook["销售数据"]

    print("Excel 原始数据：")
    print("-" * 70)

    for row in sheet.iter_rows(values_only=True):
        print(row)

    print("-" * 70)
    print()

    print("Word 原始内容：")
    print("-" * 70)
    print(read_document(str(SOURCE_WORD)))
    print("-" * 70)
    print()


def assert_source_files_unchanged() -> None:
    workbook = load_workbook(SOURCE_EXCEL, data_only=False)
    sheet = workbook["销售数据"]

    rows = list(sheet.iter_rows(values_only=True))

    expected_rows = [
        ("城市", "月份", "销售额", "订单数"),
        ("珠海", "2026-09", 128, 16),
        ("澳门", "2026-09", 186, 21),
        ("横琴", "2026-09", 154, 18),
    ]

    if rows != expected_rows:
        raise AssertionError("源 Excel 被修改。")

    source_word_text = read_document(str(SOURCE_WORD))

    required_old_values = [
        "2026年8月销售月报",
        "2026年8月总销售额为390万元。",
        "销售额最高的城市是珠海，销售额为150万元。",
        "城市销售额排名：珠海第一、澳门第二、横琴第三。",
        "报告状态：待更新。",
        "390万元",
    ]

    for text in required_old_values:
        if text not in source_word_text:
            raise AssertionError(
                f"源 Word 被修改或原始内容丢失：{text}"
            )


def assert_generated_word(output_path: Path) -> str:
    output_text = read_document(str(output_path))

    print()
    print("Agent 生成的新 Word：")
    print(output_path.resolve())
    print()
    print("修改后 Word 内容：")
    print("-" * 70)
    print(output_text)
    print("-" * 70)

    required_values = [
        "2026年9月销售月报",
        "2026年9月",
        "468万元",
        "澳门",
        "186万元",
        "澳门第一",
        "横琴第二",
        "珠海第三",
    ]

    for text in required_values:
        if text not in output_text:
            raise AssertionError(
                f"新 Word 缺少应从 Excel 推导并同步的数据：{text}"
            )

    forbidden_values = [
        "2026年8月销售月报",
        "2026年8月总销售额为390万元。",
        "销售额最高的城市是珠海，销售额为150万元。",
        "珠海第一、澳门第二、横琴第三",
    ]

    for text in forbidden_values:
        if text in output_text:
            raise AssertionError(
                f"新 Word 仍残留旧报告内容：{text}"
            )

    return output_text


def assert_agent_chain(result: Dict[str, Any]) -> None:
    tool_results = get_tool_results(result)

    if not tool_results:
        raise AssertionError("Agent 没有产生任何工具调用记录。")

    tool_names = [
        get_tool_name(item)
        for item in tool_results
    ]

    print()
    print("工具调用链：")

    for index, item in enumerate(tool_results, start=1):
        print(
            f"{index}. "
            f"{get_tool_name(item)} "
            f"SUCCESS={get_tool_success(item)}"
        )

    if "apply_word_edits" not in tool_names:
        raise AssertionError(
            "Agent 没有自主判断需要修改 Word。"
        )

    edit_index = tool_names.index("apply_word_edits")
    before_edit = tool_names[:edit_index]
    after_edit = tool_names[edit_index + 1:]

    excel_read_tools = {
        "inspect_data_files",
        "read_office_data",
    }

    if not any(
        name in excel_read_tools
        for name in before_edit
    ):
        raise AssertionError(
            "Agent 在修改 Word 前没有自主读取或检查 Excel。"
        )

    if "read_document" not in before_edit:
        raise AssertionError(
            "Agent 在修改 Word 前没有自主读取原 Word。"
        )

    if "read_document" not in after_edit:
        raise AssertionError(
            "Agent 修改 Word 后没有自主回读新 Word 进行检查。"
        )


def main() -> None:
    print("=" * 70)
    print("DataPilot v3.3 跨文件自然语言弱提示测试")
    print("=" * 70)
    print()

    reset_test_directory()
    create_test_excel()
    create_test_word()

    print("测试 Excel：", SOURCE_EXCEL.resolve())
    print("测试 Word：", SOURCE_WORD.resolve())
    print()

    print_sources()

    task = (
        "根据我提供的 Excel 最新销售数据更新 Word 销售月报，"
        "把里面过时的数据同步成最新结果。"
        "不要覆盖原文件，完成后帮我检查一下修改是否正确。"
    )

    print("用户自然语言任务：")
    print(task)
    print()

    print(
        "注意：这次没有告诉 Agent 正确月份、总额、冠军、"
        "排名，也没有指定具体工具调用顺序。"
    )
    print()

    agent = DataPilotAgent()

    print("=" * 70)
    print("开始运行真实 DataPilot Agent Loop")
    print("=" * 70)

    result = agent.execute_v31_agent_task(
        user_task=task,
        input_paths=[
            str(SOURCE_EXCEL.resolve()),
            str(SOURCE_WORD.resolve()),
        ],
        output_dir=str(TEST_ROOT.resolve()),
        max_iterations=12,
    )

    print()
    print("=" * 70)
    print("Agent Loop 返回结果")
    print("=" * 70)

    print("SUCCESS =", result.get("success"))
    print("STOP_REASON =", result.get("stop_reason"))
    print("ITERATIONS =", result.get("iterations"))
    print("TOOL_COUNT =", result.get("tool_count"))

    final_answer = (
        result.get("final_answer")
        or result.get("answer")
        or ""
    )

    print()
    print("Agent 最终回答：")
    print(final_answer)

    if not result.get("success"):
        raise AssertionError(
            "弱提示条件下 Agent Loop 没有成功完成跨文件任务。"
        )

    if result.get("stop_reason") != "completed":
        raise AssertionError(
            "Agent 没有以 completed 正常结束。"
        )

    candidates = collect_generated_word_files()

    if not candidates:
        raise AssertionError(
            "弱提示条件下没有检测到 Agent 生成的新 Word。"
        )

    output_word = candidates[0]

    assert_generated_word(output_word)

    print()
    print("开始验证源文件保护……")
    assert_source_files_unchanged()
    print("源 Excel 和源 Word 均保持不变。")

    print()
    print("开始验证自主跨文件工具链……")
    assert_agent_chain(result)

    print()
    print("=" * 70)
    print("DataPilot v3.3 跨文件自然语言弱提示测试通过！")
    print("=" * 70)
    print()
    print("已验证：")
    print("1. 用户只给出简短自然语言办公目标")
    print("2. Agent 自主识别 Excel 是最新数据来源")
    print("3. Agent 自主读取 Excel")
    print("4. Agent 自主读取 Word")
    print("5. Agent 自主推导报告期为 2026年9月")
    print("6. Agent 自主计算总销售额为 468 万元")
    print("7. Agent 自主判断销售冠军为澳门 186 万元")
    print("8. Agent 自主判断排名为澳门、横琴、珠海")
    print("9. Agent 自主选择 Word 编辑工具")
    print("10. Agent 生成新 Word，不覆盖源文件")
    print("11. Agent 编辑后自主回读新 Word 检查")
    print("12. 最终结果通过确定性数据验证")
    print("=" * 70)


if __name__ == "__main__":
    main()
