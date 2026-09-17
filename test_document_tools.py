from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

from document_tools import (
    build_document_catalog_text,
    get_document_preview,
    inspect_document,
    inspect_documents,
    read_document,
    read_text_file,
    read_word_file,
    scan_document_files,
)


def create_test_txt(folder: Path) -> Path:
    path = folder / "项目说明.txt"

    path.write_text(
        (
            "DataPilot 项目说明\n"
            "本项目用于自动处理办公文件。\n"
            "项目负责人：测试用户。\n"
            "计划完成时间：2026年9月。\n"
        ),
        encoding="utf-8",
    )

    return path


def create_test_markdown(folder: Path) -> Path:
    path = folder / "会议纪要.md"

    path.write_text(
        (
            "# 项目会议纪要\n\n"
            "## 本周任务\n"
            "- 完成 Word 文档读取\n"
            "- 完成 PDF 文档读取\n"
            "- 接入 DeepSeek 文档选择\n\n"
            "## 项目状态\n"
            "DataPilot v3.0 正在开发。\n"
        ),
        encoding="utf-8",
    )

    return path


def create_test_word(folder: Path) -> Path:
    path = folder / "销售报告.docx"

    document = Document()

    document.add_heading(
        "2026年8月销售报告",
        level=1,
    )

    document.add_paragraph(
        "2026年8月珠海销售额为120万元。"
    )

    document.add_paragraph(
        "2026年8月澳门销售额为150万元。"
    )

    document.add_paragraph(
        "本月总销售额较上月增长12%。"
    )

    table = document.add_table(
        rows=1,
        cols=3,
    )

    header = table.rows[0].cells

    header[0].text = "城市"
    header[1].text = "销售额"
    header[2].text = "同比增长"

    data = [
        ("珠海", "120万元", "10%"),
        ("澳门", "150万元", "14%"),
    ]

    for city, sales, growth in data:
        cells = table.add_row().cells

        cells[0].text = city
        cells[1].text = sales
        cells[2].text = growth

    document.save(
        str(path)
    )

    return path


def create_irrelevant_word(folder: Path) -> Path:
    path = folder / "员工通知.docx"

    document = Document()

    document.add_heading(
        "员工通知",
        level=1,
    )

    document.add_paragraph(
        "公司将于本周五下午进行办公室设备维护。"
    )

    document.add_paragraph(
        "请员工提前保存工作文件。"
    )

    document.save(
        str(path)
    )

    return path


def create_temp_word_file(folder: Path) -> Path:
    path = folder / "~$临时文件.docx"

    path.write_bytes(
        b"temporary"
    )

    return path


def test_text_reading(txt_path: Path):
    print("\n[测试 1] TXT 文档读取")

    text = read_text_file(txt_path)

    assert "DataPilot" in text
    assert "2026年9月" in text

    print("TXT 读取通过。")


def test_markdown_reading(md_path: Path):
    print("\n[测试 2] Markdown 文档读取")

    text = read_document(md_path)

    assert "项目会议纪要" in text
    assert "DeepSeek 文档选择" in text
    assert "DataPilot v3.0" in text

    print("Markdown 读取通过。")


def test_word_reading(word_path: Path):
    print("\n[测试 3] Word 正文 + 表格读取")

    text = read_word_file(word_path)

    assert "2026年8月销售报告" in text
    assert "珠海销售额为120万元" in text
    assert "澳门销售额为150万元" in text

    assert "[表格 1]" in text
    assert "城市 | 销售额 | 同比增长" in text
    assert "珠海 | 120万元 | 10%" in text
    assert "澳门 | 150万元 | 14%" in text

    print("Word 正文和表格读取通过。")


def test_document_preview(word_path: Path):
    print("\n[测试 4] 文档预览截断")

    preview = get_document_preview(
        word_path,
        max_characters=50,
    )

    assert len(preview) > 0
    assert "文档内容已截断" in preview

    print("文档预览截断通过。")


def test_document_inspection(word_path: Path):
    print("\n[测试 5] 单文档画像")

    info = inspect_document(
        word_path,
        preview_characters=100,
    )

    assert info["inspection_success"] is True
    assert info["file_name"] == "销售报告.docx"
    assert info["extension"] == ".docx"
    assert info["character_count"] > 0
    assert "2026年8月销售报告" in info["preview"]

    print("单文档画像通过。")


def test_folder_scan(
    folder: Path,
    expected_names,
):
    print("\n[测试 6] 文件夹文档扫描")

    files = scan_document_files(
        folder,
        recursive=False,
    )

    names = {
        Path(item).name
        for item in files
    }

    for expected_name in expected_names:
        assert expected_name in names

    assert "~$临时文件.docx" not in names

    print(
        f"识别到 {len(files)} 个有效办公文档。"
    )

    for name in sorted(names):
        print(f"  - {name}")

    print("Office 临时文件忽略测试通过。")

    return files


def test_batch_inspection(files):
    print("\n[测试 7] 批量文档画像")

    infos = inspect_documents(
        files,
        preview_characters=150,
    )

    assert len(infos) == len(files)

    success_infos = [
        info
        for info in infos
        if info["inspection_success"]
    ]

    assert len(success_infos) == len(files)

    print(
        f"{len(infos)} 个文档全部读取成功。"
    )

    return infos


def test_catalog(infos):
    print("\n[测试 8] LLM 文档目录生成")

    catalog = build_document_catalog_text(
        infos,
        include_preview=True,
    )

    assert "[文档 1]" in catalog
    assert "文件名：" in catalog
    assert "路径：" in catalog
    assert "内容预览：" in catalog

    assert "销售报告.docx" in catalog
    assert "员工通知.docx" in catalog
    assert "会议纪要.md" in catalog

    print("LLM 文档目录生成通过。")

    print("\n生成的文档目录预览：")
    print("-" * 70)

    print(catalog[:1500])

    if len(catalog) > 1500:
        print("\n[目录显示已截断]")

    print("-" * 70)


def main():
    print("=" * 70)
    print("DataPilot v3.0 document_tools 自动化测试")
    print("=" * 70)

    with TemporaryDirectory() as temp_dir:
        folder = Path(temp_dir)

        print(
            f"\n临时测试目录：{folder}"
        )

        txt_path = create_test_txt(folder)
        md_path = create_test_markdown(folder)
        word_path = create_test_word(folder)

        create_irrelevant_word(folder)
        create_temp_word_file(folder)

        test_text_reading(txt_path)
        test_markdown_reading(md_path)
        test_word_reading(word_path)
        test_document_preview(word_path)
        test_document_inspection(word_path)

        files = test_folder_scan(
            folder=folder,
            expected_names={
                "项目说明.txt",
                "会议纪要.md",
                "销售报告.docx",
                "员工通知.docx",
            },
        )

        infos = test_batch_inspection(files)

        test_catalog(infos)

    print("\n" + "=" * 70)
    print("全部测试通过。")

    print(
        "v3.0 多格式办公文档读取底层已具备："
    )
    print("TXT / Markdown / Word 读取")
    print("Word 正文 + 表格提取")
    print("文档预览")
    print("文件夹扫描")
    print("Office 临时文件忽略")
    print("文档画像")
    print("LLM 文档目录生成")

    print("=" * 70)


if __name__ == "__main__":
    main()