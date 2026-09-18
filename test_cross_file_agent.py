from __future__ import annotations

import os
import shutil
from pathlib import Path
from typing import Any, Dict, List

from docx import Document
from openpyxl import Workbook, load_workbook

from agent import DataPilotAgent
from document_tools import read_document


TEST_ROOT = Path("outputs") / "cross_file_agent_test"
SOURCE_EXCEL = TEST_ROOT / "销售数据.xlsx"
SOURCE_WORD = TEST_ROOT / "销售月报.docx"


def reset_test_directory() -> None:
    if TEST_ROOT.exists():
        shutil.rmtree(TEST_ROOT)

    TEST_ROOT.mkdir(parents=True, exist_ok=True)


def create_test_excel() -> None:
    workbook = Workbook()

    sales_sheet = workbook.active
    sales_sheet.title = "销售数据"

    sales_sheet.append(
        ["城市", "月份", "销售额", "订单数"]
    )
    sales_sheet.append(
        ["珠海", "2026-09", 128, 16]
    )
    sales_sheet.append(
        ["澳门", "2026-09", 186, 21]
    )
    sales_sheet.append(
        ["横琴", "2026-09", 154, 18]
    )

    note_sheet = workbook.create_sheet("说明")
    note_sheet["A1"] = "单位：万元"
    note_sheet["A2"] = "本工作簿为 DataPilot v3.3 跨文件测试数据。"

    workbook.save(SOURCE_EXCEL)


def create_test_word() -> None:
    document = Document()

    document.add_heading(
        "2026年8月销售月报",
        level=1,
    )

    document.add_paragraph(
        "本报告汇总2026年8月三个城市的销售情况。"
    )

    document.add_paragraph(
        "2026年8月总销售额为390万元。"
    )

    document.add_paragraph(
        "销售额最高的城市是珠海，销售额为150万元。"
    )

    document.add_paragraph(
        "城市销售额排名：珠海第一、澳门第二、横琴第三。"
    )

    document.add_paragraph(
        "报告状态：待更新。"
    )

    table = document.add_table(
        rows=1,
        cols=2,
    )

    table.rows[0].cells[0].text = "项目"
    table.rows[0].cells[1].text = "内容"

    row = table.add_row().cells
    row[0].text = "报告月份"
    row[1].text = "2026年8月"

    row = table.add_row().cells
    row[0].text = "总销售额"
    row[1].text = "390万元"

    row = table.add_row().cells
    row[0].text = "销售冠军"
    row[1].text = "珠海"

    document.save(SOURCE_WORD)


def print_excel_source() -> None:
    workbook = load_workbook(
        SOURCE_EXCEL,
        data_only=False,
    )

    sheet = workbook["销售数据"]

    print("Excel 原始数据：")
    print("-" * 70)

    for row in sheet.iter_rows(
        values_only=True
    ):
        print(row)

    print("-" * 70)
    print()


def print_word_source() -> None:
    print("Word 原始内容：")
    print("-" * 70)
    print(read_document(str(SOURCE_WORD)))
    print("-" * 70)
    print()


def normalize_path(value: Any) -> str:
    return os.path.normcase(
        os.path.abspath(
            os.fspath(value)
        )
    )


def collect_generated_word_files() -> List[Path]:
    source_normalized = normalize_path(
        SOURCE_WORD
    )

    candidates = []

    for path in TEST_ROOT.rglob("*.docx"):
        if normalize_path(path) == source_normalized:
            continue

        candidates.append(path)

    return sorted(
        candidates,
        key=lambda item: item.stat().st_mtime,
        reverse=True,
    )


def get_tool_results(
    result: Dict[str, Any],
) -> List[Dict[str, Any]]:
    raw_results = result.get(
        "tool_results",
        [],
    ) or []

    normalized = []

    for item in raw_results:
        if isinstance(item, dict):
            normalized.append(item)
            continue

        if hasattr(item, "to_dict"):
            normalized.append(
                item.to_dict()
            )

    return normalized


def get_tool_name(
    item: Dict[str, Any],
) -> str:
    return str(
        item.get("tool_name")
        or item.get("tool")
        or ""
    ).strip()


def get_tool_success(
    item: Dict[str, Any],
) -> bool:
    return bool(
        item.get("success")
    )


def assert_source_files_unchanged() -> None:
    workbook = load_workbook(
        SOURCE_EXCEL,
        data_only=False,
    )

    sheet = workbook["销售数据"]

    rows = list(
        sheet.iter_rows(
            values_only=True
        )
    )

    expected_rows = [
        ("城市", "月份", "销售额", "订单数"),
        ("珠海", "2026-09", 128, 16),
        ("澳门", "2026-09", 186, 21),
        ("横琴", "2026-09", 154, 18),
    ]

    if rows != expected_rows:
        raise AssertionError(
            "源 Excel 被修改。"
        )

    source_word_text = read_document(
        str(SOURCE_WORD)
    )

    required_old_values = [
        "2026年8月销售月报",
        "2026年8月总销售额为390万元。",
        "销售额最高的城市是珠海，销售额为150万元。",
        "城市销售额排名：珠海第一、澳门第二、横琴第三。",
        "报告状态：待更新。",
        "390万元",
    ]

    for text in required_old_values:
        if text not in source_word_text:
            raise AssertionError(
                f"源 Word 被修改或原始内容丢失：{text}"
            )


def assert_generated_word(
    output_path: Path,
) -> str:
    output_text = read_document(
        str(output_path)
    )

    print()
    print("Agent 生成的新 Word：")
    print(output_path.resolve())
    print()
    print("修改后 Word 内容：")
    print("-" * 70)
    print(output_text)
    print("-" * 70)

    required_values = [
        "2026年9月销售月报",
        "2026年9月",
        "468万元",
        "澳门",
        "186万元",
        "澳门第一",
        "横琴第二",
        "珠海第三",
        "报告状态：",
    ]

    for text in required_values:
        if text not in output_text:
            raise AssertionError(
                f"新 Word 缺少应同步的数据：{text}"
            )

    forbidden_values = [
        "2026年8月销售月报",
        "2026年8月总销售额为390万元。",
        "销售额最高的城市是珠海，销售额为150万元。",
        "珠海第一、澳门第二、横琴第三",
        "报告状态：待更新。",
    ]

    for text in forbidden_values:
        if text in output_text:
            raise AssertionError(
                f"新 Word 仍残留旧报告内容：{text}"
            )

    return output_text


def assert_agent_chain(
    result: Dict[str, Any],
) -> None:
    tool_results = get_tool_results(
        result
    )

    if not tool_results:
        raise AssertionError(
            "Agent 没有产生任何工具调用记录。"
        )

    tool_names = [
        get_tool_name(item)
        for item in tool_results
    ]

    print()
    print("工具调用链：")

    for index, item in enumerate(
        tool_results,
        start=1,
    ):
        print(
            f"{index}. "
            f"{get_tool_name(item)} "
            f"SUCCESS={get_tool_success(item)}"
        )

    if "apply_word_edits" not in tool_names:
        raise AssertionError(
            "Agent 没有调用 apply_word_edits。"
        )

    edit_index = tool_names.index(
        "apply_word_edits"
    )

    before_edit = tool_names[:edit_index]
    after_edit = tool_names[
        edit_index + 1:
    ]

    excel_read_tools = {
        "inspect_data_files",
        "read_office_data",
    }

    if not any(
        name in excel_read_tools
        for name in before_edit
    ):
        raise AssertionError(
            "Agent 在修改 Word 前没有先读取或检查 Excel 数据。"
        )

    if "read_document" not in before_edit:
        raise AssertionError(
            "Agent 在修改 Word 前没有读取原 Word。"
        )

    if "read_document" not in after_edit:
        raise AssertionError(
            "Agent 修改 Word 后没有重新读取新 Word 进行核验。"
        )

    if not any(
        name in excel_read_tools
        for name in after_edit
    ):
        raise AssertionError(
            "Agent 修改 Word 后没有再次读取或检查 Excel，"
            "尚未形成明确的跨文件交叉核验闭环。"
        )


def main() -> None:
    print("=" * 70)
    print("DataPilot v3.3 跨文件 Office Agent 基准测试")
    print("=" * 70)
    print()

    reset_test_directory()
    create_test_excel()
    create_test_word()

    print(
        "测试 Excel：",
        SOURCE_EXCEL.resolve(),
    )
    print(
        "测试 Word：",
        SOURCE_WORD.resolve(),
    )
    print()

    print_excel_source()
    print_word_source()

    task = (
        "请根据我提供的 Excel 最新销售数据，更新我提供的 Word 销售月报。"
        "你必须先分别读取并理解 Excel 和 Word，不能根据任务描述猜数据。"
        "Excel 的“销售数据”Sheet 是唯一真实数据来源。"
        "请把 Word 中的报告月份、总销售额、销售冠军、冠军销售额、"
        "三个城市的销售额排名以及报告状态同步为 Excel 中的最新结果。"
        "Excel 中 2026-09 的数据为当前报告期，销售额单位是万元。"
        "总销售额必须根据 Excel 三个城市真实销售额计算，"
        "城市排名必须按销售额从高到低确定。"
        "Word 正文和表格中相关的旧月份、旧数字、旧冠军都要同步更新。"
        "不要修改 Excel，不要覆盖原 Word，必须另存为新的 Word 文件。"
        "修改完成后，必须重新读取新 Word，并再次读取或检查 Excel，"
        "把新 Word 中的月份、总销售额、冠军、冠军销售额和城市排名"
        "与 Excel 的真实数据进行交叉核验。"
        "只有确认两份文件数据一致后才能结束任务。"
    )

    print("用户任务：")
    print(task)
    print()

    agent = DataPilotAgent()

    print("=" * 70)
    print("开始运行真实 DataPilot Agent Loop")
    print("=" * 70)

    result = agent.execute_v31_agent_task(
        user_task=task,
        input_paths=[
            str(SOURCE_EXCEL.resolve()),
            str(SOURCE_WORD.resolve()),
        ],
        output_dir=str(
            TEST_ROOT.resolve()
        ),
        max_iterations=12,
    )

    print()
    print("=" * 70)
    print("Agent Loop 返回结果")
    print("=" * 70)

    print(
        "SUCCESS =",
        result.get("success"),
    )
    print(
        "STOP_REASON =",
        result.get("stop_reason"),
    )
    print(
        "ITERATIONS =",
        result.get("iterations"),
    )
    print(
        "TOOL_COUNT =",
        result.get("tool_count"),
    )

    final_answer = (
        result.get("final_answer")
        or result.get("answer")
        or ""
    )

    print()
    print("Agent 最终回答：")
    print(final_answer)

    if not result.get("success"):
        raise AssertionError(
            "Agent Loop 没有成功完成跨文件任务。"
        )

    if result.get("stop_reason") != "completed":
        raise AssertionError(
            "Agent 没有以 completed 正常结束。"
        )

    candidates = collect_generated_word_files()

    if not candidates:
        raise AssertionError(
            "没有检测到 Agent 生成的新 Word 文件。"
        )

    output_word = candidates[0]

    assert_generated_word(
        output_word
    )

    print()
    print("开始验证源文件保护……")

    assert_source_files_unchanged()

    print("源 Excel 和源 Word 均保持不变。")

    print()
    print("开始验证 Agent 跨文件工具链……")

    assert_agent_chain(
        result
    )

    print()
    print("=" * 70)
    print("DataPilot v3.3 跨文件 Office Agent 基准测试通过！")
    print("=" * 70)
    print()
    print("已验证：")
    print("1. Agent 在编辑前读取 / 检查 Excel")
    print("2. Agent 在编辑前读取原 Word")
    print("3. Excel 是真实业务数据来源")
    print("4. Agent 自主计算并同步总销售额")
    print("5. Agent 自主识别销售冠军及冠军销售额")
    print("6. Agent 自主确定三个城市销售排名")
    print("7. Word 正文和表格被同步更新")
    print("8. Agent 没有修改源 Excel")
    print("9. Agent 没有覆盖源 Word")
    print("10. Agent 编辑后重新读取新 Word")
    print("11. Agent 编辑后再次读取 / 检查 Excel")
    print("12. Agent 完成跨文件交叉核验后才结束")
    print("=" * 70)


if __name__ == "__main__":
    main()
