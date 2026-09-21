from pathlib import Path

import pandas as pd

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


# ============================================================
# Word 基础格式
# ============================================================

def set_run_font(
    run,
    font_name="Microsoft YaHei",
    font_size=10,
    bold=False,
):
    """
    设置 Word 中英文统一字体。
    """
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold

    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts

    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)

    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)
    r_fonts.set(qn("w:eastAsia"), font_name)
    r_fonts.set(qn("w:cs"), font_name)


def set_cell_shading(cell, fill="D9EAF7"):
    """
    设置表格单元格背景色。
    """
    tc_pr = cell._tc.get_or_add_tcPr()

    shd = tc_pr.find(qn("w:shd"))

    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)

    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="B7C9D6", size="4"):
    """
    设置表格边框。
    """
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()

    tc_borders = tc_pr.first_child_found_in("w:tcBorders")

    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)

    for edge in (
        "top",
        "left",
        "bottom",
        "right",
        "insideH",
        "insideV",
    ):
        tag = "w:" + edge
        element = tc_borders.find(qn(tag))

        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)

        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_text(
    cell,
    text,
    font_size=9,
    bold=False,
    alignment=WD_ALIGN_PARAGRAPH.CENTER,
):
    """
    设置单元格文字。
    """
    cell.text = ""

    paragraph = cell.paragraphs[0]
    paragraph.alignment = alignment

    run = paragraph.add_run(str(text))

    set_run_font(
        run,
        font_size=font_size,
        bold=bold,
    )

    cell.vertical_alignment = (
        WD_CELL_VERTICAL_ALIGNMENT.CENTER
    )


def set_table_fixed_layout(table):
    """
    设置表格固定布局。
    """
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    tbl_pr = table._tbl.tblPr
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))

    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)

    tbl_layout.set(qn("w:type"), "fixed")


def add_heading(document, text, level=1):
    """
    添加标题。
    """
    paragraph = document.add_heading(level=level)

    run = paragraph.add_run(str(text))

    font_size = 16 if level == 1 else 13 if level == 2 else 11

    set_run_font(
        run,
        font_size=font_size,
        bold=True,
    )

    return paragraph


def add_body_paragraph(document, text):
    """
    添加正文。
    """
    paragraph = document.add_paragraph()

    run = paragraph.add_run(str(text))

    set_run_font(
        run,
        font_size=10,
    )

    paragraph.paragraph_format.space_after = Pt(5)

    return paragraph


def add_bullet(document, text):
    """
    添加项目符号。
    """
    paragraph = document.add_paragraph()

    paragraph.paragraph_format.left_indent = Inches(0.25)

    run = paragraph.add_run("• " + str(text))

    set_run_font(
        run,
        font_size=10,
    )

    return paragraph


# ============================================================
# 通用数据表
# ============================================================

def add_key_value_table(
    document,
    title,
    data,
):
    """
    将字典生成两列表格。
    """
    add_heading(
        document,
        title,
        level=2,
    )

    if not isinstance(data, dict):
        data = {}

    table = document.add_table(
        rows=1,
        cols=2,
    )

    set_table_fixed_layout(table)

    table.columns[0].width = Inches(2.5)
    table.columns[1].width = Inches(4.3)

    header_cells = table.rows[0].cells

    set_cell_text(
        header_cells[0],
        "项目",
        font_size=9,
        bold=True,
    )

    set_cell_text(
        header_cells[1],
        "结果",
        font_size=9,
        bold=True,
    )

    for cell in header_cells:
        set_cell_shading(cell)
        set_cell_border(cell)

    if not data:
        cells = table.add_row().cells

        set_cell_text(
            cells[0],
            "记录",
            font_size=9,
        )

        set_cell_text(
            cells[1],
            "没有记录",
            font_size=9,
        )

        for cell in cells:
            set_cell_border(cell)

    else:
        for key, value in data.items():
            cells = table.add_row().cells

            if isinstance(value, dict):
                value = str(value)

            if isinstance(value, list):
                value = ", ".join(str(item) for item in value)

            set_cell_text(
                cells[0],
                key,
                font_size=9,
            )

            set_cell_text(
                cells[1],
                value,
                font_size=9,
            )

            for cell in cells:
                set_cell_border(cell)

    document.add_paragraph()


def add_dataframe_table(
    document,
    title,
    dataframe,
):
    """
    将 DataFrame 生成 Word 表格。
    """
    add_heading(
        document,
        title,
        level=2,
    )

    if dataframe is None:
        add_body_paragraph(
            document,
            "没有可用的数据。",
        )
        return

    if not isinstance(dataframe, pd.DataFrame):
        try:
            dataframe = pd.DataFrame(dataframe)
        except Exception:
            add_body_paragraph(
                document,
                "数据格式无法转换为表格。",
            )
            return

    if dataframe.empty:
        add_body_paragraph(
            document,
            "数据表为空。",
        )
        return

    display_df = dataframe.copy()

    # 防止表格过宽，最多展示前 12 列
    if len(display_df.columns) > 12:
        display_df = display_df.iloc[:, :12]

    table = document.add_table(
        rows=1,
        cols=len(display_df.columns),
    )

    set_table_fixed_layout(table)

    total_width = 6.8
    column_width = total_width / len(display_df.columns)

    for column in table.columns:
        column.width = Inches(column_width)

    header_cells = table.rows[0].cells

    for index, column_name in enumerate(display_df.columns):
        set_cell_text(
            header_cells[index],
            column_name,
            font_size=8,
            bold=True,
        )

        set_cell_shading(
            header_cells[index],
            "D9EAF7",
        )

        set_cell_border(header_cells[index])

    for _, row in display_df.iterrows():
        cells = table.add_row().cells

        for index, value in enumerate(row):
            if isinstance(value, float):
                value = round(value, 4)

            set_cell_text(
                cells[index],
                value,
                font_size=8,
            )

            set_cell_border(cells[index])

    document.add_paragraph()


# ============================================================
# 批量 Word 报告主函数
# ============================================================

def generate_batch_word_report(
    task,
    result,
    output_path="outputs/batch_report.docx",
):
    """
    根据批量数据处理结果生成 Word 报告。

    参数：
        task:
            用户任务描述。

        result:
            run_batch_pipeline() 返回的结果字典。

        output_path:
            Word 报告保存路径。
    """

    if not isinstance(result, dict):
        raise TypeError(
            "generate_batch_word_report() 的 result 必须是字典。"
        )

    output_path = Path(output_path)

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    document = Document()

    # ========================================================
    # 页面设置
    # ========================================================

    section = document.sections[0]

    section.orientation = WD_ORIENT.LANDSCAPE

    section.page_width = Inches(11.69)
    section.page_height = Inches(8.27)

    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    # ========================================================
    # 报告标题
    # ========================================================

    title = document.add_heading(level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    title_run = title.add_run(
        "DataPilot 批量数据分析报告"
    )

    set_run_font(
        title_run,
        font_size=20,
        bold=True,
    )

    # ========================================================
    # 一、任务概述
    # ========================================================

    add_heading(
        document,
        "一、任务概述",
        level=1,
    )

    add_body_paragraph(
        document,
        task or "未提供任务描述。",
    )

    # ========================================================
    # 二、输入文件
    # ========================================================

    add_heading(
        document,
        "二、输入文件",
        level=1,
    )

    input_paths = result.get(
        "input_paths",
        result.get("file_paths", []),
    )

    if isinstance(input_paths, str):
        input_paths = [input_paths]

    if input_paths:
        for file_path in input_paths:
            add_bullet(
                document,
                str(file_path),
            )
    else:
        add_body_paragraph(
            document,
            "没有记录输入文件。",
        )

    # ========================================================
    # 三、数据基本情况
    # ========================================================

    add_heading(
        document,
        "三、数据基本情况",
        level=1,
    )

    original_df = result.get("original_df")
    cleaned_df = result.get("cleaned_df")

    if isinstance(original_df, pd.DataFrame):
        add_body_paragraph(
            document,
            (
                f"合并后的原始数据包含 "
                f"{original_df.shape[0]} 行、"
                f"{original_df.shape[1]} 列。"
            ),
        )

    if isinstance(cleaned_df, pd.DataFrame):
        add_body_paragraph(
            document,
            (
                f"清洗后的数据包含 "
                f"{cleaned_df.shape[0]} 行、"
                f"{cleaned_df.shape[1]} 列。"
            ),
        )

    # ========================================================
    # 四、数据质量
    # ========================================================

    add_heading(
        document,
        "四、数据质量检查",
        level=1,
    )

    before_quality = result.get(
        "before_quality",
        result.get("quality_result", {}),
    )

    after_quality = result.get(
        "after_quality",
        {},
    )

    add_key_value_table(
        document,
        "清洗前数据质量",
        before_quality,
    )

    add_key_value_table(
        document,
        "清洗后数据质量",
        after_quality,
    )

    # ========================================================
    # 五、清洗记录
    # ========================================================

    add_heading(
        document,
        "五、数据清洗过程",
        level=1,
    )

    cleaning_log = result.get(
        "cleaning_log",
        result.get("cleaning_result", {}),
    )

    add_key_value_table(
        document,
        "清洗记录",
        cleaning_log,
    )

    # ========================================================
    # 六、统计分析
    # ========================================================

    add_heading(
        document,
        "六、统计分析",
        level=1,
    )

    statistics = result.get(
        "statistics",
        result.get("statistics_result"),
    )

    add_dataframe_table(
        document,
        "统计结果",
        statistics,
    )

    # ========================================================
    # 七、趋势图
    # ========================================================

    add_heading(
        document,
        "七、趋势图",
        level=1,
    )

    chart_path = result.get(
        "chart_path",
        result.get("plot_path"),
    )

    if chart_path:
        chart_file = Path(str(chart_path))

        if chart_file.exists():
            document.add_picture(
                str(chart_file),
                width=Inches(8.5),
            )

            document.paragraphs[-1].alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )
        else:
            add_body_paragraph(
                document,
                f"趋势图文件不存在：{chart_file}",
            )
    else:
        add_body_paragraph(
            document,
            "本次任务没有生成趋势图。",
        )

    # ========================================================
    # 八、输出文件
    # ========================================================

    add_heading(
        document,
        "八、输出文件",
        level=1,
    )

    output_files = result.get(
        "output_files",
        {},
    )

    if isinstance(output_files, dict) and output_files:
        for key, value in output_files.items():
            add_bullet(
                document,
                f"{key}：{value}",
            )

    for key in (
        "excel_path",
        "cleaned_excel_path",
        "statistics_path",
        "chart_path",
        "plot_path",
    ):
        value = result.get(key)

        if value:
            add_bullet(
                document,
                f"{key}：{value}",
            )

    add_bullet(
        document,
        f"Word 报告：{output_path}",
    )

    # ========================================================
    # 保存
    # ========================================================

    document.save(output_path)

    return str(output_path)


if __name__ == "__main__":
    print("batch_report_generator.py 已加载成功。")