import os
from pathlib import Path
from typing import Any, Dict, List, Optional


TEXT_EXTENSIONS = {
    ".txt",
    ".md",
}

WORD_EXTENSIONS = {
    ".docx",
}

PDF_EXTENSIONS = {
    ".pdf",
}

SUPPORTED_DOCUMENT_EXTENSIONS = (
    TEXT_EXTENSIONS
    | WORD_EXTENSIONS
    | PDF_EXTENSIONS
)


# ============================================================
# 自动扫描排除规则
# ============================================================

# 只影响文件夹自动扫描。用户明确选择的 DOCX / PDF / TXT / MD
# 仍由 validate_document_file / read_document 正常读取。
AUTO_SCAN_EXCLUDED_DIRS = {
    "_internal",
    ".venv",
    "venv",
    "env",
    "outputs",
    "output",
    "__pycache__",
    ".git",
    ".github",
    ".idea",
    ".vscode",
    "build",
    "dist",
    "node_modules",
}

AUTO_SCAN_EXCLUDED_DIRS_NORMALIZED = {
    name.casefold()
    for name in AUTO_SCAN_EXCLUDED_DIRS
}


def is_auto_scan_excluded_path(path, scan_root) -> bool:
    """按目录组成部分判断自动扫描候选是否位于排除目录中。"""
    candidate = normalize_document_path(path)
    root = normalize_document_path(scan_root)

    try:
        relative = candidate.relative_to(root)
    except ValueError:
        return False

    parts = relative.parts[:-1] if candidate.is_file() else relative.parts

    return any(
        part.casefold() in AUTO_SCAN_EXCLUDED_DIRS_NORMALIZED
        for part in parts
    )


def iter_document_scan_candidates(folder: Path, recursive: bool):
    """遍历文档候选，并在递归阶段直接剪枝程序/输出目录。"""
    if not recursive:
        yield from folder.glob("*")
        return

    for current_root, dir_names, file_names in os.walk(folder):
        dir_names[:] = [
            name
            for name in dir_names
            if name.casefold() not in AUTO_SCAN_EXCLUDED_DIRS_NORMALIZED
            and not name.startswith(".")
        ]

        current = Path(current_root)

        for file_name in file_names:
            yield current / file_name


def normalize_document_path(file_path) -> Path:
    """
    将输入路径标准化为绝对 Path。
    """
    path = Path(file_path).expanduser()

    try:
        return path.resolve()
    except Exception:
        return path.absolute()


def validate_document_file(file_path) -> Path:
    """
    检查文件是否存在，以及是否属于当前支持的办公文档格式。
    """
    path = normalize_document_path(file_path)

    if not path.exists():
        raise FileNotFoundError(
            f"文件不存在：{path}"
        )

    if not path.is_file():
        raise ValueError(
            f"路径不是文件：{path}"
        )

    extension = path.suffix.lower()

    if extension not in SUPPORTED_DOCUMENT_EXTENSIONS:
        raise ValueError(
            f"暂不支持的文档格式：{extension}"
        )

    return path


def read_text_file(
    file_path,
    encodings: Optional[List[str]] = None,
) -> str:
    """
    读取 TXT / Markdown 文本文件。

    自动尝试常见中文编码。
    """
    path = validate_document_file(file_path)

    if path.suffix.lower() not in TEXT_EXTENSIONS:
        raise ValueError(
            f"不是 TXT / Markdown 文件：{path}"
        )

    if encodings is None:
        encodings = [
            "utf-8-sig",
            "utf-8",
            "gb18030",
            "gbk",
        ]

    errors = []

    for encoding in encodings:
        try:
            return path.read_text(
                encoding=encoding
            )

        except UnicodeDecodeError as error:
            errors.append(
                f"{encoding}: {error}"
            )

    raise UnicodeError(
        "无法识别文本文件编码："
        f"{path}\n"
        + "\n".join(errors)
    )


def read_word_file(file_path) -> str:
    """
    读取 DOCX 文件中的正文和表格文本。

    返回纯文本，供后续 LLM 理解。
    """
    path = validate_document_file(file_path)

    if path.suffix.lower() not in WORD_EXTENSIONS:
        raise ValueError(
            f"不是 DOCX 文件：{path}"
        )

    try:
        from docx import Document

    except ImportError as error:
        raise ImportError(
            "读取 Word 文件需要 python-docx。\n"
            "请运行：pip install python-docx"
        ) from error

    document = Document(
        str(path)
    )

    sections = []

    # ------------------------------------------------------------
    # 正文段落
    # ------------------------------------------------------------

    paragraph_texts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            paragraph_texts.append(
                text
            )

    if paragraph_texts:
        sections.append(
            "\n".join(paragraph_texts)
        )

    # ------------------------------------------------------------
    # 表格
    # ------------------------------------------------------------

    for table_index, table in enumerate(
        document.tables,
        start=1,
    ):
        table_lines = [
            f"[表格 {table_index}]"
        ]

        for row in table.rows:
            cells = []

            for cell in row.cells:
                cell_text = (
                    cell.text
                    .replace("\n", " ")
                    .strip()
                )

                cells.append(
                    cell_text
                )

            table_lines.append(
                " | ".join(cells)
            )

        sections.append(
            "\n".join(table_lines)
        )

    return "\n\n".join(
        sections
    ).strip()


def read_pdf_file(file_path) -> str:
    """
    读取普通文本型 PDF。

    注意：
    当前版本不做 OCR。
    扫描版 PDF 如果没有文本层，可能读取不到正文。
    """
    path = validate_document_file(file_path)

    if path.suffix.lower() not in PDF_EXTENSIONS:
        raise ValueError(
            f"不是 PDF 文件：{path}"
        )

    try:
        from pypdf import PdfReader

    except ImportError as error:
        raise ImportError(
            "读取 PDF 文件需要 pypdf。\n"
            "请运行：pip install pypdf"
        ) from error

    reader = PdfReader(
        str(path)
    )

    pages = []

    for page_number, page in enumerate(
        reader.pages,
        start=1,
    ):
        try:
            text = page.extract_text() or ""

        except Exception as error:
            text = (
                f"[第 {page_number} 页读取失败："
                f"{error}]"
            )

        text = text.strip()

        if text:
            pages.append(
                f"[第 {page_number} 页]\n{text}"
            )

    return "\n\n".join(
        pages
    ).strip()


def read_document(file_path) -> str:
    """
    根据文件扩展名自动选择读取方式。

    支持：
    - TXT
    - Markdown
    - DOCX
    - PDF
    """
    path = validate_document_file(
        file_path
    )

    extension = path.suffix.lower()

    if extension in TEXT_EXTENSIONS:
        return read_text_file(
            path
        )

    if extension in WORD_EXTENSIONS:
        return read_word_file(
            path
        )

    if extension in PDF_EXTENSIONS:
        return read_pdf_file(
            path
        )

    raise ValueError(
        f"暂不支持的文档格式：{extension}"
    )


def get_document_preview(
    file_path,
    max_characters: int = 2000,
) -> str:
    """
    获取文档预览。

    后续可以把预览交给 DeepSeek，
    让模型判断该文档是否与用户任务相关。
    """
    if max_characters <= 0:
        raise ValueError(
            "max_characters 必须大于 0。"
        )

    text = read_document(
        file_path
    )

    if len(text) <= max_characters:
        return text

    return (
        text[:max_characters]
        + "\n\n[文档内容已截断]"
    )


def inspect_document(
    file_path,
    preview_characters: int = 1000,
) -> Dict[str, Any]:
    """
    获取单个文档的基础画像。

    后续 v3.0 的文件选择 Agent
    会使用这些画像判断哪些文件需要阅读。
    """
    path = normalize_document_path(
        file_path
    )

    result = {
        "file_name": path.name,
        "file_stem": path.stem,
        "file_path": str(path),
        "extension": path.suffix.lower(),
        "size_bytes": None,
        "character_count": 0,
        "preview": "",
        "inspection_success": False,
        "inspection_error": None,
    }

    try:
        validated_path = (
            validate_document_file(
                path
            )
        )

        result["size_bytes"] = (
            validated_path.stat().st_size
        )

        text = read_document(
            validated_path
        )

        result["character_count"] = len(
            text
        )

        if len(text) > preview_characters:
            result["preview"] = (
                text[:preview_characters]
                + "\n[内容已截断]"
            )
        else:
            result["preview"] = text

        result["inspection_success"] = True

    except Exception as error:
        result["inspection_error"] = str(
            error
        )

    return result


def inspect_documents(
    file_paths,
    preview_characters: int = 1000,
) -> List[Dict[str, Any]]:
    """
    批量生成办公文档画像。
    """
    results = []

    for file_path in file_paths:
        results.append(
            inspect_document(
                file_path=file_path,
                preview_characters=preview_characters,
            )
        )

    return results


def scan_document_files(
    folder_path,
    recursive: bool = True,
) -> List[str]:
    """
    扫描文件夹中的办公文档。

    当前支持：
    - .txt
    - .md
    - .docx
    - .pdf
    """
    folder = Path(
        folder_path
    ).expanduser()

    try:
        folder = folder.resolve()
    except Exception:
        folder = folder.absolute()

    if not folder.exists():
        raise FileNotFoundError(
            f"文件夹不存在：{folder}"
        )

    if not folder.is_dir():
        raise ValueError(
            f"路径不是文件夹：{folder}"
        )

    candidates = iter_document_scan_candidates(
        folder=folder,
        recursive=recursive,
    )

    files = []

    for path in candidates:
        if not path.is_file():
            continue

        # 忽略 Word / Office 临时文件
        if path.name.startswith("~$"):
            continue

        # 忽略隐藏文件
        if path.name.startswith("."):
            continue

        if is_auto_scan_excluded_path(
            path=path,
            scan_root=folder,
        ):
            continue

        if (
            path.suffix.lower()
            not in SUPPORTED_DOCUMENT_EXTENSIONS
        ):
            continue

        try:
            resolved = str(
                path.resolve()
            )
        except Exception:
            resolved = str(
                path.absolute()
            )

        files.append(
            resolved
        )

    return sorted(
        files,
        key=lambda item: item.lower(),
    )


def build_document_catalog_text(
    document_infos: List[Dict[str, Any]],
    include_preview: bool = True,
) -> str:
    """
    将文档画像转换成适合交给 LLM 的文本目录。
    """
    if not document_infos:
        return "没有发现可读取的办公文档。"

    blocks = []

    for index, info in enumerate(
        document_infos,
        start=1,
    ):
        lines = [
            f"[文档 {index}]",
            f"文件名：{info.get('file_name', '')}",
            f"路径：{info.get('file_path', '')}",
            f"类型：{info.get('extension', '')}",
            (
                "字符数："
                f"{info.get('character_count', 0)}"
            ),
        ]

        if info.get(
            "inspection_success"
        ):
            lines.append(
                "读取状态：成功"
            )

            if include_preview:
                preview = (
                    info.get(
                        "preview",
                        "",
                    )
                    .strip()
                )

                if preview:
                    lines.extend(
                        [
                            "内容预览：",
                            preview,
                        ]
                    )

        else:
            lines.append(
                "读取状态：失败"
            )

            error = info.get(
                "inspection_error"
            )

            if error:
                lines.append(
                    f"错误：{error}"
                )

        blocks.append(
            "\n".join(lines)
        )

    return "\n\n".join(
        blocks
    )


def main():
    """
    独立底层测试。

    扫描当前目录中的 Word / PDF / TXT / Markdown，
    读取并打印文档画像。
    """
    current_folder = Path.cwd()

    print(
        "=" * 70
    )

    print(
        "DataPilot v3.0 多格式办公文档读取测试"
    )

    print(
        "=" * 70
    )

    print(
        f"\n扫描目录：{current_folder}"
    )

    document_files = scan_document_files(
        folder_path=current_folder,
        recursive=False,
    )

    print(
        f"\n发现 {len(document_files)} 个办公文档。"
    )

    if not document_files:
        print(
            "\n当前目录没有发现 "
            "TXT / Markdown / DOCX / PDF。"
        )

        print(
            "\n你可以先放一个测试文档到 "
            "F:\\DataPilot 再运行。"
        )

        print(
            "\n" + "=" * 70
        )

        return

    for file_path in document_files:
        print(
            f"  - {Path(file_path).name}"
        )

    print(
        "\n" + "-" * 70
    )

    print(
        "正在读取文档并生成画像……\n"
    )

    infos = inspect_documents(
        document_files,
        preview_characters=500,
    )

    print(
        build_document_catalog_text(
            infos,
            include_preview=True,
        )
    )

    success_count = sum(
        1
        for info in infos
        if info.get(
            "inspection_success"
        )
    )

    failed_count = (
        len(infos)
        - success_count
    )

    print(
        "\n" + "=" * 70
    )

    print(
        "读取结果："
        f"成功 {success_count} 个，"
        f"失败 {failed_count} 个"
    )

    print(
        "v3.0 多格式办公文档读取底层测试完成。"
    )

    print(
        "=" * 70
    )


if __name__ == "__main__":
    main()