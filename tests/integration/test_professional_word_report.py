from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

from professional_word_report_tools import (
    create_professional_word_report,
    inspect_professional_word_report,
)


def assert_true(condition, message):
    if not condition:
        raise AssertionError(message)


def _find_table_containing(document, *needles):
    for table in document.tables:
        text = "\n".join(
            cell.text
            for row in table.rows
            for cell in row.cells
        )
        if all(needle in text for needle in needles):
            return table
    return None


def main():
    print("=" * 72)
    print("DataPilot v5.0 Professional Word Delivery 第一阶段确定性测试")
    print("=" * 72)

    with TemporaryDirectory(prefix="datapilot_professional_word_") as temp_dir:
        output_path = Path(temp_dir) / "城市销售分析汇报.docx"

        result = create_professional_word_report(
            output_path=str(output_path),
            report_title="城市销售分析汇报",
            subtitle="按城市汇总销售额，可直接用于管理层汇报",
            metadata={
                "报告对象": "管理层",
                "统计口径": "按城市汇总销售额",
            },
            executive_summary=(
                "本次城市销售汇总显示，澳门销售额最高，为 186；"
                "横琴为 154，珠海为 128。"
            ),
            kpis=[
                {"label": "销售冠军", "value": "澳门"},
                {"label": "最高销售额", "value": 186},
                {"label": "销售总额", "value": 468},
            ],
            sections=[
                {
                    "title": "关键结论",
                    "type": "bullets",
                    "items": [
                        "澳门销售额为 186，排名第一。",
                        "横琴销售额为 154。",
                        "珠海销售额为 128。",
                    ],
                },
                {
                    "title": "城市销售汇总",
                    "type": "table",
                    "columns": ["城市", "销售额合计"],
                    "rows": [
                        {"城市": "澳门", "销售额合计": 186},
                        {"城市": "横琴", "销售额合计": 154},
                        {"城市": "珠海", "销售额合计": 128},
                    ],
                },
            ],
            source_note="数据来源：销售数据.xlsx；报告数字来自真实汇总结果。",
        )

        print("\n测试 1：创建真实 .docx")
        assert_true(result["success"] is True, result)
        assert_true(output_path.exists(), f"文件不存在：{output_path}")
        print("PASS")

        document = Document(str(output_path))
        paragraph_text = "\n".join(p.text for p in document.paragraphs)

        print("\n测试 2：标题、副标题与摘要写入")
        assert_true("城市销售分析汇报" in paragraph_text, "缺少标题。")
        assert_true("按城市汇总销售额" in paragraph_text, "缺少副标题。")
        assert_true("执行摘要" in paragraph_text, "缺少执行摘要标题。")
        assert_true("澳门销售额最高" in paragraph_text, "缺少执行摘要内容。")
        print("PASS")

        print("\n测试 3：KPI 使用独立结构化表格")
        kpi_table = _find_table_containing(
            document,
            "销售冠军",
            "澳门",
            "最高销售额",
            "186",
        )
        assert_true(kpi_table is not None, "没有找到 KPI 结构化表格。")
        kpi_text = " ".join(
            cell.text
            for row in kpi_table.rows
            for cell in row.cells
        )
        assert_true(
            "销售冠军" in kpi_text
            and "澳门" in kpi_text
            and "最高销售额" in kpi_text
            and "186" in kpi_text,
            "KPI 内容错误。",
        )
        print("PASS")

        print("\n测试 4：业务表真实生成")
        business_table = _find_table_containing(
            document,
            "城市",
            "销售额合计",
            "澳门",
            "横琴",
            "珠海",
        )
        assert_true(business_table is not None, "没有找到城市销售业务表。")
        assert_true(len(business_table.rows) == 4, "业务表行数错误。")
        print("PASS")

        print("\n测试 5：Word 标题层级真实存在")
        headings = [
            p.text.strip()
            for p in document.paragraphs
            if p.text.strip()
            and (
                getattr(p.style, "name", "") == "Title"
                or getattr(p.style, "name", "").startswith("Heading")
            )
        ]
        for expected in (
            "城市销售分析汇报",
            "执行摘要",
            "核心指标",
            "关键结论",
            "城市销售汇总",
            "数据与来源说明",
        ):
            assert_true(expected in headings, f"缺少标题层级：{expected}")
        print("PASS")

        print("\n测试 6：DataPilot 页眉存在")
        header_text = " ".join(
            p.text.strip()
            for p in document.sections[0].header.paragraphs
            if p.text.strip()
        )
        assert_true(header_text == "DataPilot", f"页眉错误：{header_text}")
        print("PASS")

        print("\n测试 7：inspect 可重新打开最终 Word")
        inspection = inspect_professional_word_report(str(output_path))
        assert_true(inspection["success"] is True, inspection)
        assert_true(inspection["table_count"] >= 2, inspection)
        print("PASS")

        print("\n测试 8：inspect 能识别标题结构")
        assert_true("执行摘要" in inspection["headings"], inspection["headings"])
        assert_true("城市销售汇总" in inspection["headings"], inspection["headings"])
        print("PASS")

        print("\n测试 9：inspect 能提供业务表证据")
        business_evidence = None
        for table in inspection["tables"]:
            flattened = " ".join(
                cell
                for row in table["preview_rows"]
                for cell in row
            )
            if "城市" in flattened and "销售额合计" in flattened:
                business_evidence = table
                break
        assert_true(business_evidence is not None, inspection["tables"])
        preview = business_evidence["preview_rows"]
        actual = {
            row[0]: row[1]
            for row in preview[1:4]
        }
        assert_true(
            actual == {"澳门": "186", "横琴": "154", "珠海": "128"},
            actual,
        )
        print("PASS")

        print("\n测试 10：来源说明存在")
        assert_true(
            "数据与来源说明" in paragraph_text
            and "销售数据.xlsx" in paragraph_text,
            "缺少来源说明。",
        )
        print("PASS")

        print("\n测试 11：非法扩展名被拒绝")
        try:
            create_professional_word_report(
                output_path=str(Path(temp_dir) / "bad.txt"),
                report_title="bad",
            )
        except ValueError:
            pass
        else:
            raise AssertionError("非法扩展名没有被拒绝。")
        print("PASS")

        print("\n测试 12：非法 section type 被拒绝")
        try:
            create_professional_word_report(
                output_path=str(Path(temp_dir) / "bad.docx"),
                report_title="bad",
                sections=[
                    {"title": "bad", "type": "unknown"}
                ],
            )
        except ValueError:
            pass
        else:
            raise AssertionError("非法 section type 没有被拒绝。")
        print("PASS")

    print("\n" + "=" * 72)
    print("Professional Word Report：12/12 PASS")
    print("=" * 72)


if __name__ == "__main__":
    main()
