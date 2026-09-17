from pathlib import Path
from typing import Any, Dict, Optional

import pandas as pd
import matplotlib.pyplot as plt

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt


def set_cell_text(cell, text: Any, bold: bool = False):
    """
    设置 Word 表格单元格文字，并处理中文字体。
    """
    cell.text = ""

    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    run = paragraph.add_run("" if text is None else str(text))
    run.bold = bold

    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9)

    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_document_font(document: Document):
    """
    设置整个 Word 文档的默认字体。
    """
    styles = document.styles

    for style_name in ["Normal", "Body Text"]:
        try:
            style = styles[style_name]
            style.font.name = "Microsoft YaHei"
            style._element.rPr.rFonts.set(
                qn("w:eastAsia"),
                "Microsoft YaHei"
            )
            style.font.size = Pt(10)
        except Exception:
            pass

    for section in document.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)


def add_heading(document: Document, text: str, level: int = 1):
    """
    添加标题并设置中文字体。
    """
    paragraph = document.add_heading(text, level=level)

    for run in paragraph.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(
            qn("w:eastAsia"),
            "Microsoft YaHei"
        )

    return paragraph


def add_paragraph(document: Document, text: str = "", bold: bool = False):
    """
    添加普通段落。
    """
    paragraph = document.add_paragraph()
    run = paragraph.add_run("" if text is None else str(text))
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        "Microsoft YaHei"
    )
    run.font.size = Pt(10)
    return paragraph


def add_key_value_table(
    document: Document,
    data: Dict[str, Any],
    title: Optional[str] = None
):
    """
    添加两列表格。
    """
    if title:
        add_heading(document, title, level=2)

    if not data:
        add_paragraph(document, "暂无数据")
        return

    table = document.add_table(
        rows=1,
        cols=2
    )

    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    set_cell_text(table.rows[0].cells[0], "项目", bold=True)
    set_cell_text(table.rows[0].cells[1], "结果", bold=True)

    for key, value in data.items():
        row_cells = table.add_row().cells
        set_cell_text(row_cells[0], key)
        set_cell_text(row_cells[1], value)

    document.add_paragraph()


def add_dataframe_table(
    document: Document,
    dataframe: Optional[pd.DataFrame],
    title: Optional[str] = None,
    max_rows: int = 20
):
    """
    将 DataFrame 添加到 Word 表格中。
    """
    if title:
        add_heading(document, title, level=2)

    if dataframe is None:
        add_paragraph(document, "暂无数据")
        return

    if not isinstance(dataframe, pd.DataFrame):
        add_paragraph(document, "数据格式不是 DataFrame，无法生成表格")
        return

    if dataframe.empty:
        add_paragraph(document, "数据为空")
        return

    display_df = dataframe.head(max_rows).copy()

    table = document.add_table(
        rows=1,
        cols=len(display_df.columns)
    )

    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    for index, column in enumerate(display_df.columns):
        set_cell_text(
            table.rows[0].cells[index],
            column,
            bold=True
        )

    for _, row in display_df.iterrows():
        cells = table.add_row().cells

        for index, value in enumerate(row):
            if pd.isna(value):
                value = ""

            if isinstance(value, float):
                value = round(value, 4)

            set_cell_text(cells[index], value)

    if len(dataframe) > max_rows:
        add_paragraph(
            document,
            f"注：表格仅展示前 {max_rows} 行，原始数据共 {len(dataframe)} 行。"
        )

    document.add_paragraph()


def format_quality_result(quality: Any) -> Dict[str, Any]:
    """
    将质量检查结果转换成适合 Word 展示的字典。
    """
    if quality is None:
        return {"结果": "暂无质量检查结果"}

    if isinstance(quality, dict):
        result = {}

        for key, value in quality.items():
            if isinstance(value, dict):
                result[key] = str(value)
            elif isinstance(value, float):
                result[key] = round(value, 4)
            else:
                result[key] = value

        return result

    return {"结果": str(quality)}


def format_cleaning_log(cleaning_log: Any) -> Dict[str, Any]:
    """
    将清洗日志转换成适合 Word 展示的字典。
    """
    if cleaning_log is None:
        return {"结果": "暂无清洗记录"}

    if isinstance(cleaning_log, dict):
        return cleaning_log

    if isinstance(cleaning_log, list):
        result = {}

        for index, item in enumerate(cleaning_log, start=1):
            result[f"清洗步骤 {index}"] = item

        return result

    return {"结果": str(cleaning_log)}


def statistics_to_dataframe(statistics: Any) -> Optional[pd.DataFrame]:
    """
    将统计结果转换为 DataFrame。
    """
    if statistics is None:
        return None

    if isinstance(statistics, pd.DataFrame):
        return statistics

    if isinstance(statistics, dict):
        rows = []

        for key, value in statistics.items():
            if isinstance(value, dict):
                for sub_key, sub_value in value.items():
                    rows.append({
                        "统计项目": f"{key}_{sub_key}",
                        "统计结果": sub_value
                    })
            else:
                rows.append({
                    "统计项目": key,
                    "统计结果": value
                })

        if rows:
            return pd.DataFrame(rows)

    if isinstance(statistics, list):
        try:
            return pd.DataFrame(statistics)
        except Exception:
            return pd.DataFrame({"统计结果": statistics})

    return pd.DataFrame({"统计结果": [str(statistics)]})


def extract_report_data(
    result: Optional[Dict[str, Any]] = None,
    original_df: Optional[pd.DataFrame] = None,
    cleaned_df: Optional[pd.DataFrame] = None,
    before_quality: Any = None,
    after_quality: Any = None,
    cleaning_log: Any = None,
    statistics_df: Optional[pd.DataFrame] = None,
    chart_path: Optional[str] = None,
    excel_path: Optional[str] = None
):
    """
    兼容新版 result 字典和旧版独立参数。
    """
    if result is None:
        result = {}

    if not isinstance(result, dict):
        result = {}

    if original_df is None:
        original_df = result.get("original_df")

    if cleaned_df is None:
        cleaned_df = result.get("cleaned_df")

    if before_quality is None:
        before_quality = result.get("before_quality")

    if after_quality is None:
        after_quality = result.get("after_quality")

    if cleaning_log is None:
        cleaning_log = result.get("cleaning_log")

    if statistics_df is None:
        statistics_value = result.get("statistics")

        if isinstance(statistics_value, pd.DataFrame):
            statistics_df = statistics_value
        else:
            statistics_df = statistics_to_dataframe(statistics_value)

    if chart_path is None:
        chart_path = (
            result.get("chart_path")
            or result.get("plot_path")
        )

    if excel_path is None:
        excel_path = result.get("excel_path")

    return {
        "original_df": original_df,
        "cleaned_df": cleaned_df,
        "before_quality": before_quality,
        "after_quality": after_quality,
        "cleaning_log": cleaning_log,
        "statistics_df": statistics_df,
        "chart_path": chart_path,
        "excel_path": excel_path,
    }


def generate_word_report(
    result: Optional[Dict[str, Any]] = None,
    output_path: Optional[str] = None,
    user_task: Optional[str] = None,
    original_df: Optional[pd.DataFrame] = None,
    cleaned_df: Optional[pd.DataFrame] = None,
    before_quality: Any = None,
    after_quality: Any = None,
    cleaning_log: Any = None,
    statistics_df: Optional[pd.DataFrame] = None,
    chart_path: Optional[str] = None,
    excel_path: Optional[str] = None
):
    """
    生成 Word 分析报告。

    兼容两种调用方式：

    方式一：新版 Agent 调用

        generate_word_report(
            result=result,
            output_path="outputs/report.docx"
        )

    方式二：旧版独立参数调用

        generate_word_report(
            original_df=original_df,
            cleaned_df=cleaned_df,
            before_quality=before_quality,
            after_quality=after_quality,
            cleaning_log=cleaning_log,
            statistics_df=statistics_df,
            output_path="outputs/report.docx"
        )
    """
    data = extract_report_data(
        result=result,
        original_df=original_df,
        cleaned_df=cleaned_df,
        before_quality=before_quality,
        after_quality=after_quality,
        cleaning_log=cleaning_log,
        statistics_df=statistics_df,
        chart_path=chart_path,
        excel_path=excel_path
    )

    original_df = data["original_df"]
    cleaned_df = data["cleaned_df"]
    before_quality = data["before_quality"]
    after_quality = data["after_quality"]
    cleaning_log = data["cleaning_log"]
    statistics_df = data["statistics_df"]
    chart_path = data["chart_path"]
    excel_path = data["excel_path"]

    if output_path is None:
        output_path = "outputs/data_analysis_report.docx"

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    set_document_font(document)

    # 标题
    title = document.add_heading(
        "DataPilot 数据分析报告",
        level=0
    )
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for run in title.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(
            qn("w:eastAsia"),
            "Microsoft YaHei"
        )
        run.font.size = Pt(20)

    if user_task:
        add_paragraph(document, f"用户任务：{user_task}")

    add_paragraph(
        document,
        f"报告生成时间：{pd.Timestamp.now().strftime('%Y-%m-%d %H:%M:%S')}"
    )

    # 一、数据概况
    add_heading(document, "一、数据概况", level=1)

    overview = {}

    if isinstance(original_df, pd.DataFrame):
        overview["原始数据行数"] = original_df.shape[0]
        overview["原始数据列数"] = original_df.shape[1]
        overview["原始数据字段"] = "、".join(
            [str(column) for column in original_df.columns]
        )
    else:
        overview["原始数据"] = "未获取到原始 DataFrame"

    if isinstance(cleaned_df, pd.DataFrame):
        overview["清洗后数据行数"] = cleaned_df.shape[0]
        overview["清洗后数据列数"] = cleaned_df.shape[1]

    add_key_value_table(document, overview)

    # 二、清洗前质量检查
    add_heading(document, "二、清洗前数据质量检查", level=1)
    add_key_value_table(
        document,
        format_quality_result(before_quality)
    )

    # 三、数据清洗记录
    add_heading(document, "三、数据清洗记录", level=1)
    add_key_value_table(
        document,
        format_cleaning_log(cleaning_log)
    )

    # 四、清洗后质量检查
    add_heading(document, "四、清洗后数据质量检查", level=1)
    add_key_value_table(
        document,
        format_quality_result(after_quality)
    )

    # 五、统计分析
    add_heading(document, "五、描述性统计分析", level=1)

    if isinstance(statistics_df, pd.DataFrame):
        add_dataframe_table(
            document,
            statistics_df,
            max_rows=50
        )
    else:
        add_paragraph(document, "暂无统计分析结果")

    # 六、清洗后数据预览
    add_heading(document, "六、清洗后数据预览", level=1)

    if isinstance(cleaned_df, pd.DataFrame):
        add_dataframe_table(
            document,
            cleaned_df,
            max_rows=15
        )
    else:
        add_paragraph(document, "暂无清洗后数据")

    # 七、统计图表
    add_heading(document, "七、统计图表", level=1)

    if chart_path:
        chart_file = Path(str(chart_path))

        if chart_file.exists():
            try:
                document.add_picture(
                    str(chart_file),
                    width=Inches(6.2)
                )

                last_paragraph = document.paragraphs[-1]
                last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            except Exception as error:
                add_paragraph(
                    document,
                    f"统计图插入失败：{error}"
                )
        else:
            add_paragraph(
                document,
                f"统计图文件不存在：{chart_file}"
            )
    else:
        add_paragraph(document, "未生成统计图")

    # 八、输出文件
    add_heading(document, "八、输出文件", level=1)

    output_files = {
        "Word 报告": str(output_path)
    }

    if chart_path:
        output_files["统计图"] = str(chart_path)

    if excel_path:
        output_files["Excel 文件"] = str(excel_path)

    add_key_value_table(document, output_files)

    document.save(str(output_path))

    return str(output_path)


if __name__ == "__main__":
    print("report_generator.py 已加载")