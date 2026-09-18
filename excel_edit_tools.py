from copy import copy
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.formula.translate import Translator


def normalize_excel_path(file_path) -> Path:
    path = Path(file_path).expanduser()

    try:
        path = path.resolve()
    except Exception:
        path = path.absolute()

    return path


def validate_excel_file(file_path) -> Path:
    """
    校验可编辑 Excel 文件。

    v3.2 编辑层只直接修改 .xlsx。
    .xls 仍可由原有数据读取工具分析，但不在本编辑器中原位修改。
    """
    path = normalize_excel_path(file_path)

    if not path.exists():
        raise FileNotFoundError(
            f"Excel 文件不存在：{path}"
        )

    if not path.is_file():
        raise ValueError(
            f"路径不是文件：{path}"
        )

    if path.suffix.lower() != ".xlsx":
        raise ValueError(
            f"当前 Excel 编辑工具只支持 .xlsx：{path}"
        )

    return path


def build_output_path(
    source_path,
    output_path=None,
    suffix="_DataPilot编辑",
) -> Path:
    """
    生成新 Excel 输出路径。

    默认另存，不覆盖源文件。
    """
    source = validate_excel_file(
        source_path
    )

    if output_path:
        target = Path(
            output_path
        ).expanduser()

        if target.suffix.lower() != ".xlsx":
            target = target.with_suffix(
                ".xlsx"
            )

        try:
            target = target.resolve()
        except Exception:
            target = target.absolute()
    else:
        target = source.with_name(
            f"{source.stem}{suffix}.xlsx"
        )

    if target == source:
        raise ValueError(
            "为保护原文件，默认禁止直接覆盖源 Excel。"
        )

    target.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    return target


def _get_sheet(workbook, sheet_name=None):
    if sheet_name is None:
        return workbook[
            workbook.sheetnames[0]
        ]

    if sheet_name not in workbook.sheetnames:
        raise ValueError(
            f"Excel 中不存在 Sheet：{sheet_name}。"
            f"可用 Sheet：{', '.join(workbook.sheetnames)}"
        )

    return workbook[
        sheet_name
    ]


def _header_map(worksheet) -> Dict[str, int]:
    result = {}

    for cell in worksheet[1]:
        if cell.value is None:
            continue

        name = str(
            cell.value
        ).strip()

        if name and name not in result:
            result[name] = cell.column

    return result


def _require_column(
    worksheet,
    column_name,
) -> int:
    headers = _header_map(
        worksheet
    )

    if column_name not in headers:
        raise ValueError(
            f"Sheet {worksheet.title} 中不存在字段："
            f"{column_name}。可用字段："
            + ", ".join(headers.keys())
        )

    return headers[
        column_name
    ]


def _copy_cell_style(
    source_cell,
    target_cell,
):
    if source_cell.has_style:
        target_cell._style = copy(
            source_cell._style
        )

    if source_cell.number_format:
        target_cell.number_format = (
            source_cell.number_format
        )

    if source_cell.font:
        target_cell.font = copy(
            source_cell.font
        )

    if source_cell.fill:
        target_cell.fill = copy(
            source_cell.fill
        )

    if source_cell.border:
        target_cell.border = copy(
            source_cell.border
        )

    if source_cell.alignment:
        target_cell.alignment = copy(
            source_cell.alignment
        )

    if source_cell.protection:
        target_cell.protection = copy(
            source_cell.protection
        )


def inspect_excel_workbook(
    file_path,
) -> Dict[str, Any]:
    """
    返回 Excel 工作簿结构，供 Agent 编辑前理解文件。
    """
    source = validate_excel_file(
        file_path
    )

    workbook = load_workbook(
        source,
        data_only=False,
    )

    sheets = []

    for worksheet in workbook.worksheets:
        headers = [
            cell.value
            for cell in worksheet[1]
            if cell.value is not None
        ]

        sheets.append(
            {
                "sheet_name": worksheet.title,
                "max_row": worksheet.max_row,
                "max_column": worksheet.max_column,
                "headers": headers,
            }
        )

    return {
        "file_path": str(source),
        "sheet_names": workbook.sheetnames,
        "sheets": sheets,
    }


def read_excel_sheet_records(
    file_path,
    sheet_name=None,
    max_rows=50,
) -> Dict[str, Any]:
    """
    读取指定 Sheet 的表头和前若干行数据，便于确定性核验。
    """
    source = validate_excel_file(
        file_path
    )

    workbook = load_workbook(
        source,
        data_only=False,
    )
    worksheet = _get_sheet(
        workbook,
        sheet_name,
    )

    headers = [
        cell.value
        for cell in worksheet[1]
    ]

    records = []

    end_row = min(
        worksheet.max_row,
        int(max_rows) + 1,
    )

    for row_index in range(
        2,
        end_row + 1,
    ):
        record = {}

        for column_index, header in enumerate(
            headers,
            start=1,
        ):
            if header is None:
                continue

            record[
                str(header)
            ] = worksheet.cell(
                row=row_index,
                column=column_index,
            ).value

        records.append(
            record
        )

    return {
        "file_path": str(source),
        "sheet_name": worksheet.title,
        "headers": [
            header
            for header in headers
            if header is not None
        ],
        "row_count": max(
            worksheet.max_row - 1,
            0,
        ),
        "records": records,
    }


def _replace_values(
    worksheet,
    old_value,
    new_value,
    column=None,
) -> int:
    target_column = None

    if column:
        target_column = _require_column(
            worksheet,
            column,
        )

    changed = 0

    for row_index in range(
        2,
        worksheet.max_row + 1,
    ):
        column_indexes = (
            [target_column]
            if target_column
            else range(
                1,
                worksheet.max_column + 1,
            )
        )

        for column_index in column_indexes:
            cell = worksheet.cell(
                row=row_index,
                column=column_index,
            )

            if cell.value == old_value:
                cell.value = new_value
                changed += 1

    return changed


def _add_column(
    worksheet,
    column_name,
    value=None,
    values=None,
) -> int:
    headers = _header_map(
        worksheet
    )

    if column_name in headers:
        raise ValueError(
            f"字段已存在：{column_name}"
        )

    new_column = (
        worksheet.max_column + 1
    )

    header_cell = worksheet.cell(
        row=1,
        column=new_column,
    )
    header_cell.value = column_name

    if worksheet.max_column > 1:
        _copy_cell_style(
            worksheet.cell(
                row=1,
                column=new_column - 1,
            ),
            header_cell,
        )

    data_row_count = max(
        worksheet.max_row - 1,
        0,
    )

    if values is not None:
        if not isinstance(
            values,
            list,
        ):
            raise TypeError(
                "add_column 的 values 必须是列表。"
            )

        if len(values) != data_row_count:
            raise ValueError(
                f"add_column 的 values 数量为 {len(values)}，"
                f"但 Sheet 数据行数为 {data_row_count}。"
            )

        for offset, item in enumerate(
            values,
            start=2,
        ):
            worksheet.cell(
                row=offset,
                column=new_column,
            ).value = item
    else:
        for row_index in range(
            2,
            worksheet.max_row + 1,
        ):
            worksheet.cell(
                row=row_index,
                column=new_column,
            ).value = value

    return new_column


def _delete_columns(
    worksheet,
    columns,
) -> List[str]:
    if isinstance(
        columns,
        str,
    ):
        columns = [
            columns
        ]

    headers = _header_map(
        worksheet
    )

    missing = [
        column
        for column in columns
        if column not in headers
    ]

    if missing:
        raise ValueError(
            "以下字段不存在："
            + ", ".join(
                missing
            )
        )

    ordered = sorted(
        (
            (
                headers[column],
                column,
            )
            for column in columns
        ),
        reverse=True,
    )

    deleted = []

    for column_index, column_name in ordered:
        worksheet.delete_cols(
            column_index,
            1,
        )
        deleted.append(
            column_name
        )

    deleted.reverse()
    return deleted


def _rename_columns(
    worksheet,
    rename_map,
) -> int:
    if not isinstance(
        rename_map,
        dict,
    ):
        raise TypeError(
            "rename_map 必须是字典。"
        )

    headers = _header_map(
        worksheet
    )
    changed = 0

    for old_name, new_name in rename_map.items():
        if old_name not in headers:
            raise ValueError(
                f"字段不存在：{old_name}"
            )

        worksheet.cell(
            row=1,
            column=headers[old_name],
        ).value = new_name

        changed += 1

    return changed


def _sort_rows(
    worksheet,
    column,
    ascending=True,
    header_row=None,
) -> int:
    """
    按指定字段排序数据行，并尽量保持 Excel 业务行的完整属性。

    v3.4 高保真策略：
    - 支持顶部标题/合并单元格，表头不再强制位于第 1 行；
    - header_row 未指定时，会在前若干行自动寻找目标字段；
    - 值、公式、样式、数字格式随业务行移动；
    - 超链接、批注随业务行移动；
    - 行高、隐藏状态、outlineLevel 等行级属性随业务行移动；
    - 公式使用 openpyxl Translator 按新位置调整相对引用；
    - 表头以上标题区、合并单元格、列宽、冻结窗格、筛选等不重建。
    """
    if header_row is not None:
        header_row = int(header_row)

        if header_row < 1 or header_row > worksheet.max_row:
            raise ValueError(
                f"header_row 超出有效范围：{header_row}"
            )

        header_cells = worksheet[
            header_row
        ]
        header_map = {}

        for cell in header_cells:
            if cell.value is None:
                continue

            name = str(
                cell.value
            ).strip()

            if name and name not in header_map:
                header_map[name] = cell.column

        if column not in header_map:
            raise ValueError(
                f"Sheet {worksheet.title} 的第 {header_row} 行"
                f"不存在字段：{column}。可用字段："
                + ", ".join(header_map.keys())
            )

        column_index = header_map[
            column
        ]

    else:
        # 真实办公 Excel 常见：
        # 第 1 行是合并标题，第 2/3 行才是真正字段名。
        # 因此针对排序目标字段自动寻找表头行。
        detected_header_row = None
        detected_column_index = None

        search_limit = min(
            worksheet.max_row,
            20,
        )

        for candidate_row in range(
            1,
            search_limit + 1,
        ):
            for cell in worksheet[
                candidate_row
            ]:
                if cell.value is None:
                    continue

                if str(
                    cell.value
                ).strip() == str(
                    column
                ).strip():
                    detected_header_row = (
                        candidate_row
                    )
                    detected_column_index = (
                        cell.column
                    )
                    break

            if detected_header_row is not None:
                break

        if detected_header_row is None:
            available = []

            for candidate_row in range(
                1,
                search_limit + 1,
            ):
                values = [
                    str(cell.value).strip()
                    for cell in worksheet[
                        candidate_row
                    ]
                    if cell.value is not None
                ]

                if values:
                    available.append(
                        f"第{candidate_row}行: "
                        + ", ".join(values)
                    )

            raise ValueError(
                f"Sheet {worksheet.title} 中找不到字段：{column}。"
                "已检查前 "
                f"{search_limit} 行。"
                + (
                    " 检查到的非空内容："
                    + " | ".join(available)
                    if available
                    else ""
                )
            )

        header_row = detected_header_row
        column_index = detected_column_index

    data_start_row = (
        header_row + 1
    )

    if worksheet.max_row < data_start_row:
        return 0

    if worksheet.max_row == data_start_row:
        return 1

    max_column = worksheet.max_column
    source_rows = []

    for row_index in range(
        data_start_row,
        worksheet.max_row + 1,
    ):
        cells = []

        for column_index_inner in range(
            1,
            max_column + 1,
        ):
            cell = worksheet.cell(
                row=row_index,
                column=column_index_inner,
            )

            cells.append(
                {
                    "value": cell.value,
                    "style": copy(cell._style),
                    "hyperlink": copy(cell.hyperlink),
                    "comment": copy(cell.comment),
                }
            )

        row_dimension = worksheet.row_dimensions[
            row_index
        ]

        source_rows.append(
            {
                "source_row": row_index,
                "cells": cells,
                "height": row_dimension.height,
                "hidden": row_dimension.hidden,
                "outline_level": row_dimension.outlineLevel,
                "collapsed": row_dimension.collapsed,
                "thick_top": row_dimension.thickTop,
                "thick_bottom": row_dimension.thickBot,
            }
        )

    def sort_key(item):
        value = item["cells"][
            column_index - 1
        ]["value"]

        if value is None:
            return (
                3,
                "",
            )

        if isinstance(
            value,
            bool,
        ):
            return (
                0,
                int(value),
            )

        if isinstance(
            value,
            (int, float),
        ):
            return (
                0,
                value,
            )

        return (
            1,
            str(value).casefold(),
        )

    non_empty = [
        item
        for item in source_rows
        if item["cells"][
            column_index - 1
        ]["value"] is not None
    ]

    empty = [
        item
        for item in source_rows
        if item["cells"][
            column_index - 1
        ]["value"] is None
    ]

    non_empty.sort(
        key=sort_key,
        reverse=not bool(
            ascending
        ),
    )

    ordered_rows = (
        non_empty
        + empty
    )

    # 只清理真正的数据区，不触碰标题区和表头区。
    for row_index in range(
        data_start_row,
        worksheet.max_row + 1,
    ):
        for column_index_inner in range(
            1,
            max_column + 1,
        ):
            target_cell = worksheet.cell(
                row=row_index,
                column=column_index_inner,
            )
            target_cell.hyperlink = None
            target_cell.comment = None

    for target_row, row_snapshot in enumerate(
        ordered_rows,
        start=data_start_row,
    ):
        source_row = row_snapshot[
            "source_row"
        ]

        for column_index_inner, cell_snapshot in enumerate(
            row_snapshot["cells"],
            start=1,
        ):
            target_cell = worksheet.cell(
                row=target_row,
                column=column_index_inner,
            )

            value = cell_snapshot[
                "value"
            ]

            if (
                isinstance(value, str)
                and value.startswith("=")
                and source_row != target_row
            ):
                source_coordinate = (
                    f"{get_column_letter(column_index_inner)}"
                    f"{source_row}"
                )
                target_coordinate = (
                    f"{get_column_letter(column_index_inner)}"
                    f"{target_row}"
                )

                try:
                    value = Translator(
                        value,
                        origin=source_coordinate,
                    ).translate_formula(
                        target_coordinate
                    )
                except Exception:
                    value = cell_snapshot[
                        "value"
                    ]

            target_cell.value = value
            target_cell._style = copy(
                cell_snapshot["style"]
            )
            target_cell.hyperlink = copy(
                cell_snapshot["hyperlink"]
            )
            target_cell.comment = copy(
                cell_snapshot["comment"]
            )

        target_dimension = worksheet.row_dimensions[
            target_row
        ]
        target_dimension.height = row_snapshot[
            "height"
        ]
        target_dimension.hidden = row_snapshot[
            "hidden"
        ]
        target_dimension.outlineLevel = row_snapshot[
            "outline_level"
        ]
        target_dimension.collapsed = row_snapshot[
            "collapsed"
        ]
        target_dimension.thickTop = row_snapshot[
            "thick_top"
        ]
        target_dimension.thickBot = row_snapshot[
            "thick_bottom"
        ]

    return len(
        ordered_rows
    )

def _delete_duplicate_rows(
    worksheet,
    columns=None,
    keep="first",
) -> int:
    headers = _header_map(
        worksheet
    )

    if columns is None:
        column_indexes = list(
            range(
                1,
                worksheet.max_column + 1,
            )
        )
    else:
        if isinstance(
            columns,
            str,
        ):
            columns = [
                columns
            ]

        missing = [
            column
            for column in columns
            if column not in headers
        ]

        if missing:
            raise ValueError(
                "以下去重字段不存在："
                + ", ".join(
                    missing
                )
            )

        column_indexes = [
            headers[column]
            for column in columns
        ]

    if keep not in {
        "first",
        "last",
    }:
        raise ValueError(
            "Excel 编辑器的 keep 当前只支持 first 或 last。"
        )

    key_to_rows = {}

    for row_index in range(
        2,
        worksheet.max_row + 1,
    ):
        key = tuple(
            worksheet.cell(
                row=row_index,
                column=column_index,
            ).value
            for column_index in column_indexes
        )

        key_to_rows.setdefault(
            key,
            [],
        ).append(
            row_index
        )

    rows_to_delete = []

    for row_indexes in key_to_rows.values():
        if len(row_indexes) <= 1:
            continue

        if keep == "first":
            rows_to_delete.extend(
                row_indexes[1:]
            )
        else:
            rows_to_delete.extend(
                row_indexes[:-1]
            )

    for row_index in sorted(
        rows_to_delete,
        reverse=True,
    ):
        worksheet.delete_rows(
            row_index,
            1,
        )

    return len(
        rows_to_delete
    )


def _fill_missing_values(
    worksheet,
    columns=None,
    fill_value="",
) -> int:
    headers = _header_map(
        worksheet
    )

    if columns is None:
        target_indexes = list(
            headers.values()
        )
    else:
        if isinstance(
            columns,
            str,
        ):
            columns = [
                columns
            ]

        missing = [
            column
            for column in columns
            if column not in headers
        ]

        if missing:
            raise ValueError(
                "以下缺失值处理字段不存在："
                + ", ".join(
                    missing
                )
            )

        target_indexes = [
            headers[column]
            for column in columns
        ]

    changed = 0

    for row_index in range(
        2,
        worksheet.max_row + 1,
    ):
        for column_index in target_indexes:
            cell = worksheet.cell(
                row=row_index,
                column=column_index,
            )

            if cell.value is None:
                cell.value = fill_value
                changed += 1

    return changed


def _append_rows(
    worksheet,
    rows,
) -> int:
    if not isinstance(
        rows,
        list,
    ):
        raise TypeError(
            "append_rows 的 rows 必须是列表。"
        )

    headers = _header_map(
        worksheet
    )

    if not headers:
        raise ValueError(
            "目标 Sheet 没有可识别的表头。"
        )

    header_names = list(
        headers.keys()
    )

    appended = 0

    for row in rows:
        if isinstance(
            row,
            dict,
        ):
            values = [
                row.get(
                    header
                )
                for header in header_names
            ]
        elif isinstance(
            row,
            (list, tuple),
        ):
            if len(row) > len(
                header_names
            ):
                raise ValueError(
                    "追加行的字段数量超过现有表头数量。"
                )

            values = list(
                row
            )
        else:
            raise TypeError(
                "每个追加行必须是字典、列表或元组。"
            )

        worksheet.append(
            values
        )
        appended += 1

    return appended


def _update_cell(
    worksheet,
    row,
    column=None,
    column_name=None,
    value=None,
) -> Dict[str, Any]:
    row = int(
        row
    )

    if row < 2:
        raise ValueError(
            "数据行 row 必须从 2 开始；第 1 行是表头。"
        )

    if column_name:
        column_index = _require_column(
            worksheet,
            column_name,
        )
    elif column is not None:
        column_index = int(
            column
        )

        if column_index < 1:
            raise ValueError(
                "column 必须从 1 开始。"
            )
    else:
        raise ValueError(
            "update_cell 必须提供 column_name 或 column。"
        )

    cell = worksheet.cell(
        row=row,
        column=column_index,
    )
    old_value = cell.value
    cell.value = value

    return {
        "row": row,
        "column": column_index,
        "column_letter": get_column_letter(
            column_index
        ),
        "old_value": old_value,
        "new_value": value,
    }


def apply_excel_edits(
    file_path,
    operations: List[Dict[str, Any]],
    output_path=None,
) -> Dict[str, Any]:
    """
    对现有 XLSX 执行多个编辑操作，并另存为新文件。

    每个 operation 可单独指定 sheet_name。
    如果不指定，则使用工作簿第一个 Sheet。

    支持 action：
    - replace_values
    - add_column
    - delete_columns
    - rename_columns
    - sort_rows
    - delete_duplicate_rows
    - fill_missing_values
    - append_rows
    - update_cell

    设计原则：
    - 使用 openpyxl 直接编辑原工作簿副本；
    - 未操作的其他 Sheet 会保留；
    - 默认不覆盖源文件；
    - 一次执行多个操作，只保存一次。
    """
    source = validate_excel_file(
        file_path
    )
    target = build_output_path(
        source,
        output_path=output_path,
    )

    if not isinstance(
        operations,
        list,
    ):
        raise TypeError(
            "operations 必须是列表。"
        )

    workbook = load_workbook(
        source,
        data_only=False,
    )

    logs = []

    for index, operation in enumerate(
        operations,
        start=1,
    ):
        if not isinstance(
            operation,
            dict,
        ):
            raise TypeError(
                f"第 {index} 个操作必须是字典。"
            )

        action = str(
            operation.get(
                "action"
            )
            or ""
        ).strip()

        sheet_name = operation.get(
            "sheet_name"
        )

        worksheet = _get_sheet(
            workbook,
            sheet_name,
        )

        if action == "replace_values":
            if "old_value" not in operation:
                raise ValueError(
                    f"第 {index} 个 replace_values "
                    "缺少 old_value。"
                )

            changed = _replace_values(
                worksheet,
                old_value=operation.get(
                    "old_value"
                ),
                new_value=operation.get(
                    "new_value"
                ),
                column=operation.get(
                    "column"
                ),
            )

            logs.append(
                {
                    "action": action,
                    "sheet_name": worksheet.title,
                    "changed_cells": changed,
                }
            )

        elif action == "add_column":
            column_name = operation.get(
                "column_name"
            )

            if not column_name:
                raise ValueError(
                    f"第 {index} 个 add_column "
                    "缺少 column_name。"
                )

            column_index = _add_column(
                worksheet,
                column_name=str(
                    column_name
                ),
                value=operation.get(
                    "value"
                ),
                values=operation.get(
                    "values"
                ),
            )

            logs.append(
                {
                    "action": action,
                    "sheet_name": worksheet.title,
                    "column_name": str(
                        column_name
                    ),
                    "column_index": column_index,
                }
            )

        elif action == "delete_columns":
            columns = operation.get(
                "columns"
            )

            if not columns:
                raise ValueError(
                    f"第 {index} 个 delete_columns "
                    "缺少 columns。"
                )

            deleted = _delete_columns(
                worksheet,
                columns,
            )

            logs.append(
                {
                    "action": action,
                    "sheet_name": worksheet.title,
                    "deleted_columns": deleted,
                }
            )

        elif action == "rename_columns":
            rename_map = operation.get(
                "rename_map"
            )

            if not rename_map:
                raise ValueError(
                    f"第 {index} 个 rename_columns "
                    "缺少 rename_map。"
                )

            changed = _rename_columns(
                worksheet,
                rename_map,
            )

            logs.append(
                {
                    "action": action,
                    "sheet_name": worksheet.title,
                    "renamed_columns": changed,
                }
            )

        elif action == "sort_rows":
            column = operation.get(
                "column"
            )

            if not column:
                raise ValueError(
                    f"第 {index} 个 sort_rows "
                    "缺少 column。"
                )

            sorted_rows = _sort_rows(
                worksheet,
                column=str(
                    column
                ),
                ascending=operation.get(
                    "ascending",
                    True,
                ),
                header_row=operation.get(
                    "header_row"
                ),
            )

            logs.append(
                {
                    "action": action,
                    "sheet_name": worksheet.title,
                    "column": str(
                        column
                    ),
                    "ascending": bool(
                        operation.get(
                            "ascending",
                            True,
                        )
                    ),
                    "sorted_rows": sorted_rows,
                }
            )

        elif action == "delete_duplicate_rows":
            deleted = _delete_duplicate_rows(
                worksheet,
                columns=operation.get(
                    "columns"
                ),
                keep=operation.get(
                    "keep",
                    "first",
                ),
            )

            logs.append(
                {
                    "action": action,
                    "sheet_name": worksheet.title,
                    "deleted_rows": deleted,
                }
            )

        elif action == "fill_missing_values":
            changed = _fill_missing_values(
                worksheet,
                columns=operation.get(
                    "columns"
                ),
                fill_value=operation.get(
                    "fill_value",
                    "",
                ),
            )

            logs.append(
                {
                    "action": action,
                    "sheet_name": worksheet.title,
                    "filled_cells": changed,
                }
            )

        elif action == "append_rows":
            rows = operation.get(
                "rows"
            )

            if rows is None:
                raise ValueError(
                    f"第 {index} 个 append_rows "
                    "缺少 rows。"
                )

            appended = _append_rows(
                worksheet,
                rows,
            )

            logs.append(
                {
                    "action": action,
                    "sheet_name": worksheet.title,
                    "appended_rows": appended,
                }
            )

        elif action == "update_cell":
            if "row" not in operation:
                raise ValueError(
                    f"第 {index} 个 update_cell "
                    "缺少 row。"
                )

            cell_result = _update_cell(
                worksheet,
                row=operation.get(
                    "row"
                ),
                column=operation.get(
                    "column"
                ),
                column_name=operation.get(
                    "column_name"
                ),
                value=operation.get(
                    "value"
                ),
            )

            logs.append(
                {
                    "action": action,
                    "sheet_name": worksheet.title,
                    **cell_result,
                }
            )

        else:
            raise ValueError(
                f"暂不支持的 Excel 编辑 action：{action}"
            )

    workbook.save(
        target
    )

    return {
        "success": True,
        "source_path": str(source),
        "output_path": str(target),
        "sheet_names": workbook.sheetnames,
        "operation_count": len(
            operations
        ),
        "operations": logs,
    }


def main():
    print("=" * 70)
    print("DataPilot v3.2 Excel 编辑工具")
    print("=" * 70)
    print(
        "支持：值替换、新增列、删除列、重命名列、排序、"
        "去重、缺失值填充、追加行、单元格修改。"
    )
    print("保留未修改 Sheet，默认另存新文件，不覆盖源 Excel。")
    print("=" * 70)


if __name__ == "__main__":
    main()
