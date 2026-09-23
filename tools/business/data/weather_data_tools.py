from __future__ import annotations

import json
import time
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Dict, Optional
from zoneinfo import ZoneInfo

import pandas as pd
import requests
import matplotlib.pyplot as plt

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


HISTORICAL_ENDPOINT = "https://archive-api.open-meteo.com/v1/archive"
FORECAST_ENDPOINT = "https://api.open-meteo.com/v1/forecast"

HISTORICAL_DAILY_VARIABLES = [
    "temperature_2m_mean",
    "temperature_2m_max",
    "temperature_2m_min",
    "relative_humidity_2m_mean",
    "precipitation_sum",
    "wind_speed_10m_mean",
    "wind_speed_10m_max",
    "wind_gusts_10m_max",
]

FORECAST_DAILY_VARIABLES = HISTORICAL_DAILY_VARIABLES + [
    "precipitation_probability_max",
]

COLUMN_MAP = {
    "time": "date",
    "temperature_2m_mean": "temperature_mean_c",
    "temperature_2m_max": "temperature_max_c",
    "temperature_2m_min": "temperature_min_c",
    "relative_humidity_2m_mean": "relative_humidity_mean_pct",
    "precipitation_sum": "precipitation_sum_mm",
    "wind_speed_10m_mean": "wind_speed_mean_kmh",
    "wind_speed_10m_max": "wind_speed_max_kmh",
    "wind_gusts_10m_max": "wind_gusts_max_kmh",
    "precipitation_probability_max": "precipitation_probability_max_pct",
}

DEFAULT_REQUEST_ATTEMPTS = 3
DEFAULT_RETRY_BACKOFF = 0.75


def _request_json(
    url: str,
    *,
    params: Dict[str, Any],
    timeout: int,
    attempts: int = DEFAULT_REQUEST_ATTEMPTS,
    retry_backoff: float = DEFAULT_RETRY_BACKOFF,
) -> tuple[Dict[str, Any], str]:
    """
    请求 JSON，并对临时网络错误做有界重试。

    仅对 requests 层的连接/超时错误与 HTTP 429/5xx 重试。
    HTTP 4xx 参数错误直接抛出，避免把错误 URL 重试多次。
    """
    attempts = max(1, int(attempts))
    retry_backoff = max(0.0, float(retry_backoff))

    last_error: Optional[Exception] = None

    for attempt in range(1, attempts + 1):
        try:
            response = requests.get(
                url,
                params=params,
                timeout=timeout,
                headers={
                    "User-Agent": (
                        "DataPilot/6.6 weather-data-acquisition"
                    ),
                    "Accept": "application/json",
                },
            )

            status_code = int(
                getattr(response, "status_code", 200)
                or 200
            )

            if status_code == 429 or status_code >= 500:
                if attempt < attempts:
                    if retry_backoff > 0:
                        time.sleep(
                            retry_backoff * attempt
                        )
                    continue

            response.raise_for_status()

            payload = response.json()
            if not isinstance(payload, dict):
                raise ValueError(
                    "Open-Meteo 返回类型异常："
                    f"{type(payload).__name__}"
                )

            if payload.get("error"):
                raise RuntimeError(
                    "Open-Meteo API error: "
                    + str(
                        payload.get("reason")
                        or payload.get("error")
                    )
                )

            return payload, str(response.url)

        except (
            requests.Timeout,
            requests.ConnectionError,
        ) as error:
            last_error = error

            if attempt >= attempts:
                break

            if retry_backoff > 0:
                time.sleep(
                    retry_backoff * attempt
                )

    if last_error is not None:
        raise last_error

    raise RuntimeError(
        "Open-Meteo 请求未返回有效结果。"
    )


def _daily_payload_to_dataframe(
    payload: Dict[str, Any],
    *,
    data_kind: str,
) -> pd.DataFrame:
    daily = payload.get("daily")
    if not isinstance(daily, dict):
        raise ValueError(
            "Open-Meteo 返回中缺少 daily 数据。"
        )

    times = daily.get("time")
    if not isinstance(times, list) or not times:
        raise ValueError(
            "Open-Meteo daily.time 为空。"
        )

    expected_length = len(times)
    data: Dict[str, Any] = {
        "time": times,
    }

    for key, values in daily.items():
        if key == "time":
            continue
        if not isinstance(values, list):
            continue
        if len(values) != expected_length:
            raise ValueError(
                f"Open-Meteo daily.{key} 长度 "
                f"{len(values)} 与 time 长度 "
                f"{expected_length} 不一致。"
            )
        data[key] = values

    frame = pd.DataFrame(data).rename(
        columns=COLUMN_MAP
    )

    if "date" not in frame.columns:
        raise ValueError(
            "天气数据缺少 date 字段。"
        )

    frame["date"] = pd.to_datetime(
        frame["date"],
        errors="raise",
    )
    frame.insert(
        1,
        "data_kind",
        str(data_kind),
    )

    preferred = [
        "date",
        "data_kind",
        "temperature_mean_c",
        "temperature_max_c",
        "temperature_min_c",
        "relative_humidity_mean_pct",
        "precipitation_sum_mm",
        "precipitation_probability_max_pct",
        "wind_speed_mean_kmh",
        "wind_speed_max_kmh",
        "wind_gusts_max_kmh",
    ]

    ordered = [
        column
        for column in preferred
        if column in frame.columns
    ]
    remaining = [
        column
        for column in frame.columns
        if column not in ordered
    ]

    return frame[
        ordered + remaining
    ].copy()


def _resolve_reference_date(
    *,
    reference_date: Optional[str],
    timezone: str,
) -> date:
    if reference_date:
        try:
            return date.fromisoformat(
                str(reference_date).strip()
            )
        except ValueError as error:
            raise ValueError(
                "reference_date 必须是 YYYY-MM-DD。"
            ) from error

    try:
        return datetime.now(
            ZoneInfo(timezone)
        ).date()
    except Exception as error:
        raise ValueError(
            f"无效 timezone：{timezone}"
        ) from error


def _split_combined_payload(
    payload: Dict[str, Any],
    *,
    reference_date: date,
    historical_days: int,
    forecast_days: int,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """
    把 Forecast API 的 past_days + forecast_days 单次结果
    按 reference_date 拆成历史段与未来段。
    """
    frame = _daily_payload_to_dataframe(
        payload,
        data_kind="combined",
    )

    reference_timestamp = pd.Timestamp(
        reference_date
    )

    historical = (
        frame[
            frame["date"] < reference_timestamp
        ]
        .sort_values("date")
        .tail(historical_days)
        .copy()
    )

    forecast = (
        frame[
            frame["date"] >= reference_timestamp
        ]
        .sort_values("date")
        .head(forecast_days)
        .copy()
    )

    historical["data_kind"] = "historical"
    forecast["data_kind"] = "forecast"

    if len(historical) < historical_days:
        raise ValueError(
            "Forecast API past_days 返回的历史天数不足："
            f"需要 {historical_days} 天，"
            f"实际 {len(historical)} 天。"
        )

    if len(forecast) < forecast_days:
        raise ValueError(
            "Forecast API 返回的未来预报天数不足："
            f"需要 {forecast_days} 天，"
            f"实际 {len(forecast)} 天。"
        )

    return (
        historical.reset_index(drop=True),
        forecast.reset_index(drop=True),
    )


def _fetch_combined_recent_weather(
    *,
    latitude: float,
    longitude: float,
    historical_days: int,
    forecast_days: int,
    timezone: str,
    reference_date: date,
    timeout: int,
) -> tuple[
    pd.DataFrame,
    pd.DataFrame,
    Dict[str, Any],
    str,
]:
    """
    首选路径：
    使用 Forecast API 的 past_days + forecast_days
    一次获取最近历史段与未来预报段。

    这样可以减少网络请求次数，并避免历史 endpoint
    与 forecast endpoint 任意一个超时导致整条链失败。
    """
    params = {
        "latitude": latitude,
        "longitude": longitude,
        "past_days": historical_days,
        "forecast_days": forecast_days,
        "daily": ",".join(
            FORECAST_DAILY_VARIABLES
        ),
        "timezone": timezone,
    }

    payload, request_url = _request_json(
        FORECAST_ENDPOINT,
        params=params,
        timeout=timeout,
    )

    historical, forecast = (
        _split_combined_payload(
            payload,
            reference_date=reference_date,
            historical_days=historical_days,
            forecast_days=forecast_days,
        )
    )

    return (
        historical,
        forecast,
        payload,
        request_url,
    )


def _fetch_split_weather(
    *,
    latitude: float,
    longitude: float,
    historical_days: int,
    forecast_days: int,
    timezone: str,
    reference_date: date,
    timeout: int,
) -> tuple[
    pd.DataFrame,
    pd.DataFrame,
    Dict[str, Any],
    Dict[str, Any],
    str,
    str,
]:
    """
    第二路径：
    单请求模式不可用时，退回 Historical API + Forecast API。
    """
    historical_end = (
        reference_date
        - timedelta(days=1)
    )
    historical_start = (
        historical_end
        - timedelta(
            days=historical_days - 1
        )
    )

    historical_params = {
        "latitude": latitude,
        "longitude": longitude,
        "start_date": (
            historical_start.isoformat()
        ),
        "end_date": (
            historical_end.isoformat()
        ),
        "daily": ",".join(
            HISTORICAL_DAILY_VARIABLES
        ),
        "timezone": timezone,
    }

    forecast_params = {
        "latitude": latitude,
        "longitude": longitude,
        "forecast_days": forecast_days,
        "daily": ",".join(
            FORECAST_DAILY_VARIABLES
        ),
        "timezone": timezone,
    }

    historical_payload, historical_url = (
        _request_json(
            HISTORICAL_ENDPOINT,
            params=historical_params,
            timeout=timeout,
        )
    )

    forecast_payload, forecast_url = (
        _request_json(
            FORECAST_ENDPOINT,
            params=forecast_params,
            timeout=timeout,
        )
    )

    historical = (
        _daily_payload_to_dataframe(
            historical_payload,
            data_kind="historical",
        )
        .sort_values("date")
        .tail(historical_days)
        .reset_index(drop=True)
    )

    forecast = (
        _daily_payload_to_dataframe(
            forecast_payload,
            data_kind="forecast",
        )
        .sort_values("date")
        .head(forecast_days)
        .reset_index(drop=True)
    )

    if len(historical) < historical_days:
        raise ValueError(
            "Historical Weather API 返回历史天数不足："
            f"需要 {historical_days} 天，"
            f"实际 {len(historical)} 天。"
        )

    if len(forecast) < forecast_days:
        raise ValueError(
            "Forecast API 返回预报天数不足："
            f"需要 {forecast_days} 天，"
            f"实际 {len(forecast)} 天。"
        )

    return (
        historical,
        forecast,
        historical_payload,
        forecast_payload,
        historical_url,
        forecast_url,
    )


def fetch_weather_dataset(
    latitude: float,
    longitude: float,
    historical_days: int = 30,
    forecast_days: int = 7,
    timezone: str = "Asia/Shanghai",
    reference_date: Optional[str] = None,
    timeout: int = 30,
    output_dir: Optional[str] = None,
) -> Dict[str, Any]:
    """
    获取一个地点的近期历史天气 + 未来天气预报。

    v6.6 fix35:
    1. 优先使用 Forecast API 的 past_days + forecast_days，
       用一个 HTTP 请求同时取得近期历史与未来预报；
    2. 单请求路径失败后，自动退回
       Historical Weather API + Forecast API；
    3. 每个 HTTP 请求内部有界重试；
    4. metadata 明确记录实际采用的数据获取路径和历史数据语义。

    注意：
    - combined 模式中的过去数据属于 Forecast API 的 archived
      / past-days 数据，不应伪装成地面站实测观测；
    - split 模式的历史段来自 Historical Weather API。
    """
    try:
        latitude = float(latitude)
        longitude = float(longitude)
    except (
        TypeError,
        ValueError,
    ) as error:
        raise ValueError(
            "latitude / longitude 必须是数值。"
        ) from error

    if not -90 <= latitude <= 90:
        raise ValueError(
            "latitude 必须在 -90 到 90 之间。"
        )
    if not -180 <= longitude <= 180:
        raise ValueError(
            "longitude 必须在 -180 到 180 之间。"
        )

    historical_days = int(
        historical_days
    )
    forecast_days = int(
        forecast_days
    )
    timeout = int(timeout)

    # Forecast API 官方支持 past_days 0-92。
    # 为保证 combined 首选路径可用，当前工具把近期历史上限设为 92。
    if not 1 <= historical_days <= 92:
        raise ValueError(
            "historical_days 必须在 1-92 之间。"
        )
    if not 1 <= forecast_days <= 16:
        raise ValueError(
            "forecast_days 必须在 1-16 之间；"
            "Open-Meteo Forecast API "
            "不支持 forecast_days=37 这类请求。"
        )
    if timeout <= 0:
        raise ValueError(
            "timeout 必须大于 0。"
        )

    ref_date = _resolve_reference_date(
        reference_date=reference_date,
        timezone=timezone,
    )

    historical_end = (
        ref_date
        - timedelta(days=1)
    )
    historical_start = (
        historical_end
        - timedelta(
            days=historical_days - 1
        )
    )

    acquisition_mode = ""
    fallback_reason = ""
    combined_payload: Dict[str, Any] = {}
    historical_payload: Dict[str, Any] = {}
    forecast_payload: Dict[str, Any] = {}
    historical_url = ""
    forecast_url = ""

    try:
        (
            historical,
            forecast,
            combined_payload,
            combined_url,
        ) = _fetch_combined_recent_weather(
            latitude=latitude,
            longitude=longitude,
            historical_days=historical_days,
            forecast_days=forecast_days,
            timezone=timezone,
            reference_date=ref_date,
            timeout=timeout,
        )

        acquisition_mode = (
            "forecast_api_past_days_combined"
        )
        historical_url = combined_url
        forecast_url = combined_url
        forecast_payload = combined_payload

    except Exception as combined_error:
        fallback_reason = (
            f"{type(combined_error).__name__}: "
            f"{combined_error}"
        )

        (
            historical,
            forecast,
            historical_payload,
            forecast_payload,
            historical_url,
            forecast_url,
        ) = _fetch_split_weather(
            latitude=latitude,
            longitude=longitude,
            historical_days=historical_days,
            forecast_days=forecast_days,
            timezone=timezone,
            reference_date=ref_date,
            timeout=timeout,
        )

        acquisition_mode = (
            "historical_plus_forecast_split"
        )

    combined = pd.concat(
        [historical, forecast],
        ignore_index=True,
        sort=False,
    )

    acquired_at = datetime.now(
        ZoneInfo(timezone)
    ).isoformat(
        timespec="seconds"
    )

    if (
        acquisition_mode
        == "forecast_api_past_days_combined"
    ):
        historical_units = (
            combined_payload.get(
                "daily_units"
            )
            if isinstance(
                combined_payload.get(
                    "daily_units"
                ),
                dict,
            )
            else {}
        )
        historical_semantics = (
            "Open-Meteo Forecast API past_days "
            "返回的近期归档/回填天气数据；"
            "它不是本地地面站逐日实测观测。"
        )
        method_note = (
            "首选单请求模式：使用 Open-Meteo "
            "Forecast API 的 past_days + forecast_days "
            "一次获取最近历史段与未来预报段，再按日期拆分。"
        )
    else:
        historical_units = (
            historical_payload.get(
                "daily_units"
            )
            if isinstance(
                historical_payload.get(
                    "daily_units"
                ),
                dict,
            )
            else {}
        )
        historical_semantics = (
            "Open-Meteo Historical Weather API "
            "提供的历史天气数据。"
        )
        method_note = (
            "单请求模式失败后自动降级："
            "历史段使用 Historical Weather API，"
            "未来段使用 Forecast API。"
        )

    forecast_units = (
        forecast_payload.get(
            "daily_units"
        )
        if isinstance(
            forecast_payload.get(
                "daily_units"
            ),
            dict,
        )
        else {}
    )

    metadata: Dict[str, Any] = {
        "provider": "Open-Meteo",
        "latitude": latitude,
        "longitude": longitude,
        "timezone": timezone,
        "reference_date": (
            ref_date.isoformat()
        ),
        "historical_start_date": (
            historical_start.isoformat()
        ),
        "historical_end_date": (
            historical_end.isoformat()
        ),
        "historical_days_requested": (
            historical_days
        ),
        "historical_rows": int(
            len(historical)
        ),
        "forecast_start_date": (
            forecast["date"]
            .min()
            .date()
            .isoformat()
            if not forecast.empty
            else None
        ),
        "forecast_end_date": (
            forecast["date"]
            .max()
            .date()
            .isoformat()
            if not forecast.empty
            else None
        ),
        "forecast_days_requested": (
            forecast_days
        ),
        "forecast_rows": int(
            len(forecast)
        ),
        "acquisition_mode": (
            acquisition_mode
        ),
        "fallback_reason": (
            fallback_reason or None
        ),
        "historical_endpoint": (
            HISTORICAL_ENDPOINT
            if acquisition_mode
            == "historical_plus_forecast_split"
            else FORECAST_ENDPOINT
        ),
        "forecast_endpoint": (
            FORECAST_ENDPOINT
        ),
        "historical_request_url": (
            historical_url
        ),
        "forecast_request_url": (
            forecast_url
        ),
        "acquired_at": acquired_at,
        "historical_daily_units": (
            historical_units
        ),
        "forecast_daily_units": (
            forecast_units
        ),
        "historical_data_semantics": (
            historical_semantics
        ),
        "method_note": method_note,
    }

    result: Dict[str, Any] = {
        "historical": historical,
        "forecast": forecast,
        "combined": combined,
        "metadata": metadata,
    }

    if output_dir:
        directory = Path(
            str(output_dir)
        ).expanduser()
        directory.mkdir(
            parents=True,
            exist_ok=True,
        )

        historical_path = (
            directory
            / "open_meteo_historical_daily.csv"
        )
        forecast_path = (
            directory
            / "open_meteo_forecast_daily.csv"
        )
        metadata_path = (
            directory
            / "open_meteo_weather_metadata.json"
        )

        historical.to_csv(
            historical_path,
            index=False,
            encoding="utf-8-sig",
        )
        forecast.to_csv(
            forecast_path,
            index=False,
            encoding="utf-8-sig",
        )
        metadata_path.write_text(
            json.dumps(
                metadata,
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )

        result.update(
            {
                "historical_file": str(
                    historical_path.resolve()
                ),
                "forecast_file": str(
                    forecast_path.resolve()
                ),
                "metadata_file": str(
                    metadata_path.resolve()
                ),
            }
        )

    return result

# ============================================================
# v6.6 fix38：Weather deterministic analysis + delivery
# ============================================================

_RISK_ORDER = {
    "低": 0,
    "中": 1,
    "高": 2,
}


def _set_weather_chart_font() -> None:
    plt.rcParams["font.sans-serif"] = [
        "Microsoft YaHei",
        "SimHei",
        "Arial Unicode MS",
        "DejaVu Sans",
    ]
    plt.rcParams["axes.unicode_minus"] = False


def _ensure_weather_dataframe(
    value: Any,
    *,
    name: str,
) -> pd.DataFrame:
    if not isinstance(value, pd.DataFrame):
        raise TypeError(
            f"{name} 必须是 pandas DataFrame，"
            f"实际为 {type(value).__name__}。"
        )

    if value.empty:
        raise ValueError(
            f"{name} 不能为空。"
        )

    frame = value.copy()

    if "date" not in frame.columns:
        raise ValueError(
            f"{name} 缺少 date 字段。"
        )

    frame["date"] = pd.to_datetime(
        frame["date"],
        errors="raise",
    )

    return (
        frame
        .sort_values("date")
        .reset_index(drop=True)
    )


def _weather_missing_dates(
    frame: pd.DataFrame,
) -> list[str]:
    if frame.empty:
        return []

    start = frame["date"].min().normalize()
    end = frame["date"].max().normalize()

    expected = pd.date_range(
        start=start,
        end=end,
        freq="D",
    )
    actual = pd.DatetimeIndex(
        frame["date"].dt.normalize().unique()
    )

    missing = expected.difference(actual)

    return [
        item.date().isoformat()
        for item in missing
    ]


def _weather_abnormal_counts(
    frame: pd.DataFrame,
) -> Dict[str, int]:
    checks: Dict[str, int] = {}

    bounds = {
        "temperature_mean_c": (-80.0, 60.0),
        "temperature_max_c": (-80.0, 60.0),
        "temperature_min_c": (-80.0, 60.0),
        "relative_humidity_mean_pct": (0.0, 100.0),
        "precipitation_sum_mm": (0.0, 1000.0),
        "precipitation_probability_max_pct": (0.0, 100.0),
        "wind_speed_mean_kmh": (0.0, 500.0),
        "wind_speed_max_kmh": (0.0, 500.0),
        "wind_gusts_max_kmh": (0.0, 500.0),
    }

    for column, (low, high) in bounds.items():
        if column not in frame.columns:
            continue

        series = pd.to_numeric(
            frame[column],
            errors="coerce",
        )

        count = int(
            (
                (series < low)
                | (series > high)
            ).sum()
        )

        if count:
            checks[column] = count

    return checks


def _clean_weather_frame(
    frame: pd.DataFrame,
) -> tuple[pd.DataFrame, Dict[str, Any]]:
    cleaned = frame.copy()
    original_rows = int(len(cleaned))

    duplicate_rows = int(
        cleaned.duplicated(
            subset=["date"],
            keep="last",
        ).sum()
    )

    if duplicate_rows:
        cleaned = (
            cleaned
            .drop_duplicates(
                subset=["date"],
                keep="last",
            )
            .copy()
        )

    abnormal_before = _weather_abnormal_counts(
        cleaned
    )

    numeric_interpolation_columns = [
        "temperature_mean_c",
        "temperature_max_c",
        "temperature_min_c",
        "relative_humidity_mean_pct",
        "wind_speed_mean_kmh",
        "wind_speed_max_kmh",
        "wind_gusts_max_kmh",
    ]

    # First convert impossible values to NA.
    for column, count in abnormal_before.items():
        if count <= 0 or column not in cleaned.columns:
            continue

        if column.startswith("temperature_"):
            low, high = -80.0, 60.0
        elif column in {
            "relative_humidity_mean_pct",
            "precipitation_probability_max_pct",
        }:
            low, high = 0.0, 100.0
        elif column == "precipitation_sum_mm":
            low, high = 0.0, 1000.0
        else:
            low, high = 0.0, 500.0

        numeric = pd.to_numeric(
            cleaned[column],
            errors="coerce",
        )
        invalid = (
            (numeric < low)
            | (numeric > high)
        )
        cleaned.loc[
            invalid,
            column,
        ] = pd.NA

    missing_before = {
        column: int(count)
        for column, count in (
            cleaned.isna().sum().to_dict()
        ).items()
        if int(count) > 0
    }

    filled_counts: Dict[str, int] = {}

    # Deterministic interpolation only for continuous atmospheric variables.
    for column in numeric_interpolation_columns:
        if column not in cleaned.columns:
            continue

        before = int(
            cleaned[column].isna().sum()
        )

        if before <= 0:
            continue

        numeric = pd.to_numeric(
            cleaned[column],
            errors="coerce",
        )

        numeric = (
            numeric
            .interpolate(
                method="linear",
                limit_direction="both",
            )
            .ffill()
            .bfill()
        )

        cleaned[column] = numeric

        after = int(
            cleaned[column].isna().sum()
        )
        filled = max(
            0,
            before - after,
        )
        if filled:
            filled_counts[column] = filled

    # Precipitation / probability missingness is not silently fabricated.
    unresolved_missing = {
        column: int(count)
        for column, count in (
            cleaned.isna().sum().to_dict()
        ).items()
        if int(count) > 0
    }

    cleaned = (
        cleaned
        .sort_values("date")
        .reset_index(drop=True)
    )

    result = {
        "original_rows": original_rows,
        "cleaned_rows": int(len(cleaned)),
        "duplicate_rows_removed": duplicate_rows,
        "missing_before": missing_before,
        "filled_by_linear_interpolation": filled_counts,
        "unresolved_missing": unresolved_missing,
        "abnormal_values_replaced_with_missing": abnormal_before,
        "method": (
            "按 date 去重；对温度、湿度、风速等连续变量的缺失/异常值"
            "使用线性插值并以前后有效值补边界；"
            "降水量与降水概率不静默填造。"
        ),
    }

    return cleaned, result


def _risk_level_max(
    values: list[str],
) -> str:
    return max(
        values,
        key=lambda item: _RISK_ORDER.get(
            str(item),
            -1,
        ),
    )


def _build_weather_risk_table(
    forecast: pd.DataFrame,
) -> pd.DataFrame:
    rows = []

    for _, row in forecast.iterrows():
        t_mean = float(
            row.get("temperature_mean_c")
            if pd.notna(
                row.get("temperature_mean_c")
            )
            else 0.0
        )
        t_max = float(
            row.get("temperature_max_c")
            if pd.notna(
                row.get("temperature_max_c")
            )
            else t_mean
        )
        rh = float(
            row.get("relative_humidity_mean_pct")
            if pd.notna(
                row.get("relative_humidity_mean_pct")
            )
            else 0.0
        )
        precip = float(
            row.get("precipitation_sum_mm")
            if pd.notna(
                row.get("precipitation_sum_mm")
            )
            else 0.0
        )
        precip_prob = float(
            row.get(
                "precipitation_probability_max_pct"
            )
            if pd.notna(
                row.get(
                    "precipitation_probability_max_pct"
                )
            )
            else 0.0
        )

        wind_value = row.get(
            "wind_gusts_max_kmh"
        )
        if pd.isna(wind_value):
            wind_value = row.get(
                "wind_speed_max_kmh"
            )
        if pd.isna(wind_value):
            wind_value = row.get(
                "wind_speed_mean_kmh"
            )
        wind = float(
            wind_value
            if pd.notna(wind_value)
            else 0.0
        )

        if t_max >= 35.0:
            high_temp = "高"
        elif t_max >= 33.0:
            high_temp = "中"
        else:
            high_temp = "低"

        if precip >= 50.0:
            heavy_rain = "高"
        elif (
            precip >= 25.0
            or precip_prob >= 70.0
        ):
            heavy_rain = "中"
        else:
            heavy_rain = "低"

        if wind >= 62.0:
            strong_wind = "高"
        elif wind >= 39.0:
            strong_wind = "中"
        else:
            strong_wind = "低"

        if (
            t_mean >= 30.0
            and rh >= 85.0
        ):
            humid_heat = "高"
        elif (
            t_mean >= 28.0
            and rh >= 80.0
        ):
            humid_heat = "中"
        else:
            humid_heat = "低"

        overall = _risk_level_max(
            [
                high_temp,
                heavy_rain,
                strong_wind,
                humid_heat,
            ]
        )

        basis = (
            f"最高温 {t_max:.1f}℃；"
            f"平均湿度 {rh:.0f}%；"
            f"降水 {precip:.1f} mm，"
            f"降水概率 {precip_prob:.0f}%；"
            f"最大风/阵风 {wind:.1f} km/h。"
        )

        rows.append(
            {
                "日期": pd.Timestamp(
                    row["date"]
                ).date().isoformat(),
                "高温风险": high_temp,
                "强降雨风险": heavy_rain,
                "大风风险": strong_wind,
                "高湿闷热风险": humid_heat,
                "综合风险": overall,
                "判断依据": basis,
            }
        )

    return pd.DataFrame(rows)


def analyze_weather_dataset(
    historical: pd.DataFrame,
    forecast: pd.DataFrame,
    metadata: Optional[Dict[str, Any]] = None,
) -> Dict[str, Any]:
    """
    对“近期历史逐日天气 + 未来逐日预报”执行确定性气象分析。

    该 Tool 是 Processing 层能力：
    - 严格分开历史段和未来段，避免把 37 天混成“最近30天”统计；
    - 检查缺失、重复、时间连续性、明显异常值；
    - 记录清洗方法；
    - 计算最近历史段统计；
    - 对未来逐日生成高温/强降雨/大风/高湿闷热风险。
    """
    raw_historical = _ensure_weather_dataframe(
        historical,
        name="historical",
    )
    raw_forecast = _ensure_weather_dataframe(
        forecast,
        name="forecast",
    )

    clean_historical, historical_cleaning = (
        _clean_weather_frame(
            raw_historical
        )
    )
    clean_forecast, forecast_cleaning = (
        _clean_weather_frame(
            raw_forecast
        )
    )

    historical_missing_dates = (
        _weather_missing_dates(
            clean_historical
        )
    )
    forecast_missing_dates = (
        _weather_missing_dates(
            clean_forecast
        )
    )

    historical_abnormal = (
        _weather_abnormal_counts(
            clean_historical
        )
    )
    forecast_abnormal = (
        _weather_abnormal_counts(
            clean_forecast
        )
    )

    quality_rows = [
        {
            "检查项目": "历史原始行数",
            "检查结果": int(
                len(raw_historical)
            ),
        },
        {
            "检查项目": "历史清洗后行数",
            "检查结果": int(
                len(clean_historical)
            ),
        },
        {
            "检查项目": "历史重复日期删除数",
            "检查结果": historical_cleaning[
                "duplicate_rows_removed"
            ],
        },
        {
            "检查项目": "历史缺失值（清洗前）",
            "检查结果": str(
                historical_cleaning[
                    "missing_before"
                ]
            ),
        },
        {
            "检查项目": "历史缺失值（清洗后）",
            "检查结果": str(
                historical_cleaning[
                    "unresolved_missing"
                ]
            ),
        },
        {
            "检查项目": "历史缺失日期",
            "检查结果": (
                ", ".join(
                    historical_missing_dates
                )
                if historical_missing_dates
                else "无"
            ),
        },
        {
            "检查项目": "历史明显异常值",
            "检查结果": (
                str(historical_abnormal)
                if historical_abnormal
                else "无"
            ),
        },
        {
            "检查项目": "未来预报行数",
            "检查结果": int(
                len(clean_forecast)
            ),
        },
        {
            "检查项目": "未来缺失日期",
            "检查结果": (
                ", ".join(
                    forecast_missing_dates
                )
                if forecast_missing_dates
                else "无"
            ),
        },
        {
            "检查项目": "清洗方法",
            "检查结果": historical_cleaning[
                "method"
            ],
        },
    ]

    quality_report = pd.DataFrame(
        quality_rows
    )

    hist = clean_historical.copy()

    if "temperature_mean_c" in hist.columns:
        mean_temp = float(
            hist[
                "temperature_mean_c"
            ].mean()
        )
    else:
        mean_temp = float("nan")

    max_temp = float(
        hist[
            "temperature_max_c"
        ].max()
    )
    min_temp = float(
        hist[
            "temperature_min_c"
        ].min()
    )

    daily_range = (
        pd.to_numeric(
            hist["temperature_max_c"],
            errors="coerce",
        )
        - pd.to_numeric(
            hist["temperature_min_c"],
            errors="coerce",
        )
    )

    mean_humidity = float(
        hist[
            "relative_humidity_mean_pct"
        ].mean()
    )
    precip_total = float(
        hist[
            "precipitation_sum_mm"
        ].sum()
    )
    precip_days = int(
        (
            pd.to_numeric(
                hist[
                    "precipitation_sum_mm"
                ],
                errors="coerce",
            )
            > 0
        ).sum()
    )
    mean_wind = float(
        hist[
            "wind_speed_mean_kmh"
        ].mean()
    )

    wind_candidates = []
    for column in (
        "wind_gusts_max_kmh",
        "wind_speed_max_kmh",
        "wind_speed_mean_kmh",
    ):
        if column in hist.columns:
            wind_candidates.append(
                float(
                    pd.to_numeric(
                        hist[column],
                        errors="coerce",
                    ).max()
                )
            )
    max_wind = max(
        wind_candidates
        or [float("nan")]
    )

    trend_description = (
        "样本不足，无法判断明显趋势。"
    )
    trend_slope = float("nan")

    if (
        len(hist) >= 3
        and "temperature_mean_c"
        in hist.columns
    ):
        values = pd.to_numeric(
            hist["temperature_mean_c"],
            errors="coerce",
        ).dropna()

        if len(values) >= 3:
            x = pd.Series(
                range(len(values)),
                dtype="float64",
            )
            y = values.reset_index(
                drop=True
            )
            x_mean = float(x.mean())
            y_mean = float(y.mean())
            denominator = float(
                (
                    (x - x_mean) ** 2
                ).sum()
            )
            if denominator > 0:
                trend_slope = float(
                    (
                        (
                            (x - x_mean)
                            * (y - y_mean)
                        ).sum()
                    )
                    / denominator
                )

                if trend_slope > 0.08:
                    trend_description = (
                        "最近历史段存在较明显升温趋势。"
                    )
                elif trend_slope < -0.08:
                    trend_description = (
                        "最近历史段存在较明显降温趋势。"
                    )
                else:
                    trend_description = (
                        "最近历史段平均气温未出现明显单向升降温趋势。"
                    )

    precipitation_description = (
        "累计降水不足以支持集中性判断。"
    )
    if precip_total > 0:
        top3 = (
            pd.to_numeric(
                hist[
                    "precipitation_sum_mm"
                ],
                errors="coerce",
            )
            .fillna(0.0)
            .nlargest(
                min(3, len(hist))
            )
            .sum()
        )
        share = float(
            top3 / precip_total
        )
        if share >= 0.5:
            precipitation_description = (
                "降水较集中：降水量最大的3天贡献了"
                f"约 {share * 100:.1f}% 的累计降水。"
            )
        else:
            precipitation_description = (
                "降水分布相对分散，最大3个降水日占比"
                f"约 {share * 100:.1f}%。"
            )

    statistics_report = pd.DataFrame(
        [
            {
                "指标": "平均气温",
                "数值": round(
                    mean_temp,
                    2,
                ),
                "单位": "℃",
            },
            {
                "指标": "最高气温",
                "数值": round(
                    max_temp,
                    2,
                ),
                "单位": "℃",
            },
            {
                "指标": "最低气温",
                "数值": round(
                    min_temp,
                    2,
                ),
                "单位": "℃",
            },
            {
                "指标": "平均日温差",
                "数值": round(
                    float(
                        daily_range.mean()
                    ),
                    2,
                ),
                "单位": "℃",
            },
            {
                "指标": "平均相对湿度",
                "数值": round(
                    mean_humidity,
                    2,
                ),
                "单位": "%",
            },
            {
                "指标": "累计降水量",
                "数值": round(
                    precip_total,
                    2,
                ),
                "单位": "mm",
            },
            {
                "指标": "降水日数",
                "数值": precip_days,
                "单位": "天",
            },
            {
                "指标": "平均风速",
                "数值": round(
                    mean_wind,
                    2,
                ),
                "单位": "km/h",
            },
            {
                "指标": "最大风速/阵风",
                "数值": round(
                    max_wind,
                    2,
                ),
                "单位": "km/h",
            },
        ]
    )

    risk_report = (
        _build_weather_risk_table(
            clean_forecast
        )
    )

    forecast_report = (
        clean_forecast.copy()
    )

    risk_lookup = (
        risk_report
        .set_index("日期")
        .to_dict(
            orient="index"
        )
    )

    forecast_report[
        "风险等级"
    ] = forecast_report["date"].map(
        lambda value: risk_lookup.get(
            pd.Timestamp(
                value
            ).date().isoformat(),
            {},
        ).get(
            "综合风险",
            "低",
        )
    )
    forecast_report[
        "备注"
    ] = forecast_report["date"].map(
        lambda value: risk_lookup.get(
            pd.Timestamp(
                value
            ).date().isoformat(),
            {},
        ).get(
            "判断依据",
            "",
        )
    )

    raw_combined = pd.concat(
        [
            raw_historical,
            raw_forecast,
        ],
        ignore_index=True,
        sort=False,
    )

    cleaned_combined = pd.concat(
        [
            clean_historical,
            clean_forecast,
        ],
        ignore_index=True,
        sort=False,
    )

    high_risk_days = int(
        (
            risk_report[
                "综合风险"
            ]
            == "高"
        ).sum()
    )
    medium_risk_days = int(
        (
            risk_report[
                "综合风险"
            ]
            == "中"
        ).sum()
    )

    if high_risk_days:
        risk_summary = (
            f"未来7天有 {high_risk_days} 天综合风险为高，"
            "应优先关注对应日期。"
        )
    elif medium_risk_days:
        risk_summary = (
            f"未来7天有 {medium_risk_days} 天综合风险为中，"
            "主要关注降雨、大风或高湿闷热条件。"
        )
    else:
        risk_summary = (
            "未来7天综合风险均为低，但仍需结合临近预报滚动更新。"
        )

    summary = {
        "historical_rows": int(
            len(clean_historical)
        ),
        "forecast_rows": int(
            len(clean_forecast)
        ),
        "mean_temperature_c": round(
            mean_temp,
            2,
        ),
        "max_temperature_c": round(
            max_temp,
            2,
        ),
        "min_temperature_c": round(
            min_temp,
            2,
        ),
        "precipitation_total_mm": round(
            precip_total,
            2,
        ),
        "precipitation_days": precip_days,
        "temperature_trend_slope_c_per_day": (
            None
            if pd.isna(
                trend_slope
            )
            else round(
                trend_slope,
                4,
            )
        ),
        "temperature_trend_description": (
            trend_description
        ),
        "precipitation_distribution_description": (
            precipitation_description
        ),
        "forecast_risk_summary": risk_summary,
        "cleaning_historical": (
            historical_cleaning
        ),
        "cleaning_forecast": (
            forecast_cleaning
        ),
    }

    return {
        # 专业阶段完成信号由 Python 产生，AgentLoop 可确定性推进到 Delivery。
        "processing_complete": True,
        # 'dataframe' is intentionally present so AgentLoop DataState
        # tracks this result as a real Processing DataFrame.
        "dataframe": cleaned_combined,
        "raw_data": raw_combined,
        "cleaned_data": cleaned_combined,
        "historical_cleaned": clean_historical,
        "forecast_cleaned": clean_forecast,
        "quality_report": quality_report,
        "statistics_report": statistics_report,
        "forecast_report": forecast_report,
        "risk_report": risk_report,
        "summary": summary,
        "metadata": dict(
            metadata or {}
        ),
    }


def _write_dataframe_table_to_word(
    document: Document,
    frame: pd.DataFrame,
    *,
    max_rows: int = 12,
) -> None:
    if frame.empty:
        paragraph = document.add_paragraph(
            "无可展示数据。"
        )
        paragraph.paragraph_format.space_after = Pt(4)
        return

    preview = frame.head(
        max_rows
    ).copy()

    columns = [
        str(column)
        for column in preview.columns
    ]

    table = document.add_table(
        rows=1,
        cols=len(columns),
    )
    table.style = "Table Grid"

    for index, column in enumerate(columns):
        table.rows[0].cells[
            index
        ].text = column

    for _, record in preview.iterrows():
        cells = table.add_row().cells

        for index, column in enumerate(columns):
            value = record[column]

            if isinstance(
                value,
                pd.Timestamp,
            ):
                text = (
                    value
                    .date()
                    .isoformat()
                )
            elif pd.isna(value):
                text = ""
            else:
                text = str(value)

            cells[index].text = text


WEATHER_DISPLAY_COLUMN_NAMES = {
    "date": "日期",
    "data_kind": "数据类型",
    "temperature_mean_c": "平均气温（℃）",
    "temperature_max_c": "最高气温（℃）",
    "temperature_min_c": "最低气温（℃）",
    "relative_humidity_mean_pct": "平均相对湿度（%）",
    "precipitation_sum_mm": "降水量（mm）",
    "precipitation_probability_max_pct": "最大降水概率（%）",
    "wind_speed_mean_kmh": "平均风速（km/h）",
    "wind_speed_max_kmh": "最大风速（km/h）",
    "wind_gusts_max_kmh": "最大阵风（km/h）",
}


def _weather_display_frame(
    frame: pd.DataFrame,
) -> pd.DataFrame:
    return frame.rename(
        columns={
            column: WEATHER_DISPLAY_COLUMN_NAMES.get(
                str(column),
                str(column),
            )
            for column in frame.columns
        }
    )


def _save_weather_chart(
    *,
    frame: pd.DataFrame,
    x_column: str,
    y_columns: list[str],
    title: str,
    ylabel: str,
    output_path: Path,
    kind: str = "line",
) -> str:
    _set_weather_chart_font()

    data = frame.copy()
    data[x_column] = pd.to_datetime(
        data[x_column],
        errors="coerce",
    )
    data = (
        data
        .dropna(
            subset=[x_column]
        )
        .sort_values(
            x_column
        )
    )

    plt.figure(
        figsize=(10, 5.5)
    )

    if kind == "bar":
        if len(y_columns) != 1:
            raise ValueError(
                "bar 图当前要求仅提供一个 y_column。"
            )

        plt.bar(
            data[x_column],
            pd.to_numeric(
                data[
                    y_columns[0]
                ],
                errors="coerce",
            ),
        )
    else:
        for column in y_columns:
            plt.plot(
                data[x_column],
                pd.to_numeric(
                    data[column],
                    errors="coerce",
                ),
                marker="o",
                markersize=3,
                label=WEATHER_DISPLAY_COLUMN_NAMES.get(
                    str(column),
                    str(column),
                ),
            )

        if len(y_columns) > 1:
            plt.legend()

    plt.title(title)
    plt.xlabel("日期")
    plt.ylabel(ylabel)
    plt.grid(
        True,
        alpha=0.25,
    )
    plt.xticks(
        rotation=35,
        ha="right",
    )
    plt.tight_layout()

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    plt.savefig(
        output_path,
        dpi=160,
        bbox_inches="tight",
    )
    plt.close()

    return str(
        output_path.resolve()
    )


def _normalize_weather_compare_value(
    value: Any,
) -> Any:
    if value is None:
        return None

    if isinstance(
        value,
        pd.Timestamp,
    ):
        return value.date().isoformat()

    try:
        if pd.isna(value):
            return None
    except Exception:
        pass

    if isinstance(
        value,
        (int, float),
    ) and not isinstance(
        value,
        bool,
    ):
        return round(
            float(value),
            6,
        )

    text = str(value).strip()

    # Date-like values exported/read back by Excel.
    try:
        parsed = pd.to_datetime(
            text,
            errors="raise",
        )
        if (
            len(text) >= 8
            and any(
                separator in text
                for separator in (
                    "-",
                    "/",
                )
            )
        ):
            return parsed.date().isoformat()
    except Exception:
        pass

    # Numeric strings from Word tables.
    try:
        return round(
            float(text),
            6,
        )
    except Exception:
        return text


def _compare_weather_tables_by_columns(
    expected: pd.DataFrame,
    actual: pd.DataFrame,
    *,
    columns: list[str],
) -> tuple[bool, list[str]]:
    missing_columns = [
        column
        for column in columns
        if (
            column not in expected.columns
            or column not in actual.columns
        )
    ]

    if missing_columns:
        return (
            False,
            [
                "缺少字段："
                + ", ".join(
                    missing_columns
                )
            ],
        )

    if len(expected) != len(actual):
        return (
            False,
            [
                (
                    "行数不一致："
                    f"expected={len(expected)}, "
                    f"actual={len(actual)}"
                )
            ],
        )

    mismatches: list[str] = []

    for row_index in range(
        len(expected)
    ):
        for column in columns:
            left = (
                _normalize_weather_compare_value(
                    expected.iloc[
                        row_index
                    ][column]
                )
            )
            right = (
                _normalize_weather_compare_value(
                    actual.iloc[
                        row_index
                    ][column]
                )
            )

            if left == right:
                continue

            if (
                isinstance(
                    left,
                    float,
                )
                and isinstance(
                    right,
                    float,
                )
                and abs(
                    left - right
                )
                <= 1e-6
            ):
                continue

            mismatches.append(
                (
                    f"row={row_index}, "
                    f"column={column}, "
                    f"expected={left!r}, "
                    f"actual={right!r}"
                )
            )

            if len(mismatches) >= 12:
                return (
                    False,
                    mismatches,
                )

    return (
        not mismatches,
        mismatches,
    )


def _extract_word_forecast_table(
    document: Document,
) -> pd.DataFrame:
    required_headers = {
        "日期",
        "平均气温（℃）",
        "最高气温（℃）",
        "最低气温（℃）",
    }

    for table in document.tables:
        if not table.rows:
            continue

        headers = [
            cell.text.strip()
            for cell in (
                table.rows[0].cells
            )
        ]

        if not required_headers.issubset(
            set(headers)
        ):
            continue

        rows = []

        for word_row in table.rows[1:]:
            values = [
                cell.text.strip()
                for cell in (
                    word_row.cells
                )
            ]

            if not any(values):
                continue

            rows.append(
                dict(
                    zip(
                        headers,
                        values,
                    )
                )
            )

        return pd.DataFrame(
            rows
        )

    return pd.DataFrame()


def _verify_weather_cross_deliverable_consistency(
    *,
    excel_path: Path,
    verified_word: Document,
    frames: Dict[str, pd.DataFrame],
    summary: Dict[str, Any],
    chart_paths: list[str],
    png_readable_count: int,
) -> Dict[str, Any]:
    """
    对最终 Excel / Word / PNG 与同一份 analysis_result 做确定性一致性验证。

    Excel:
    - 重新读取“30天统计分析”；
    - 重新读取“未来7天预测”；
    - 与内存中的确定性分析结果逐字段比较。

    Word:
    - 重新读取未来7天预测表；
    - 与同一份 forecast_report 比较；
    - 检查正文关键统计数字是否与 summary 一致。

    PNG:
    - 图像文件必须真实可读；
    - 四张图均由当前函数直接使用 historical_cleaned /
      forecast_cleaned 生成，因此记录生成源行数作为 provenance 证据。
    """
    stats_actual = pd.read_excel(
        excel_path,
        sheet_name="30天统计分析",
    )
    forecast_actual = pd.read_excel(
        excel_path,
        sheet_name="未来7天预测",
    )

    stats_expected = (
        frames[
            "statistics_report"
        ]
        .copy()
        .reset_index(
            drop=True
        )
    )
    forecast_expected = (
        _weather_display_frame(
            frames[
                "forecast_report"
            ]
        )
        .copy()
        .reset_index(
            drop=True
        )
    )

    stats_columns = [
        "指标",
        "数值",
        "单位",
    ]

    forecast_columns = [
        column
        for column in (
            "日期",
            "平均气温（℃）",
            "最高气温（℃）",
            "最低气温（℃）",
            "平均相对湿度（%）",
            "降水量（mm）",
            "最大降水概率（%）",
            "平均风速（km/h）",
            "风险等级",
        )
        if column in forecast_expected.columns
    ]

    (
        excel_stats_match,
        excel_stats_mismatches,
    ) = _compare_weather_tables_by_columns(
        stats_expected,
        stats_actual,
        columns=stats_columns,
    )

    (
        excel_forecast_match,
        excel_forecast_mismatches,
    ) = _compare_weather_tables_by_columns(
        forecast_expected,
        forecast_actual,
        columns=forecast_columns,
    )

    word_forecast = (
        _extract_word_forecast_table(
            verified_word
        )
    )

    (
        word_forecast_match,
        word_forecast_mismatches,
    ) = _compare_weather_tables_by_columns(
        forecast_expected,
        word_forecast,
        columns=forecast_columns,
    )

    word_text = "\n".join(
        [
            paragraph.text
            for paragraph
            in verified_word.paragraphs
        ]
    )

    summary_phrases = [
        (
            "平均气温 "
            f"{summary.get('mean_temperature_c')}℃"
        ),
        (
            "最高气温 "
            f"{summary.get('max_temperature_c')}℃"
        ),
        (
            "最低气温 "
            f"{summary.get('min_temperature_c')}℃"
        ),
        (
            "累计降水量 "
            f"{summary.get('precipitation_total_mm')} mm"
        ),
        (
            "降水日数 "
            f"{summary.get('precipitation_days')} 天"
        ),
    ]

    missing_word_summary_phrases = [
        phrase
        for phrase in summary_phrases
        if phrase not in word_text
    ]

    word_summary_match = (
        len(
            missing_word_summary_phrases
        )
        == 0
    )

    chart_source_rows = {
        "01_最近30天气温变化.png": int(
            len(
                frames[
                    "historical_cleaned"
                ]
            )
        ),
        "02_最近30天逐日降水量.png": int(
            len(
                frames[
                    "historical_cleaned"
                ]
            )
        ),
        "03_未来7天气温预报.png": int(
            len(
                frames[
                    "forecast_cleaned"
                ]
            )
        ),
        "04_未来7天降水预报.png": int(
            len(
                frames[
                    "forecast_cleaned"
                ]
            )
        ),
    }

    png_source_consistency = bool(
        len(chart_paths) >= 4
        and png_readable_count >= 4
        and list(
            chart_source_rows.values()
        ) == [
            30,
            30,
            7,
            7,
        ]
    )

    evidence = [
        (
            "Excel 30天统计分析 "
            + (
                "与 analysis_result 一致"
                if excel_stats_match
                else "存在不一致"
            )
        ),
        (
            "Excel 未来7天预测 "
            + (
                "与 analysis_result 一致"
                if excel_forecast_match
                else "存在不一致"
            )
        ),
        (
            "Word 未来7天预测表 "
            + (
                "与 analysis_result 一致"
                if word_forecast_match
                else "存在不一致"
            )
        ),
        (
            "Word 关键统计正文 "
            + (
                "与 analysis_result 一致"
                if word_summary_match
                else "存在不一致"
            )
        ),
        (
            "PNG 均由同一 analysis_result 的 "
            "30/30/7/7 行源数据生成且重新读取成功"
            if png_source_consistency
            else "PNG 源数据或可读性验证未通过"
        ),
    ]

    mismatches = (
        excel_stats_mismatches
        + excel_forecast_mismatches
        + word_forecast_mismatches
        + [
            (
                "Word 缺少关键统计短语："
                + phrase
            )
            for phrase
            in missing_word_summary_phrases
        ]
    )

    passed = bool(
        excel_stats_match
        and excel_forecast_match
        and word_forecast_match
        and word_summary_match
        and png_source_consistency
    )

    return {
        "passed": passed,
        "excel_statistics_match": (
            excel_stats_match
        ),
        "excel_forecast_match": (
            excel_forecast_match
        ),
        "word_forecast_match": (
            word_forecast_match
        ),
        "word_summary_match": (
            word_summary_match
        ),
        "png_source_consistency": (
            png_source_consistency
        ),
        "chart_source_rows": (
            chart_source_rows
        ),
        "evidence": evidence,
        "mismatches": mismatches[
            :20
        ],
    }


def create_weather_analysis_package(
    analysis_result: Dict[str, Any],
    output_dir: str,
    excel_filename: str = "珠海气象分析与7天天气预测.xlsx",
    word_filename: str = "珠海气象分析与7天天气预测报告.docx",
    report_title: str = "珠海近期气象数据分析与未来7天天气预测",
) -> Dict[str, Any]:
    """
    将 analyze_weather_dataset 的真实输出确定性生成：
    - 4 张独立 PNG；
    - 6 Sheet Excel；
    - 插入 4 张 PNG 的 Word 报告；
    - 写后重新读取并返回基本验证证据。

    该 Tool 是 Delivery 层能力，避免本地小模型把大量表格/正文直接
    序列化进一次超长 JSON 工具调用。
    """
    if not isinstance(
        analysis_result,
        dict,
    ):
        raise TypeError(
            "analysis_result 必须是 analyze_weather_dataset 返回的字典。"
        )

    required_frames = {
        "raw_data": "原始气象数据",
        "cleaned_data": "清洗后数据",
        "quality_report": "数据质量报告",
        "statistics_report": "30天统计分析",
        "forecast_report": "未来7天预测",
        "risk_report": "未来7天天气风险",
        "historical_cleaned": "历史清洗数据",
        "forecast_cleaned": "未来预报数据",
    }

    frames: Dict[str, pd.DataFrame] = {}

    for key, description in required_frames.items():
        value = analysis_result.get(
            key
        )
        if not isinstance(
            value,
            pd.DataFrame,
        ):
            raise TypeError(
                f"analysis_result.{key} 缺失或不是 DataFrame（{description}）。"
            )
        frames[key] = value.copy()

    summary = analysis_result.get(
        "summary"
    )
    if not isinstance(
        summary,
        dict,
    ):
        summary = {}

    metadata = analysis_result.get(
        "metadata"
    )
    if not isinstance(
        metadata,
        dict,
    ):
        metadata = {}

    target_dir = Path(
        str(output_dir)
    ).expanduser()
    target_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    historical = frames[
        "historical_cleaned"
    ]
    forecast = frames[
        "forecast_cleaned"
    ]

    chart_paths = []

    chart_paths.append(
        _save_weather_chart(
            frame=historical,
            x_column="date",
            y_columns=[
                "temperature_mean_c",
                "temperature_max_c",
                "temperature_min_c",
            ],
            title="最近30天气温变化",
            ylabel="气温（℃）",
            output_path=(
                target_dir
                / "01_最近30天气温变化.png"
            ),
            kind="line",
        )
    )

    chart_paths.append(
        _save_weather_chart(
            frame=historical,
            x_column="date",
            y_columns=[
                "precipitation_sum_mm",
            ],
            title="最近30天逐日降水量",
            ylabel="降水量（mm）",
            output_path=(
                target_dir
                / "02_最近30天逐日降水量.png"
            ),
            kind="bar",
        )
    )

    chart_paths.append(
        _save_weather_chart(
            frame=forecast,
            x_column="date",
            y_columns=[
                "temperature_mean_c",
                "temperature_max_c",
                "temperature_min_c",
            ],
            title="未来7天气温预报",
            ylabel="气温（℃）",
            output_path=(
                target_dir
                / "03_未来7天气温预报.png"
            ),
            kind="line",
        )
    )

    chart_paths.append(
        _save_weather_chart(
            frame=forecast,
            x_column="date",
            y_columns=[
                "precipitation_sum_mm",
            ],
            title="未来7天降水预报",
            ylabel="降水量（mm）",
            output_path=(
                target_dir
                / "04_未来7天降水预报.png"
            ),
            kind="bar",
        )
    )

    excel_path = (
        target_dir
        / excel_filename
    ).resolve()

    sheets = {
        "原始气象数据": _weather_display_frame(
            frames[
                "raw_data"
            ]
        ),
        "清洗后数据": _weather_display_frame(
            frames[
                "cleaned_data"
            ]
        ),
        "数据质量报告": frames[
            "quality_report"
        ],
        "30天统计分析": frames[
            "statistics_report"
        ],
        "未来7天预测": _weather_display_frame(
            frames[
                "forecast_report"
            ]
        ),
        "未来7天天气风险": frames[
            "risk_report"
        ],
    }

    with pd.ExcelWriter(
        excel_path,
        engine="openpyxl",
    ) as writer:
        for sheet_name, frame in sheets.items():
            export_frame = frame.copy()

            for column in export_frame.columns:
                if pd.api.types.is_datetime64_any_dtype(
                    export_frame[column]
                ):
                    export_frame[
                        column
                    ] = export_frame[
                        column
                    ].dt.strftime(
                        "%Y-%m-%d"
                    )

            export_frame.to_excel(
                writer,
                sheet_name=sheet_name,
                index=False,
            )

    word_path = (
        target_dir
        / word_filename
    ).resolve()

    document = Document()

    styles = document.styles
    styles["Normal"].font.name = (
        "Microsoft YaHei"
    )
    styles["Normal"].font.size = Pt(
        10.5
    )

    title = document.add_heading(
        report_title,
        level=0,
    )
    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    source_name = str(
        metadata.get("provider")
        or "Open-Meteo"
    )
    hist_start = str(
        metadata.get(
            "historical_start_date"
        )
        or ""
    )
    hist_end = str(
        metadata.get(
            "historical_end_date"
        )
        or ""
    )
    forecast_start = str(
        metadata.get(
            "forecast_start_date"
        )
        or ""
    )
    forecast_end = str(
        metadata.get(
            "forecast_end_date"
        )
        or ""
    )

    document.add_heading(
        "任务说明",
        level=1,
    )
    document.add_paragraph(
        "本报告基于公开气象数据，对近期逐日天气进行质量检查、"
        "清洗与统计，并使用真实天气预报数据描述未来7天趋势与风险。"
    )

    document.add_heading(
        "数据来源",
        level=1,
    )
    document.add_paragraph(
        f"数据源：{source_name}。"
        f"历史时间范围：{hist_start} 至 {hist_end}；"
        f"未来预报范围：{forecast_start} 至 {forecast_end}。"
        f"获取时间：{metadata.get('acquired_at', '')}。"
    )
    if metadata.get(
        "historical_data_semantics"
    ):
        document.add_paragraph(
            str(
                metadata[
                    "historical_data_semantics"
                ]
            )
        )

    document.add_heading(
        "数据质量与清洗过程",
        level=1,
    )
    _write_dataframe_table_to_word(
        document,
        frames[
            "quality_report"
        ],
        max_rows=20,
    )

    document.add_heading(
        "最近30天气象特征",
        level=1,
    )
    document.add_paragraph(
        f"历史段有效记录 {summary.get('historical_rows', len(historical))} 条。"
        f"平均气温 {summary.get('mean_temperature_c', '')}℃，"
        f"最高气温 {summary.get('max_temperature_c', '')}℃，"
        f"最低气温 {summary.get('min_temperature_c', '')}℃；"
        f"累计降水量 {summary.get('precipitation_total_mm', '')} mm，"
        f"降水日数 {summary.get('precipitation_days', '')} 天。"
    )
    document.add_paragraph(
        str(
            summary.get(
                "temperature_trend_description",
                "",
            )
        )
    )
    document.add_paragraph(
        str(
            summary.get(
                "precipitation_distribution_description",
                "",
            )
        )
    )

    document.add_heading(
        "未来7天天气预测",
        level=1,
    )
    forecast_display_columns = [
        column
        for column in [
            "date",
            "temperature_mean_c",
            "temperature_max_c",
            "temperature_min_c",
            "relative_humidity_mean_pct",
            "precipitation_sum_mm",
            "precipitation_probability_max_pct",
            "wind_speed_mean_kmh",
            "风险等级",
            "备注",
        ]
        if column in frames[
            "forecast_report"
        ].columns
    ]
    forecast_word_frame = _weather_display_frame(
        frames[
            "forecast_report"
        ][
            forecast_display_columns
        ]
    )

    _write_dataframe_table_to_word(
        document,
        forecast_word_frame,
        max_rows=7,
    )

    document.add_heading(
        "天气风险分析",
        level=1,
    )
    document.add_paragraph(
        str(
            summary.get(
                "forecast_risk_summary",
                "",
            )
        )
    )
    _write_dataframe_table_to_word(
        document,
        frames[
            "risk_report"
        ],
        max_rows=7,
    )

    document.add_heading(
        "预测方法与假设",
        level=1,
    )
    document.add_paragraph(
        "未来7天预测优先采用数据源直接提供的天气预报值，"
        "本报告没有把统计外推结果冒充为确定天气事实。"
        "风险等级由固定阈值规则从预报温度、湿度、降水和风速中确定性计算。"
    )

    document.add_heading(
        "预测不确定性",
        level=1,
    )
    document.add_paragraph(
        "天气预报随预报时效增加而不确定性上升；"
        "降水落区、短时强降水和阵风通常比大尺度温度趋势更难精确预测。"
        "业务使用时应结合临近预报和最新官方预警滚动更新。"
    )

    document.add_heading(
        "主要结论",
        level=1,
    )
    document.add_paragraph(
        str(
            summary.get(
                "forecast_risk_summary",
                "",
            )
        )
    )

    document.add_heading(
        "主要图表",
        level=1,
    )

    for index, chart_path in enumerate(
        chart_paths,
        start=1,
    ):
        paragraph = document.add_paragraph()
        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )
        run = paragraph.add_run()
        run.add_picture(
            chart_path,
            width=Cm(16.0),
        )

        caption = document.add_paragraph(
            f"图 {index}"
        )
        caption.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

    document.save(
        str(word_path)
    )

    # Write-back verification performed inside the deterministic delivery tool.
    excel_book = pd.ExcelFile(
        excel_path
    )
    verified_sheet_names = list(
        excel_book.sheet_names
    )

    verified_word = Document(
        str(word_path)
    )
    word_image_count = int(
        len(
            verified_word.inline_shapes
        )
    )

    missing_files = [
        path
        for path in (
            [
                str(excel_path),
                str(word_path),
            ]
            + list(chart_paths)
        )
        if not Path(path).is_file()
    ]

    png_readable_count = 0
    for chart_path in chart_paths:
        try:
            image_data = plt.imread(
                chart_path
            )
            if (
                getattr(
                    image_data,
                    "size",
                    0,
                )
                > 0
            ):
                png_readable_count += 1
        except Exception:
            pass

    cross_deliverable_consistency = (
        _verify_weather_cross_deliverable_consistency(
            excel_path=excel_path,
            verified_word=verified_word,
            frames=frames,
            summary=summary,
            chart_paths=chart_paths,
            png_readable_count=png_readable_count,
        )
    )

    verification = {
        "cross_deliverable_consistency": (
            cross_deliverable_consistency
        ),
        "all_files_exist": (
            len(missing_files) == 0
        ),
        "missing_files": missing_files,
        "excel_sheet_names": (
            verified_sheet_names
        ),
        "excel_sheet_count": int(
            len(
                verified_sheet_names
            )
        ),
        "expected_excel_sheet_count": 6,
        "png_count": int(
            len(chart_paths)
        ),
        "png_readable_count": int(
            png_readable_count
        ),
        "word_readable": True,
        "word_image_count": (
            word_image_count
        ),
        "forecast_rows": int(
            len(
                frames[
                    "forecast_report"
                ]
            )
        ),
        "risk_rows": int(
            len(
                frames[
                    "risk_report"
                ]
            )
        ),
    }

    package_success = bool(
        verification[
            "all_files_exist"
        ]
        and verification[
            "excel_sheet_count"
        ] >= 6
        and verification[
            "png_count"
        ] >= 4
        and verification[
            "png_readable_count"
        ] >= 4
        and verification[
            "word_image_count"
        ] >= 4
        and verification[
            "forecast_rows"
        ] == 7
        and verification[
            "risk_rows"
        ] == 7
        and verification[
            "cross_deliverable_consistency"
        ][
            "passed"
        ]
        is True
    )

    return {
        "success": package_success,
        # 专业阶段完成信号只有在内部写后验证全部通过时才为 True。
        "delivery_complete": package_success,
        "excel_path": str(
            excel_path
        ),
        "word_path": str(
            word_path
        ),
        "chart_paths": list(
            chart_paths
        ),
        "deliverable_paths": (
            [
                str(excel_path),
                str(word_path),
            ]
            + list(chart_paths)
        ),
        "verification": verification,
        "summary": summary,
        "metadata": metadata,
    }

