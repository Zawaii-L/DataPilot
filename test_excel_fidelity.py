from __future__ import annotations

import shutil
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from excel_edit_tools import apply_excel_edits


TEST_ROOT = Path("outputs") / "excel_fidelity_test"
SOURCE_EXCEL = TEST_ROOT / "excel_fidelity_source.xlsx"
OUTPUT_EXCEL = TEST_ROOT / "excel_fidelity_output.xlsx"


def reset_test_directory() -> None:
    if TEST_ROOT.exists():
        shutil.rmtree(TEST_ROOT)
    TEST_ROOT.mkdir(parents=True, exist_ok=True)


def create_source_excel() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "销售数据"

    headers = ["城市", "销售额", "奖金", "官网", "备注"]
    sheet.append(headers)
    sheet.append(["珠海", 128, "=B2*10%", "https://example.com/zhuhai", "重点客户"])
    sheet.append(["澳门", 186, "=B3*10%", "https://example.com/macau", "冠军"])
    sheet.append(["横琴", 154, "=B4*10%", "https://example.com/hengqin", "增长"])

    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    header_font = Font(bold=True, name="微软雅黑", size=11)
    thin = Side(style="thin", color="808080")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal="center")

    row_fills = {
        2: "FFF2CC",
        3: "E2F0D9",
        4: "FCE4D6",
    }

    for row_index in range(2, 5):
        for column_index in range(1, 6):
            cell = sheet.cell(row=row_index, column=column_index)
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor=row_fills[row_index],
            )
            cell.border = border

        sheet.cell(row=row_index, column=2).number_format = '0.0" 万元"'
        sheet.cell(row=row_index, column=3).number_format = '0.0" 万元"'

        link_cell = sheet.cell(row=row_index, column=4)
        link_cell.hyperlink = link_cell.value
        link_cell.style = "Hyperlink"

        note_cell = sheet.cell(row=row_index, column=5)
        note_cell.comment = Comment(
            f"{sheet.cell(row=row_index, column=1).value}备注",
            "DataPilot",
        )

    sheet.column_dimensions["A"].width = 16
    sheet.column_dimensions["B"].width = 14
    sheet.column_dimensions["C"].width = 14
    sheet.column_dimensions["D"].width = 34
    sheet.column_dimensions["E"].width = 18

    sheet.row_dimensions[1].height = 24
    sheet.row_dimensions[2].height = 30
    sheet.row_dimensions[3].height = 36
    sheet.row_dimensions[4].height = 42

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:E4"

    note_sheet = workbook.create_sheet("说明")
    note_sheet["A1"] = "此 Sheet 不应被编辑。"
    note_sheet["A1"].font = Font(bold=True)
    note_sheet.column_dimensions["A"].width = 28

    workbook.save(SOURCE_EXCEL)


def fill_color(cell) -> str | None:
    color = cell.fill.fgColor
    if color is None:
        return None
    return color.rgb or color.indexed or color.theme


def capture_row(sheet, row_index: int) -> dict:
    city = sheet.cell(row=row_index, column=1).value
    return {
        "city": city,
        "sales": sheet.cell(row=row_index, column=2).value,
        "bonus": sheet.cell(row=row_index, column=3).value,
        "sales_number_format": sheet.cell(row=row_index, column=2).number_format,
        "bonus_number_format": sheet.cell(row=row_index, column=3).number_format,
        "fill": fill_color(sheet.cell(row=row_index, column=1)),
        "hyperlink": (
            sheet.cell(row=row_index, column=4).hyperlink.target
            if sheet.cell(row=row_index, column=4).hyperlink
            else None
        ),
        "comment": (
            sheet.cell(row=row_index, column=5).comment.text
            if sheet.cell(row=row_index, column=5).comment
            else None
        ),
        "height": sheet.row_dimensions[row_index].height,
    }


def main() -> None:
    print("=" * 70)
    print("DataPilot v3.4 Excel 高保真编辑基准测试")
    print("=" * 70)

    reset_test_directory()
    create_source_excel()

    source = load_workbook(SOURCE_EXCEL, data_only=False)
    source_sheet = source["销售数据"]

    source_by_city = {
        source_sheet.cell(row=row, column=1).value: capture_row(source_sheet, row)
        for row in range(2, 5)
    }

    source_widths = {
        column: source_sheet.column_dimensions[column].width
        for column in ("A", "B", "C", "D", "E")
    }
    source_freeze = source_sheet.freeze_panes
    source_filter = source_sheet.auto_filter.ref
    source_note = source["说明"]["A1"].value

    print()
    print("[测试 1] 排序后业务行属性跟随数据")

    result = apply_excel_edits(
        str(SOURCE_EXCEL),
        operations=[
            {
                "action": "sort_rows",
                "sheet_name": "销售数据",
                "column": "销售额",
                "ascending": False,
            }
        ],
        output_path=str(OUTPUT_EXCEL),
    )

    if not result.get("success"):
        raise AssertionError("Excel 编辑执行失败。")

    output = load_workbook(OUTPUT_EXCEL, data_only=False)
    sheet = output["销售数据"]

    actual_order = [
        sheet.cell(row=row, column=1).value
        for row in range(2, 5)
    ]
    expected_order = ["澳门", "横琴", "珠海"]

    if actual_order != expected_order:
        raise AssertionError(
            f"销售额排序错误：{actual_order}"
        )

    for row_index in range(2, 5):
        current = capture_row(sheet, row_index)
        city = current["city"]
        before = source_by_city[city]

        if current["fill"] != before["fill"]:
            raise AssertionError(
                f"{city} 排序后填充色没有跟随业务行。"
            )

        if current["sales_number_format"] != before["sales_number_format"]:
            raise AssertionError(
                f"{city} 排序后销售额数字格式丢失。"
            )

        if current["bonus_number_format"] != before["bonus_number_format"]:
            raise AssertionError(
                f"{city} 排序后奖金数字格式丢失。"
            )

        if current["hyperlink"] != before["hyperlink"]:
            raise AssertionError(
                f"{city} 排序后超链接没有跟随业务行。"
            )

        if current["comment"] != before["comment"]:
            raise AssertionError(
                f"{city} 排序后批注没有跟随业务行。"
            )

        if current["height"] != before["height"]:
            raise AssertionError(
                f"{city} 排序后行高没有跟随业务行。"
            )

    print("业务行样式、超链接、批注、行高跟随检查通过。")

    print()
    print("[测试 2] 排序后公式语义保持正确")

    expected_bonus_formula = {
        2: "=B2*10%",
        3: "=B3*10%",
        4: "=B4*10%",
    }

    for row_index, expected_formula in expected_bonus_formula.items():
        actual_formula = sheet.cell(row=row_index, column=3).value
        if actual_formula != expected_formula:
            raise AssertionError(
                "排序后公式引用没有随新行位置调整："
                f"C{row_index}={actual_formula!r}，"
                f"期望 {expected_formula!r}"
            )

    print("公式语义检查通过。")

    print()
    print("[测试 3] 工作表级属性和未编辑 Sheet 保持")

    for column, width in source_widths.items():
        if sheet.column_dimensions[column].width != width:
            raise AssertionError(
                f"列宽 {column} 被破坏。"
            )

    if sheet.freeze_panes != source_freeze:
        raise AssertionError("冻结窗格被破坏。")

    if sheet.auto_filter.ref != source_filter:
        raise AssertionError("自动筛选范围被破坏。")

    if output["说明"]["A1"].value != source_note:
        raise AssertionError("未编辑 Sheet 内容被修改。")

    print("列宽、冻结窗格、筛选范围和未编辑 Sheet 保持通过。")

    print()
    print("[测试 4] 源文件保护")

    source_check = load_workbook(SOURCE_EXCEL, data_only=False)
    source_order = [
        source_check["销售数据"].cell(row=row, column=1).value
        for row in range(2, 5)
    ]

    if source_order != ["珠海", "澳门", "横琴"]:
        raise AssertionError("源 Excel 被覆盖。")

    print("源文件保护通过。")

    print()
    print("[测试 5] 合并单元格区域保持")

    merged_source = TEST_ROOT / "merged_source.xlsx"
    merged_output = TEST_ROOT / "merged_output.xlsx"

    merged_workbook = Workbook()
    merged_sheet = merged_workbook.active
    merged_sheet.title = "销售数据"
    merged_sheet.merge_cells("A1:E1")
    merged_sheet["A1"] = "2026年9月销售汇总"
    merged_sheet["A1"].font = Font(
        bold=True,
        name="微软雅黑",
        size=16,
    )
    merged_sheet["A1"].alignment = Alignment(
        horizontal="center",
    )

    merged_sheet.append(
        ["城市", "销售额", "奖金", "官网", "备注"]
    )
    merged_sheet.append(
        ["珠海", 128, "=B3*10%", "https://example.com/zhuhai", "正常"]
    )
    merged_sheet.append(
        ["澳门", 186, "=B4*10%", "https://example.com/macau", "冠军"]
    )
    merged_sheet.append(
        ["横琴", 154, "=B5*10%", "https://example.com/hengqin", "增长"]
    )
    merged_workbook.save(merged_source)

    try:
        apply_excel_edits(
            str(merged_source),
            operations=[
                {
                    "action": "sort_rows",
                    "sheet_name": "销售数据",
                    "column": "销售额",
                    "ascending": False,
                }
            ],
            output_path=str(merged_output),
        )
    except Exception as exc:
        raise AssertionError(
            "包含顶部合并标题的工作表排序失败。"
            "当前 Excel 编辑工具默认把第 1 行当表头，"
            "需要明确处理合并单元格/非首行表头场景。"
        ) from exc

    merged_check = load_workbook(
        merged_output,
        data_only=False,
    )
    merged_sheet_check = merged_check["销售数据"]

    if "A1:E1" not in {
        str(item)
        for item in merged_sheet_check.merged_cells.ranges
    }:
        raise AssertionError("排序后顶部合并单元格区域丢失。")

    if merged_sheet_check["A1"].value != "2026年9月销售汇总":
        raise AssertionError("排序后合并标题内容被破坏。")

    print("合并单元格区域保持通过。")

    print()
    print("=" * 70)
    print("DataPilot v3.4 Excel 高保真编辑基准测试通过！")
    print("=" * 70)


if __name__ == "__main__":
    main()
